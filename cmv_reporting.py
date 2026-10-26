"""Svenska CMV-rapporter med tydlig sammanfattning, täckningsstatus och indelgranskning."""
import csv
import os
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .cmv_clinical_annotations import classify
from .cmv_thresholds import MIN_ALT_FRACTION

from .filesystem import (
    ensure_dir, run_samples_folder, safe_report_filename,
    find_hbv_report_logo, image_file_to_data_uri,
)

STATUS_SV = {
    "COVERED": "Bedömbar enligt täckningskraven",
    "PARTIAL": "Delvis bedömbar",
    "NOT_COVERED": "Ej bedömbar – otillräcklig täckning",
    "EXCLUDED_BY_VARIANT_COMPLEXITY": "Ej bedömbar – komplex variantbild",
    "MIXED_FRAME_UNRESOLVED": "Läsramsstatus kan inte fastställas",
    "READ_LEVEL_REVIEW_INSUFFICIENT_CLEAN_SUPPORT": "Otillräckligt stöd i kvalitetsfiltrerade läsningar",
    "READ_PHASED_EXPLORATORY": "Stöd i läsdata – granskningsfynd",
    "BOTH_CALLERS": "Samstämmigt fynd i två analysmetoder",
    "LEGACY_WITH_READ_LEVEL_REVIEW": "Fynd från primäranalys, kontrollerat mot läsdata",
    "EXPLORATORY_UNPHASED_SNV": "Preliminärt substitutionsfynd, ej fasbestämt",
    "DOMINANT_REFERENCE_DIFFERENCE": "Dominerande variant jämfört med referensen",
    "MINORITY_ALLELE": "Mindre allel i provet",
    "NOT_RUN": "Ej körd",
    "DISABLED": "Avstängd",
    "REVIEW_ONLY": "Kandidatvarianter framtagna – granskning krävs",
    "REVIEW_FAILED": "Kunde inte slutföras",
    "NOT_RECORDED": "Inte registrerad",
    "UNKNOWN": "Okänd",
    "COMPLETED": "Slutförd",
    "HIGH": "Hög",
    "MEDIUM": "Måttlig",
    "LOW": "Låg",
    "PASS": "Godkänd",
    "FAIL": "Underkänd",
    "WARN": "Varning",
    "OK": "Godkänd",
    "UNPHASED_SNV": "SNV ej fasad",
    "EXPLORATORY_UNPHASED_SNV": "Preliminärt substitutionsfynd, ej fasbestämt",
    "READ_CODON_AUDIT_FAILED": "Läsbaserad kodonanalys misslyckades",
    "NOT_ASSESSED": "Ej bedömd",
    "ASSESSED_CLEAN_NO_ALT": "Bedömd; ingen alternativ variant i kvalitetsfiltrerade läsningar",
    "ASSESSED_CLEAN_WITH_ALT": "Alternativ variant stöds i kvalitetsfiltrerade läsningar",
    "NOT_ASSESSABLE_LOW_CLEAN_DEPTH": "Ej bedömbar – för få kvalitetsfiltrerade läsningar",
    "NO_ELIGIBLE_EVIDENCE": "Saknar tillräckligt bedömningsunderlag",
    "INVALID_REFERENCE_CODON": "Ogiltigt referenskodon",
}

HEADER_SV = {
    "Sample": "Prov-ID", "Mapped reads": "Mappade läsningar", "Read-yield QC": "Sekvenseringsutbyte",
    "Genome QC": "Genom-QC", "Selected reference": "Vald referens", "Reference coverage %": "Referenstäckning %",
    "Observed covered bases": "Observerade täckta baser", "Median depth over covered bases": "Medianläsdjup på täckta baser",
    "Covered bases >=30x": "Baser med minst 30× djup", "Covered-base breadth >=30x %": "Andel baser med minst 30× djup %",
    "Assessed positions >=30x": "Bedömda positioner vid minst 30×", "Minor SNVs >=10%": "Mindre SNV-fynd ≥10 %",
    "Interpretation": "Samlad tolkning", "Target QC status": "Mål-QC-status", "Target QC note": "Mål-QC-kommentar",
    "Assessment status": "Bedömningsstatus", "Protein variant records": "Proteinvarianter", "Published mutation matches": "Träffar i mutationspanelen",
    "BCFtools status": "BCFtools-status", "BCFtools candidate records": "BCFtools-kandidater", "BCFtools VCF": "BCFtools-VCF",
    "BCFtools manifest": "BCFtools-körningsfil", "Gene": "Gen", "Mutation": "Mutation", "Reference position": "Referensposition",
    "Depth": "Läsdjup", "Alt reads": "Läsningar med variant", "Frequency %": "Variantandel %", "Panel match": "Träff i panel",
    "Status": "Status", "Allele class": "Alleleklass", "Reference codon": "Referenskodon", "Alternate codon": "Alternativt kodon",
    "Reference accession": "Referens-ID", "Evidence source": "Evidenskälla", "Drug codes": "Antivirala läkemedel (koder)",
    "Source row": "Källrad", "Comment": "Kommentar", "Caller status": "Anroparstatus", "Position (1-based)": "Position (1-baserad)",
    "REF": "REF", "ALT": "ALT", "QUAL": "Kvalitetspoäng", "FILTER": "Filter", "Candidate TSV": "Kandidat-TSV",
    "Reference position (1-based)": "Referensposition (1-baserad)", "Major base": "Vanligaste bas", "Minor base": "Alternativ bas",
    "Minor reads": "Läsningar med alternativ bas", "Minor frequency %": "Andel alternativ bas %", "Reference position (1-based anchor)": "Referensposition (ankare, 1-baserad)",
    "Type": "Typ", "Length bp": "Längd (bp)", "Supporting reads": "Stödjande läsningar", "Eligible depth": "Bedömbart läsdjup",
    "Comparison": "Jämförelse", "Observed codon": "Observerat kodon", "Supporting reads": "Stödjande läsningar",
    "Frameshift reads excluded": "Läsningar exkluderade p.g.a. läsram", "Local-indel alt reads": "Variantläsningar med lokal indel",
    "Restored-frame alt reads": "Variantläsningar med återställd läsram", "Unresolved-frame alt reads": "Variantläsningar med oklart läsramsstöd",
    "Low-baseQ reads": "Läsningar med låg baskvalitet", "Incomplete reads": "Ofullständiga läsningar", "Reconstructed depth": "Återskapat läsdjup (granskning)",
    "Reconstructed alt reads": "Återskapade variantläsningar (granskning)", "Reconstructed alt %": "Återskapad variantandel % (granskning)",
    "Reconstructed threshold passed": "Återskapad tröskel passerad (inte ett eget variantfynd)", "Codon": "Kodon", "Eligible clean depth": "Kvalitetsfiltrerat läsdjup",
    "Clean reference-AA reads": "Kvalitetsfiltrerade referensläsningar", "Local-indel reads": "Läsningar med lokal indel",
    "Restored-frame reads": "Läsningar med återställd läsram", "Unresolved-frame reads": "Läsningar med oklart läsramsstöd",
    "Assessment": "Bedömning", "Reconstructed reference-AA reads": "Återskapade referensläsningar (granskning)",
    "Bases at >=30x": "Baser med minst 30× djup", "CDS length": "Kodande sekvens (bp)", "CDS coverage %": "Täckning av kodande sekvens %",
    "Covered codons >=30x": "Kodon med minst 30× djup", "Excluded codons": "Exkluderade kodon", "Mixed-frame flagged codons": "Kodon med oklart läsramsstöd",
    "Assessable codons": "Bedömbara kodon", "Total codons": "Totalt antal kodon",
}


def _text(value):
    return "" if value is None else str(value)


def _status(value):
    raw = _text(value)
    if not raw:
        return "Ej angivet"
    if raw in STATUS_SV:
        return STATUS_SV[raw]
    # Translate known status prefix while preserving any useful error detail.
    for key, sv in STATUS_SV.items():
        if raw.startswith(key + ":"):
            return sv + raw[len(key):]
    raw = raw.replace("NOT ASSESSED", "Ej bedömd").replace("EXPLORATORY ONLY", "Utforskande analys")
    phrase_map = {
        "protein SNVs assessed where complete codons have >=30x; no clinical resistance conclusion.":
            "Protein-SNV bedöms endast vid kompletta kodon med minst 30× läsdjup; ingen slutsats om läkemedelsresistens dras.",
        "Exploratory codon-level SNV screening; coordinate TSV must be independently verified.":
            "Utforskande SNV-analys på kodonnivå; genkoordinaterna ska verifieras separat.",
        "Local indel-affected codons excluded; downstream mixed-frame candidates flagged and withheld from published mutation matches. No phased haplotype or indel protein effects.":
            "Kodon nära indels exkluderas; nedströms fynd med oklart läsramsstöd markeras och räknas inte som panelträffar. Haplotype och indelns proteineffekt är inte fastställda.",
        "No reference-specific CMV mutation panel supplied; resistance interpretation unavailable.":
            "Ingen referensspecifik CMV-mutationspanel har angetts; resistens tolkas därför inte.",
    }
    for old, new in phrase_map.items():
        raw = raw.replace(old, new)
    # Convert remaining internal labels into readable text instead of exposing snake_case codes.
    if "_" in raw and raw.upper() == raw:
        return raw.replace("_", " ").capitalize()
    return raw


def _header(value):
    return HEADER_SV.get(value, value)


def _bcftools_tsv_path(row):
    explicit = row.get("cmv_bcftools_candidates_tsv")
    if explicit and Path(str(explicit)).is_file():
        return Path(explicit)
    vcf = row.get("cmv_bcftools_vcf")
    if vcf:
        guess = Path(str(vcf)).with_name(f"{row.get('sample', 'sample')}.merlin.bcftools.candidates.tsv")
        if guess.is_file():
            return guess
    return None


def _read_bcftools_rows(row):
    path = _bcftools_tsv_path(row)
    if not path:
        return []
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle, delimiter="\t"))
    except (OSError, csv.Error):
        return []


def _protein_records(row):
    records, seen = [], set()
    for item in row.get("cmv_protein_variant_calls", []):
        rec = dict(item)
        rec.setdefault("evidence_source", "Proteinvariantanalys")
        key = (str(rec.get("gene", "")).upper(), str(rec.get("mutation", "")).upper(), str(rec.get("position", "")))
        seen.add(key)
        records.append(rec)
    for item in row.get("cmv_read_codon_calls", []):
        mutation = str(item.get("mutation", "")).strip()
        if mutation:
            key = (str(item.get("gene", "")).upper(), mutation.upper(), str(item.get("position", "")))
            if key not in seen:
                rec = dict(item)
                rec.setdefault("evidence_source", "Granskning av kodon på läsnivå")
                rec.setdefault("status", rec.get("comparison", "READ_LEVEL"))
                records.append(rec)
                seen.add(key)
    return records


def _html_table(headers, rows, empty_text="Inga rader att visa.", min_width=760):
    head = "".join(f"<th>{escape(_text(h))}</th>" for h in headers)
    body = []
    for row in rows:
        cells = []
        for i, value in enumerate(row):
            display = escape(_text(value))
            if headers and headers[i] == "Bedömning" and str(value).startswith("Resistensassocierad"):
                display = f'<span class="assessment assessment-resistance">{display}</span>'
            elif headers and headers[i] == "Bedömning" and ("ej fastställd" in str(value).lower() or "ingen säkerställd" in str(value).lower() or "ofullständigt" in str(value).lower()):
                display = f'<span class="assessment assessment-uncertain">{display}</span>'
            elif headers and headers[i] == "Bedömning":
                display = f'<span class="assessment assessment-neutral">{display}</span>'
            cells.append(f"<td>{display}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    if not body:
        body = [f'<tr><td colspan="{len(headers)}" class="empty">{escape(empty_text)}</td></tr>']
    return f'<div class="table-wrap"><table style="--table-min:{min_width}px"><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'


def _style_workbook(workbook):
    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="185B8D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_cells in ws.columns:
            letter = column_cells[0].column_letter
            width = min(max(len(_text(c.value)) for c in column_cells) + 2, 46)
            ws.column_dimensions[letter].width = max(12, width)
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def _indel_summary(item):
    try:
        length = abs(int(item.get("length", 0)))
    except (TypeError, ValueError):
        length = 0
    kind = _text(item.get("type", "indel")).lower()
    if length and length % 3:
        consequence = "Kan förskjuta läsramen och ändra efterföljande aminosyror om händelsen ligger i en kodande region."
        tag = "Möjlig läsramspåverkan"
    elif length:
        consequence = "Ändrar sekvenslängden med hela kodon; effekten beror på exakt position och gen."
        tag = "Möjlig proteinpåverkan"
    else:
        consequence = "Konsekvensen kan inte bedömas från den här tabellen ensam."
        tag = "Behöver granskas"
    return tag, consequence


def _indel_review_priority(item):
    try:
        freq = float(item.get("frequency", 0) or 0)
        support = int(item.get("alt_reads", 0) or 0)
        depth = int(item.get("depth", 0) or 0)
    except (ValueError, TypeError):
        freq, support, depth = 0, 0, 0
    if support >= 10 and freq >= 50 and depth >= 30:
        return "Prioritera granskning: hög andel och flera stödjande läsningar"
    if support >= 5 and freq >= 15 and depth >= 30:
        return "Granska: stöd i flera läsningar"
    return "Kontrollera visuellt i alignment"


def _indel_explanation(item):
    tag, consequence = _indel_summary(item)
    return f"{tag}. {consequence} Detta är en läsbaserad signal; gen/proteinkonsekvens är inte fastställd."


DRUG_NAMES = {"GCV":"Ganciklovir", "VGCV":"Valganciklovir", "FOS":"Foskarnet", "CDV":"Cidofovir", "LET":"Letermovir", "MBV":"Maribavir"}

def _drug_label(value):
    """Show only interpretable panel codes; never infer susceptibility from malformed cells."""
    import re
    raw = str(value or '').strip()
    if not raw:
        return "Läkemedelsassociation saknas i källpanelen"
    entries = re.findall(r"(?i)(?<![A-Z])(GCV|VGCV|FOS|CDV|LET|MBV)\s*:\s*([A-Z0-9.+-]+)", raw)
    valid = []
    for drug, code in entries:
        # The historical panel's R/K and other symbols have not been clinically validated.
        if code.upper() in {"R", "K", "S", "I"}:
            valid.append(f"{DRUG_NAMES[drug.upper()]} (historisk panelkod {code.upper()}; ej validerad)")
    if valid:
        return "; ".join(dict.fromkeys(valid))
    if re.fullmatch(r"(?i)\s*(?:GCV|VGCV|FOS|CDV|LET|MBV)(?:\s*[;,/]\s*(?:GCV|VGCV|FOS|CDV|LET|MBV))*\s*", raw):
        return ", ".join(dict.fromkeys(DRUG_NAMES[x.upper()] for x in re.split(r"\s*[;,/]\s*", raw)))
    return "Ej fastställd – otydlig läkemedelsannotation i källpanelen"


def _clinical_coverage(coverage):
    lines = []
    for gene in ("UL97", "UL54", "UL56"):
        item = next((g for g in coverage if str(g.get("gene", "")).upper() == gene), None)
        if not item:
            lines.append(f"{gene}: Ej bedömbar – täckningsuppgift saknas")
            continue
        try:
            n, total = int(item.get("assessable_codons_30x")), int(item.get("total_codons"))
            if total <= 0 or n <= 0:
                lines.append(f"{gene}: Ej bedömbar – inga bedömbara kodon")
            elif n / total < .05:
                lines.append(f"{gene}: Mycket begränsad analys ({n}/{total} kodon); ej tillförlitlig för genövergripande slutsatser")
            elif n < total:
                lines.append(f"{gene}: Delvis bedömbar ({n}/{total} kodon)")
            else:
                lines.append(f"{gene}: Bedömbara kodon {n}/{total}")
        except (ValueError, TypeError):
            lines.append(f"{gene}: Ej bedömbar – otillräcklig QC-information")
    return "; ".join(lines)


def _excluded_candidates(row, published):
    accepted = {(str(p.get("gene", "")).upper(), str(p.get("mutation", "")).upper()) for p in published}
    excluded = {}
    for p in row.get("cmv_protein_variant_calls", []):
        key = (str(p.get("gene", "")).upper(), str(p.get("mutation", "")).upper())
        if p.get("panel_match") and key not in accepted:
            excluded[key] = "Otillräckligt stöd i kvalitetsfiltrerade läsningar"
    return sorted((g, m, reason) for (g, m), reason in excluded.items())


def _target_qc_status(value):
    raw = str(value or '').strip()
    if 'target qc not validated' in raw.lower() or 'not validated' in raw.lower():
        return 'Målregions-QC ej validerad'
    return _status(value)


def _target_qc_note(value):
    raw = str(value or "").strip()
    if any(x in raw.lower() for x in ("target qc not validated", "no assay bed", "whole-genome qc verdict is retained", "expected amplicon intervals are unknown")):
        return ("Målregions-QC är ännu inte validerad. Täckning och läsdjup redovisas per gen. "
                "Analysens fullständighet kan inte fastställas utan definierade målregioner och validerade kvalitetskriterier.")
    return raw

def _drug_assessments(published, coverage):
    """Conservative drug-level summary; no negative calls from gene breadth alone."""
    dependencies = [
        ("Ganciklovir / valganciklovir", ("UL97", "UL54")),
        ("Foskarnet", ("UL54",)),
        ("Cidofovir", ("UL54",)),
        ("Maribavir", ("UL97",)),
        ("Letermovir", ("UL56", "UL51", "UL89")),
    ]
    coverage_map = {str(g.get("gene", "")).upper(): g for g in coverage}
    out = []
    for drug, genes in dependencies:
        positives = []
        for hit in published:
            ann = classify(hit.get("gene"), hit.get("mutation"))
            if ann["category"] != "resistance":
                continue
            drug_text = ann["drugs"].lower()
            synonyms = {"Ganciklovir / valganciklovir": ("ganciklovir", "valganciklovir")}
            tokens = synonyms.get(drug, (drug.lower(),))
            if any(token in drug_text for token in tokens):
                positives.append(f"{hit.get('gene')} {hit.get('mutation')}")
        if positives:
            out.append((drug, "Resistensassocierad mutation påvisad", ", ".join(positives), "resistance"))
            continue
        observed = [coverage_map.get(g) for g in genes]
        if not any(observed):
            status, reason = "Ej bedömbar", "Täckningsuppgifter för relevanta gener saknas"
        elif any(not x or not int(x.get("assessable_codons_30x") or 0) for x in observed):
            status, reason = "Ej bedömbar", "En eller flera relevanta gener saknar bedömbara kodon"
        else:
            status, reason = "Ofullständigt bedömd", "Ingen resistensassocierad mutation påvisad i rapporterade fynd; fullständig positionsbaserad bedömning saknas"
        out.append((drug, status, reason, "uncertain"))
    return out


def _indel_clinical_relevance(indels, coverage):
    """Escalate only explicit, independently annotated clinical effects or QC impact.

    Raw indel coordinates alone are not sufficient to assign a gene or resistance.
    """
    relevant = []
    for item in indels:
        if item.get("clinical_association_verified") is True and item.get("clinical_association") and item.get("clinical_source"):
            relevant.append(("association", item))
        elif item.get("resistance_region_qc_impact") is True:
            relevant.append(("qc", item))
    return relevant


def _indel_clinical_note(indels, coverage):
    relevant = _indel_clinical_relevance(indels, coverage)
    if not relevant:
        # No unsupported clinical interpretation from a raw CIGAR indel.
        return "Inga indelfynd med dokumenterad klinisk resistensassociation eller verifierad påverkan på resistensanalysens bedömbarhet har rapporterats."
    notes = []
    for kind, item in relevant:
        position = _text(item.get("position", "okänd"))
        if kind == "association":
            notes.append(f"Position {position}: {_text(item['clinical_association'])} (källa: {_text(item['clinical_source'])}).")
        else:
            notes.append(f"Position {position}: påverkar bedömbarheten av en resistensrelevant region.")
    return " ".join(notes)


def _clinical_result_summary(published):
    resistant = []
    neutral = []
    for hit in published:
        ann = classify(hit.get("gene"), hit.get("mutation"))
        name = f"{hit.get('gene')} {hit.get('mutation')}"
        if ann["category"] == "resistance":
            resistant.append(f"{name} ({ann['drugs']})")
        elif ann["category"] == "polymorphism":
            neutral.append(name)
    sentences = []
    if resistant:
        sentences.append("Resistensassocierad mutation påvisad: " + ", ".join(resistant) + ".")
    else:
        sentences.append("Ingen kuraterad resistensassocierad mutation påvisad bland rapporterade varianter.")
    if neutral:
        sentences.append("Övriga kända polymorfismer: " + ", ".join(neutral) + ".")
    sentences.append("Bedömningen per läkemedel och täckningsbegränsningar redovisas ovan.")
    return " ".join(sentences)


CMV_RESULT_GENES = ("UL97", "UL54", "UL56", "UL89", "UL51")


def _result_number(value):
    try:
        return float(value) if value is not None and str(value).strip() else None
    except (ValueError, TypeError):
        return None


def _simple_result_rows(row):
    """All genes, all observed protein substitutions; no unassessed gene is called negative.

    Protein-level records are the primary summary. Read-codon evidence is attached
    by gene, mutation and codon where available; read-only mutations are retained.
    """
    sample = str(row.get("sample", ""))
    published = {(str(x.get("gene", "")).upper(), str(x.get("mutation", "")).upper())
                 for x in (row.get("cmv_protein_mutation_matches") or [])}
    coverage = {str(x.get("gene", "")).upper(): x
                for x in (row.get("cmv_gene_coverage") or [])}
    audits = {}
    for item in (row.get("cmv_read_codon_calls") or []):
        gene = str(item.get("gene", "")).upper()
        mutation = str(item.get("mutation", "")).upper()
        codon = str(item.get("alternate_codon", "")).upper()
        if gene and mutation:
            audits[(gene, mutation, codon)] = item
    seen = set()
    output = []
    for gene in CMV_RESULT_GENES:
        gene_items = []
        for item in (row.get("cmv_protein_variant_calls") or []):
            if str(item.get("gene", "")).upper() != gene:
                continue
            mutation = str(item.get("mutation", "")).upper()
            codon = str(item.get("alternate_codon", "")).upper()
            if not mutation or (gene, mutation, codon) in seen:
                continue
            seen.add((gene, mutation, codon))
            gene_items.append((item, audits.get((gene, mutation, codon))))
        for (agene, mutation, codon), audit in audits.items():
            if agene == gene and (gene, mutation, codon) not in seen:
                seen.add((gene, mutation, codon))
                gene_items.append((None, audit))
        cov = coverage.get(gene)
        assessable = _result_number(cov.get("assessable_codons_30x")) if cov else None
        total = _result_number(cov.get("total_codons")) if cov else None
        coverage_text = (f"{int(assessable)}/{int(total)} kodon" if assessable is not None and total is not None
                         else "Täckningsuppgift saknas")
        if not gene_items:
            if assessable is None or assessable <= 0:
                finding = "Ej bedömbar – inga rapporterade varianter"
            elif total is not None and assessable < total:
                finding = "Inga rapporterade varianter – delvis bedömbar"
            else:
                finding = "Inga rapporterade varianter i bedömda positioner"
            output.append([sample, gene, "–", None, "–", finding, "–", "–", "–", coverage_text])
            continue
        for protein, audit in sorted(gene_items, key=lambda pair: (str((pair[0] or pair[1]).get("mutation", "")), str((pair[0] or pair[1]).get("alternate_codon", "")))):
            item = protein or audit
            mutation = str(item.get("mutation", ""))
            freq = _result_number(item.get("frequency"))
            alt = item.get("alt_reads")
            depth = item.get("depth")
            support = f"{alt}/{depth}" if alt is not None and depth is not None else "–"
            clean = (f"{audit.get('alt_reads', 0)}/{audit.get('depth', 0)}" if audit else "–")
            reconstructed = (f"{audit.get('reconstructed_alt_reads', 0)}/{audit.get('reconstructed_depth', 0)}" if audit else "–")
            confirmed = (gene, mutation.upper()) in published
            panel = bool(protein.get("panel_match")) if protein else False
            if confirmed:
                assessment = "Panelträff – uppfyller läsnivåkrav"
            elif panel:
                assessment = "Panelkandidat – ej bekräftad enligt kraven"
            elif audit and audit.get("status") == "READ_PHASED_EXPLORATORY":
                assessment = "Påvisad proteinvariant – ej klassad resistens"
            else:
                assessment = "Proteinvariant – granskning krävs"
            output.append([sample, gene, mutation, freq, support, assessment,
                           "Ja" if panel or confirmed else "Nej/ej angivet", clean,
                           reconstructed, coverage_text])
    return output


def generate_cmv_reports(rows, output_folder, run_tag):
    """Write Swedish per-run Excel and per-sample HTML CMV reports."""
    directory = os.path.join(run_samples_folder(output_folder, run_tag), "CMV_reports")
    ensure_dir(directory, must_be_writable=True)
    workbook = Workbook()
    results_ws = workbook.active
    results_ws.title = "Resultat"
    results_ws.append(["Prov-ID", "Gen", "Mutation", "Variantandel %", "Variantläsningar/djup",
                       "Bedömning", "Träff i panel", "Kvalitetsfiltrerade läsningar",
                       "Återskapade läsningar (granskning)", "Bedömbara kodon"])
    summary = workbook.create_sheet("Sammanfattning")
    summary.append(["Prov-ID", "Mappade läsningar", "Sekvenseringsutbyte", "Genom-QC", "Vald referens", "Referenstäckning %", "Täckta baser", "Medianläsdjup", "Baser ≥30×", "Bredd ≥30× %", "Bedömda positioner ≥30×", "Mindre SNV-fynd", "Samlad tolkning", "Mål-QC-status", "Mål-QC-kommentar", "Bedömningsstatus", "Proteinvarianter", "Träffar i mutationspanelen", "BCFtools-status", "BCFtools-kandidater", "BCFtools-VCF", "BCFtools-körningsfil"])
    priority_ws = workbook.create_sheet("Resistensfynd", 1)
    priority_ws.append(["Prov-ID", "Gen", "Mutation", "Läkemedel", "Variantandel %", "Läsdjup", "Status", "Kommentar", "Källrad", "Klinisk bedömning", "Källa (PMID/DOI)"])
    drug_ws = workbook.create_sheet("Läkemedelsbedömning", 1)
    drug_ws.append(["Prov-ID", "Läkemedel", "Bedömning", "Underlag"])
    excluded_ws = workbook.create_sheet("Ej rapporterade panelkandidater")
    excluded_ws.append(["Prov-ID", "Gen", "Mutation", "Orsak"])
    protein_ws = workbook.create_sheet("Alla proteinvarianter")
    protein_ws.append(["Prov-ID", "Gen", "Mutation", "Referensposition", "Läsdjup", "Variantläsningar", "Variantandel %", "Träff i panel", "Status", "Alleleklass", "Referenskodon", "Alternativt kodon", "Referens-ID", "Evidenskälla"])
    published_ws = workbook.create_sheet("Resistenspanel")
    published_ws.append(["Prov-ID", "Gen", "Mutation", "Variantandel %", "Läsdjup", "Antivirala läkemedel (koder)", "Källrad", "Status", "Kommentar", "Referenskodon", "Alternativt kodon", "Variantläsningar", "Referens-ID"])
    bcf_ws = workbook.create_sheet("BCFtools-varianter")
    bcf_ws.append(["Prov-ID", "Anroparstatus", "Referens", "Position (1-baserad)", "REF", "ALT", "Kvalitetspoäng", "Filter", "VCF", "Kandidat-TSV", "Körningsfil"])
    snv_ws = workbook.create_sheet("Mindre SNV")
    snv_ws.append(["Prov-ID", "Referensposition (1-baserad)", "Vanligaste bas", "Alternativ bas", "Läsdjup", "Alternativa läsningar", "Variantandel %", "Tolkning"])
    indel_ws = workbook.create_sheet("Insertioner och deletioner")
    indel_ws.append(["Prov-ID", "Position (ankare, 1-baserad)", "Typ", "REF", "ALT", "Längd (bp)", "Läsdjup", "Stödjande läsningar", "Variantandel %", "Bedömningsrelevans", "Kort förklaring", "Prioritet för granskning"])
    audit_ws = workbook.create_sheet("Kodon på läsnivå")
    audit_ws.append(["Prov-ID", "Gen", "Mutation", "Jämförelse", "Referenskodon", "Observerat kodon", "Bedömbart läsdjup", "Stödjande läsningar", "Variantandel %", "Exkluderade läsningar p.g.a. läsram", "Lokal indel", "Återställd läsram", "Oklart läsramsstöd", "Låg baskvalitet", "Ofullständiga läsningar", "Status", "Återskapat läsdjup (granskning)", "Återskapade variantläsningar (granskning)", "Återskapad variantandel % (granskning)", "Återskapad tröskel passerad (inte ett eget fynd)"])
    detail_ws = workbook.create_sheet("Diagnostiskt läsunderlag")
    detail_ws.append(["Prov-ID", "Gen", "Mutation", "Referenskodon", "Observerat kodon", "Kvalitetsfiltrerat läsdjup", "Alternativa läsningar", "Variantandel %", "Lokal indel", "Återställd läsram", "Oklart läsramsstöd", "Låg baskvalitet", "Ofullständiga läsningar", "Finns i äldre resultat", "Godkänd läsnivå", "Status", "Återskapat läsdjup (granskning)", "Återskapade variantläsningar (granskning)", "Återskapad variantandel % (granskning)", "Återskapad tröskel passerad (inte ett eget fynd)"])
    qc_ws = workbook.create_sheet("Kodon-QC")
    qc_ws.append(["Prov-ID", "Gen", "Kodon", "Referenskodon", "Kvalitetsfiltrerat djup", "Referensläsningar", "Läsningar med lokal indel", "Återställd läsram", "Oklart läsramsstöd", "Låg baskvalitet", "Ofullständiga läsningar", "Bedömning", "Återskapat djup (granskning)", "Återskapade referensläsningar (granskning)"])
    coverage_ws = workbook.create_sheet("Gentäckning")
    coverage_ws.append(["Prov-ID", "Gen", "Baser ≥30×", "Kodande sekvens (bp)", "Täckning %", "Kodon ≥30×", "Exkluderade kodon", "Kodon med oklart läsramsstöd", "Bedömbara kodon", "Totalt antal kodon", "Status"])

    for row in rows:
        sample = row.get("sample", "")
        for result_line in _simple_result_rows(row):
            results_ws.append(result_line)
        calls = row.get("cmv_variant_calls", [])
        proteins = _protein_records(row)
        published = row.get("cmv_protein_mutation_matches", [])
        published = sorted(published, key=lambda p: ({"UL97": 0, "UL54": 1, "UL56": 2}.get(str(p.get("gene", "")).upper(), 3), str(p.get("mutation", ""))))
        for hit in published:
            priority_ws.append([sample, hit.get("gene", ""), hit.get("mutation", ""), classify(hit.get("gene"), hit.get("mutation"))["drugs"], hit.get("frequency", ""), hit.get("depth", ""), _status(hit.get("status", "")), hit.get("comment", ""), hit.get("source_row", ""), classify(hit.get("gene"), hit.get("mutation"))["label"], classify(hit.get("gene"), hit.get("mutation"))["source"]])
        drug_results = _drug_assessments(published, row.get("cmv_gene_coverage", []))
        for drug, assessment, reason, _ in drug_results:
            drug_ws.append([sample, drug, assessment, reason])
        excluded_candidates = _excluded_candidates(row, published)
        for gene, mutation, reason in excluded_candidates:
            excluded_ws.append([sample, gene, mutation, reason])
        bcf_rows = _read_bcftools_rows(row)
        vcf_path = str(row.get("cmv_bcftools_vcf", ""))
        candidates_tsv = str(_bcftools_tsv_path(row) or "")
        manifest = str(row.get("cmv_bcftools_manifest", ""))
        indels = sorted(row.get("cmv_indel_calls", []), key=lambda x: (int(x.get("alt_reads", 0) or 0), float(x.get("frequency", 0) or 0)), reverse=True)
        bcf_status = row.get("cmv_bcftools_status", "NOT_RUN")
        summary.append([sample, row.get("mapped_reads", ""), _status(row.get("cmv_read_yield_qc", "UNKNOWN")), _status(row.get("qc_verdict", "")), row.get("best_reference", ""), row.get("reference_coverage_percent", row.get("target_or_genome_coverage_percent", "")), row.get("cmv_covered_bases", ""), row.get("cmv_covered_median_depth", ""), row.get("cmv_covered_bases_30x", ""), row.get("cmv_covered_bases_30x_percent", ""), row.get("cmv_assessed_positions", 0), len(calls), _status(row.get("cmv_interpretation", "")), _target_qc_status(row.get("cmv_qc_status", "NOT_RECORDED")), _target_qc_note(row.get("cmv_qc_note", "")), _status(row.get("cmv_assessment_status", "")), len(proteins), len(published), _status(bcf_status), row.get("cmv_bcftools_candidates", len(bcf_rows)), vcf_path, manifest])
        for item in proteins:
            protein_ws.append([sample, item.get("gene", ""), item.get("mutation", ""), item.get("position", ""), item.get("depth", ""), item.get("alt_reads", ""), item.get("frequency", ""), "Ja" if item.get("panel_match") else "Nej", _status(item.get("status", "")), _status(item.get("allele_class", "")), item.get("reference_codon", ""), item.get("alternate_codon", ""), row.get("cmv_resistance_reference", ""), item.get("evidence_source", "Proteinvariantanalys")])
        for item in published:
            published_ws.append([sample, item.get("gene", ""), item.get("mutation", ""), item.get("frequency", ""), item.get("depth", ""), item.get("drug_codes", ""), item.get("source_row", ""), _status(item.get("status", "")), item.get("comment", ""), item.get("reference_codon", ""), item.get("alternate_codon", ""), item.get("alt_reads", ""), row.get("cmv_resistance_reference", "")])
        if bcf_rows:
            for item in bcf_rows:
                bcf_ws.append([sample, _status(bcf_status), item.get("chrom", "NC_006273.2"), item.get("position", ""), item.get("ref", ""), item.get("alt", ""), item.get("qual", ""), item.get("filter", ""), vcf_path, candidates_tsv, manifest])
        else:
            bcf_ws.append([sample, _status(bcf_status), "NC_006273.2", "", "", "", "", "", vcf_path, candidates_tsv, manifest])
        for item in calls:
            snv_ws.append([sample, item.get("position"), item.get("major"), item.get("alt"), item.get("depth"), item.get("alt_reads"), item.get("frequency"), item.get("interpretation")])
        for item in row.get("cmv_indel_calls", []):
            tag, explanation = _indel_summary(item)
            indel_ws.append([sample, item.get("position"), item.get("type"), item.get("ref"), item.get("alt"), item.get("length"), item.get("depth"), item.get("alt_reads"), item.get("frequency"), tag, _indel_explanation(item), _indel_review_priority(item)])
        for item in row.get("cmv_read_codon_calls", []):
            audit_ws.append([sample, item.get("gene"), item.get("mutation"), _status(item.get("comparison")), item.get("reference_codon"), item.get("alternate_codon"), item.get("depth"), item.get("alt_reads"), item.get("frequency"), item.get("frameshift_reads_excluded"), item.get("local_indel_support"), item.get("restored_frame_support"), item.get("unresolved_frame_support"), item.get("low_baseq_reads"), item.get("incomplete_reads"), _status(item.get("status")), item.get("reconstructed_depth"), item.get("reconstructed_alt_reads"), item.get("reconstructed_frequency"), item.get("reconstructed_threshold_pass")])
        for item in row.get("cmv_read_codon_diagnostics", []):
            detail_ws.append([sample, item.get("gene"), item.get("mutation"), item.get("reference_codon"), item.get("alternate_codon"), item.get("depth"), item.get("alt_reads"), item.get("frequency"), item.get("local_indel_support"), item.get("restored_frame_support"), item.get("unresolved_frame_support"), item.get("low_baseq_reads"), item.get("incomplete_reads"), item.get("legacy_present"), item.get("read_level_pass"), _status(item.get("status")), item.get("reconstructed_depth"), item.get("reconstructed_alt_reads"), item.get("reconstructed_frequency"), item.get("reconstructed_threshold_pass")])
        for item in row.get("cmv_read_codon_qc", []):
            qc_ws.append([sample, item.get("gene"), item.get("codon"), item.get("reference_codon"), item.get("clean_depth"), item.get("clean_reference_reads"), item.get("local_indel_reads"), item.get("restored_frame_reads"), item.get("unresolved_frame_reads"), item.get("low_baseq_reads"), item.get("incomplete_reads"), _status(item.get("assessment")), item.get("reconstructed_depth"), item.get("reconstructed_reference_reads")])
        gene_coverage = row.get("cmv_gene_coverage", [])
        for item in gene_coverage:
            coverage_ws.append([sample, item.get("gene"), item.get("covered_bases_30x"), item.get("length"), item.get("coverage_pct"), item.get("covered_codons_30x"), item.get("excluded_codons"), item.get("mixed_frame_codons", 0), item.get("assessable_codons_30x"), item.get("total_codons"), _status(item.get("status"))])

        logo_uri = image_file_to_data_uri(find_hbv_report_logo(row, os.path.join(directory, "report.html")))
        logo_html = f'<img class="report-logo" src="{logo_uri}" alt="Aurora-logotyp">' if logo_uri else ""
        coverage_rows = [[g.get("gene"), g.get("covered_bases_30x"), g.get("length"), g.get("coverage_pct"), g.get("covered_codons_30x"), g.get("excluded_codons"), g.get("mixed_frame_codons", 0), f'{g.get("assessable_codons_30x", "")}/{g.get("total_codons", "")}', _status(g.get("status"))] for g in gene_coverage]
        protein_rows = [[p.get("gene"), p.get("mutation"), p.get("position"), p.get("reference_codon"), p.get("alternate_codon"), p.get("depth"), p.get("alt_reads"), p.get("frequency"), "Ja" if p.get("panel_match") else "Nej", _status(p.get("status")), p.get("evidence_source", "Proteinvariantanalys")] for p in proteins]
        published_rows = []
        for p in published:
            ann = classify(p.get("gene"), p.get("mutation"))
            published_rows.append([p.get("gene"), p.get("mutation"), ann["drugs"],
                                   p.get("frequency"), f"{p.get('alt_reads', '–')}/{p.get('depth', '–')}",
                                   ann["label"], ann["source"]])
        audit_rows = [[a.get("gene"), a.get("mutation"), _status(a.get("comparison")), a.get("depth"), a.get("alt_reads"), a.get("frequency"), a.get("frameshift_reads_excluded"), a.get("local_indel_support"), a.get("restored_frame_support"), a.get("unresolved_frame_support"), _status(a.get("status"))] for a in row.get("cmv_read_codon_calls", [])]
        bcf_html_rows = [[x.get("chrom", "NC_006273.2"), x.get("position"), x.get("ref"), x.get("alt"), x.get("qual"), x.get("filter")] for x in bcf_rows]
        indel_html_rows = [[i.get("position"), "Insertion" if str(i.get("type", "")).lower().startswith("ins") else "Deletion", i.get("ref"), i.get("alt"), i.get("length"), i.get("alt_reads"), i.get("depth"), i.get("frequency"), _indel_summary(i)[0], _indel_explanation(i), _indel_review_priority(i)] for i in indels]
        minor_rows = [[c.get("position"), c.get("major"), c.get("alt"), c.get("depth"), c.get("frequency")] for c in calls]
        published_highlights = []
        for p in published[:8]:
            published_highlights.append(f"<li><strong>{escape(_text(p.get('gene')))} {escape(_text(p.get('mutation')))}</strong> – {escape(_text(p.get('frequency', '')))} % variantandel. {escape(_text(p.get('comment', 'Träff i den inlästa mutationspanelen.')))}</li>")
        if not published_highlights:
            published_highlights.append('<li>Inga träffar i den inlästa mutationspanelen. Bedöm detta tillsammans med gentäckning och antal bedömbara kodon.</li>')
        indel_highlights = []
        # Keep the front-page summary short; the complete event table remains below and in Excel.
        for item in sorted(indels, key=lambda x: (int(x.get("alt_reads", 0) or 0), float(x.get("frequency", 0) or 0)), reverse=True)[:5]:
            tag, consequence = _indel_summary(item)
            indel_highlights.append(f"<li><strong>Position {escape(_text(item.get('position')))}: {escape('insertion' if str(item.get('type','')).lower().startswith('ins') else 'deletion')}</strong>, {escape(_text(item.get('length')))} bp, {escape(_text(item.get('frequency')))} % ({escape(_text(item.get('alt_reads')))} av {escape(_text(item.get('depth')))} läsningar). <strong>{escape(tag)}.</strong> {escape(consequence)} <em>Gen/proteinkonsekvens är inte fastställd.</em></li>")
        if not indel_highlights:
            indel_highlights.append('<li>Inga insertioner eller deletioner passerade de aktuella urvalskraven. Det utesluter inte att sådana varianter finns.</li>')
        assay = _status(row.get("cmv_bcftools_status", "NOT_RUN"))
        qc_note = _target_qc_note(row.get("cmv_qc_note", ""))
        excluded_html = _html_table(["Gen", "Mutation", "Orsak till att fyndet inte rapporteras som panelträff"], excluded_candidates,
                                    "Inga ytterligare panelmatchande kandidater exkluderades.", 720)
        html = """<!doctype html><html lang="sv"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>CMV – __SAMPLE__</title><style>
:root{--blue:#dbeeff;--blue-dark:#185b8d;--line:#c9d3dc;--text:#1e2a32;--muted:#5d6b75;--pale:#f4f8fb}
*{box-sizing:border-box}body{font:15px/1.55 Arial,Helvetica,sans-serif;color:var(--text);margin:0;background:#eef2f5}
.report{max-width:1380px;width:calc(100% - 32px);margin:24px auto;background:#fff;padding:clamp(18px,3.5vw,42px);box-shadow:0 3px 18px #0002}
.header{display:flex;justify-content:space-between;align-items:flex-start;gap:24px;margin-bottom:20px}.title-block{min-width:0;flex:1}h1{margin:0 0 5px;color:#123e5f;font-size:clamp(23px,2.5vw,30px);line-height:1.22;overflow-wrap:normal}.subtitle,.small{color:var(--muted)}.report-logo{max-width:210px;max-height:90px;object-fit:contain}
h2{font-size:19px;line-height:1.3;margin:30px 0 10px;border-bottom:2px solid var(--blue-dark);padding-bottom:7px;break-after:avoid}
p{margin:8px 0 12px}.meta{width:100%;border-collapse:collapse;table-layout:auto;margin:12px 0 18px}.meta td{border:1px solid var(--line);padding:9px 11px;vertical-align:top}.meta td:first-child{width:30%;min-width:180px;background:var(--pale);font-weight:bold;white-space:normal}
.table-wrap{width:100%;overflow-x:auto;margin:10px 0 20px;border:1px solid var(--line);border-radius:6px}.table-wrap table{width:100%;min-width:var(--table-min,760px);border-collapse:collapse;font-size:13px;table-layout:auto;margin:0}.table-wrap th{background:var(--blue);text-align:left;font-weight:700;white-space:normal}.table-wrap th,.table-wrap td{border-bottom:1px solid var(--line);border-right:1px solid var(--line);padding:9px 10px;vertical-align:top;text-align:left;word-break:normal;overflow-wrap:normal;hyphens:none;line-height:1.4}.table-wrap tr:last-child td{border-bottom:0}.table-wrap th:last-child,.table-wrap td:last-child{border-right:0}.table-wrap td:first-child,.table-wrap th:first-child{white-space:nowrap}.table-wrap tbody tr:nth-child(even){background:#fbfcfd}.empty{color:var(--muted);font-style:italic}
.clinical{border:1px solid #c9d3dc;padding:14px 17px;border-radius:6px;margin:14px 0}.clinical h2{margin-top:0}.summary{background:#f4f8fb;border-left:5px solid var(--blue-dark);padding:14px 17px;margin:16px 0 20px;border-radius:0 6px 6px 0}.summary h2{border:0;margin:0 0 8px;padding:0;font-size:20px}.summary h3{font-size:15px;margin:14px 0 4px}.summary ul{margin:5px 0 0;padding-left:22px}.summary li{margin:7px 0}.summary-grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px;margin:14px 0}.metric{background:#fff;border:1px solid var(--line);border-radius:6px;padding:12px;min-width:0}.metric strong{display:block;font-size:20px;color:#123e5f}.metric span{font-size:12px;color:var(--muted)}.indel-note{background:#fff8e8;border-left:4px solid #c99b00;padding:10px 12px;margin:12px 0}.assessment{display:inline-block;padding:3px 8px;border-radius:5px;font-weight:700}.assessment-resistance{background:#fbe2e2;color:#8d1c1c}.assessment-uncertain{background:#fff0c2;color:#765100}.assessment-neutral{background:#edf1f5;color:#394b59}.footer{font-size:11px;color:var(--muted);margin-top:30px;border-top:1px solid var(--line);padding-top:12px}
@media(max-width:720px){.report{width:100%;margin:0;padding:18px}.header{gap:12px}.report-logo{max-width:100px;max-height:65px}.summary-grid{grid-template-columns:1fr}.meta td:first-child{min-width:125px;width:38%}h2{font-size:18px}}
@media print{body{background:#fff}.report{box-shadow:none;margin:0;width:100%;max-width:none;padding:12mm}.table-wrap{overflow:visible;border-radius:0}.table-wrap table{min-width:0!important;font-size:10px}.table-wrap th,.table-wrap td{padding:5px;word-break:normal;overflow-wrap:anywhere}.report-logo{max-height:65px}tr{break-inside:avoid}.summary{break-inside:avoid}}
</style></head><body><main class="report">
<div class="header"><div class="title-block"><h1>Cytomegalovirus – sekvensanalys</h1><div class="subtitle">Aurora · Körning __RUN__</div></div>__LOGO__</div>
<table class="meta"><tr><td>Prov-ID</td><td>__SAMPLE__</td></tr><tr><td>Körning</td><td>__RUN__</td></tr><tr><td>Vald aligneringsreferens</td><td>__REFERENCE__</td></tr><tr><td>Referens för resistenspanel</td><td>__RESREF__</td></tr><tr><td>Svarsstatus</td><td>__RELEASESTATUS__</td></tr></table>
<section class="clinical"><h2>Resistensbedömning per läkemedel</h2>__DRUGTABLE__<h2>Påvisade mutationer och bedömning (UL97, UL54, UL56)</h2>__PUBLISHEDTABLE__<h3>Sammanfattande bedömning</h3><p>__CLINICALSUMMARY__</p><p><strong>Genvis bedömbarhet:</strong> __GENEASSESS__</p><h3>Panelmatchande kandidater som inte uppfyller rapporteringskraven</h3><p>Fynden nedan uppfyller inte de aktuella rapporteringskriterierna.</p>__EXCLUDED__<p><strong>Indels och resistensanalys:</strong> __INDELCLINICAL__</p><p><strong>Analysbegränsning:</strong> Genpositioner med otillräckligt eller oklart läsramsstöd kan inte bedömas säkert.</p></section><h2>Teknisk analysinformation</h2><table class="meta"><tr><td>Vald referens</td><td>__REFERENCE__</td></tr><tr><td>Referens för mutationspanel</td><td>__RESREF__</td></tr><tr><td>Mappade läsningar</td><td>__MAPPED__</td></tr><tr><td>Sekvenseringsutbyte</td><td>__YIELD__</td></tr><tr><td>Helgenom-QC (teknisk information, ej mål-QC)</td><td>__GENOMEQC__</td></tr><tr><td>Mål-QC</td><td>__TARGETQC__</td></tr><tr><td>Referenstäckning</td><td>__REFCOV__ %</td></tr><tr><td>Täckta baser / medianläsdjup</td><td>__BASES__ / __MEDIAN__×</td></tr><tr><td>Baser med minst 30× djup</td><td>__BASES30__ (__PCT30__ %)</td></tr><tr><td>Bedömda positioner vid minst 30×</td><td>__ASSESSED__</td></tr><tr><td>Proteinvarianter / panelträffar</td><td>__PROTEINS__ / __PUBLISHED__</td></tr><tr><td>Oberoende BCFtools-analys</td><td>__BCFSTATUS__ – __BCFCOUNT__ kandidater</td></tr></table>
<section class="summary"><h2>Sammanfattning av viktiga fynd</h2><div class="summary-grid"><div class="metric"><strong>__PROTEINS__</strong><span>proteinvarianter rapporterade</span></div><div class="metric"><strong>__PUBLISHED__</strong><span>träffar i mutationspanelen</span></div><div class="metric"><strong>__INDELCOUNT__</strong><span>insertioner/deletioner för granskning</span></div></div>
<h3>Panelträffar (teknisk sammanfattning)</h3><p>De kliniskt prioriterade fynden redovisas ovan. Detta är en separat granskningsöversikt.</p><ul>__PUBLISHEDLIST__</ul>
<h3>Insertioner och deletioner att granska</h3><p>Indels kan påverka proteinsekvensen. En längd som inte är delbar med tre kan förskjuta läsramen om händelsen ligger i en kodande region. Position och gen måste fastställas innan den biologiska betydelsen kan bedömas.</p><ul>__INDELLIST__</ul>
<div class="indel-note"><strong>Hur indels prioriteras:</strong> listan sorteras efter antal stödjande läsningar och variantandel. Detta är granskningsprioritering, inte en klinisk klassificering. Kontrollera misstänkta fynd i alignment och bedöm deras påverkan på kodon och protein separat.</div></section>
<h2>Alla proteinvarianter</h2><p>Samtliga proteinvarianter listas här, även när de inte finns i den inlästa mutationspanelen. UL54 och UL97 visas tillsammans med övriga rapporterade gener.</p>__PROTEINTABLE__

<h2>Gentäckning och bedömbara kodon</h2><p>30× används som miniminivå för denna täckningsöversikt. <strong>Ej bedömbar</strong> betyder att tillräckliga data saknas enligt nuvarande krav; det betyder inte att genen saknar mutationer.</p>__COVERAGETABLE__
<h2>Insertioner och deletioner</h2><p>Fynden är signaler från läsningarnas alignering. De kan vara relevanta eftersom de ibland påverkar läsram eller proteinlängd, men effekten kan inte bestämmas från längd och frekvens ensam. Tabellen visar därför en kort granskningsförklaring och prioritet.</p>__INDELTABLE__
<h2>Oberoende BCFtools-variantkandidater</h2><p>Status: __BCFSTATUS__. Kandidaterna visas separat från Aurora's proteinvarianter och ändrar inte dessa automatiskt. Jämför med genkoordinater och läsbaserat stöd.</p>__BCFTABLE__
<h2>Jämförelse av kodon på läsnivå</h2>__AUDITTABLE__
<h2>Kodon-QC</h2>__QCTABLE__
<h2>Diagnostiskt läsunderlag</h2>__DETAILTABLE__
<h2>Mindre SNV-fynd</h2>__SNVTABLE__
<div class="footer">Aurora · CMV · Körning __RUN__ · Resultat och källfiler anges i rapporten och Excel-arbetsboken.</div></main></body></html>"""
        substitutions = {
            "__SAMPLE__": escape(_text(sample)), "__RUN__": escape(_text(run_tag)), "__LOGO__": logo_html,
            "__RELEASESTATUS__": escape(_text(row.get("cmv_release_status", "Ej signerad"))),
            "__DRUGTABLE__": _html_table(["Läkemedel", "Bedömning", "Underlag"], [[drug, status, reason] for drug, status, reason, category in drug_results], "Ingen läkemedelsbedömning tillgänglig.", 760),
            "__REFERENCE__": escape(_text(row.get("best_reference", ""))), "__RESREF__": escape(_text(row.get("cmv_resistance_reference", ""))),
            "__MAPPED__": escape(_text(row.get("mapped_reads", ""))), "__YIELD__": escape(_status(row.get("cmv_read_yield_qc", "UNKNOWN"))),
            "__GENOMEQC__": escape(_status(row.get("qc_verdict", ""))),
            "__TARGETQC__": escape(_target_qc_status(row.get("cmv_qc_status", "NOT_RECORDED")) + (" – " + qc_note if qc_note else "")),
            "__REFCOV__": escape(_text(row.get("target_or_genome_coverage_percent", row.get("reference_coverage_percent", "")))),
            "__BASES__": escape(_text(row.get("cmv_covered_bases", ""))), "__MEDIAN__": escape(_text(row.get("cmv_covered_median_depth", ""))),
            "__BASES30__": escape(_text(row.get("cmv_covered_bases_30x", ""))), "__PCT30__": escape(_text(row.get("cmv_covered_bases_30x_percent", ""))),
            "__ASSESSED__": escape(_text(row.get("cmv_assessed_positions", 0))), "__PROTEINS__": str(len(proteins)), "__PUBLISHED__": str(len(published)),
            "__BCFSTATUS__": escape(_status(bcf_status)), "__BCFCOUNT__": escape(_text(row.get("cmv_bcftools_candidates", len(bcf_rows)))),
            "__EXCLUDED__": excluded_html, "__INDELCLINICAL__": escape(_indel_clinical_note(indels, gene_coverage)), "__CLINICALSUMMARY__": escape(_clinical_result_summary(published)), "__GENEASSESS__": escape(_clinical_coverage(gene_coverage)), "__INDELCOUNT__": str(len(indels)), "__PUBLISHEDLIST__": "".join(published_highlights), "__INDELLIST__": "".join(indel_highlights),
            "__PROTEINTABLE__": _html_table(["Gen", "Mutation", "Referensposition", "Referenskodon", "Alternativt kodon", "Läsdjup", "Variantläsningar", "Variantandel %", "Panelträff", "Status", "Evidenskälla"], protein_rows, "Inga proteinvarianter rapporterades.", 1050),
            "__PUBLISHEDTABLE__": _html_table(["Gen", "Mutation", "Läkemedelsassociation", "Variantandel %", "Variantläsningar / djup", "Bedömning", "Källa (PMID)"], published_rows, "Inga panelträffar identifierade. Detta utesluter inte resistens vid ofullständig täckning.", 760),
            "__COVERAGETABLE__": _html_table(["Gen", "Baser ≥30×", "Kodande sekvens (bp)", "Täckning %", "Kodon ≥30×", "Exkluderade kodon", "Oklart läsramsstöd", "Bedömbara / totalt kodon", "Status"], coverage_rows, "Gentäckningsdata saknas.", 940),
            "__INDELTABLE__": _html_table(["Position", "Typ", "REF", "ALT", "Längd (bp)", "Stödjande läsningar", "Läsdjup", "Variantandel %", "Möjlig betydelse", "Kort förklaring", "Granskningsprioritet"], indel_html_rows, "Inga indels passerade de aktuella urvalskraven.", 1220),
            "__BCFTABLE__": _html_table(["Referens", "Position", "REF", "ALT", "Kvalitetspoäng", "Filter"], bcf_html_rows, "Inga BCFtools-kandidater tillgängliga. Kontrollera analysstatus och filsökväg.", 700),
            "__AUDITTABLE__": _html_table(["Gen", "Mutation", "Jämförelse", "Bedömbart djup", "Variantläsningar", "Variantandel %", "Exkluderade läsningar p.g.a. läsram", "Lokal indel", "Återställd läsram", "Oklart läsramsstöd", "Status"], audit_rows, "Inga kodonfynd från läsnivåjämförelsen.", 1200),
            "__QCTABLE__": _html_table(["Gen", "Kodon", "Referenskodon", "Kvalitetsfiltrerat djup", "Referensläsningar", "Lokal indel", "Återställd läsram", "Oklart läsramsstöd", "Låg baskvalitet", "Ofullständiga läsningar", "Bedömning"], [[q.get("gene"),q.get("codon"),q.get("reference_codon"),q.get("clean_depth"),q.get("clean_reference_reads"),q.get("local_indel_reads"),q.get("restored_frame_reads"),q.get("unresolved_frame_reads"),q.get("low_baseq_reads"),q.get("incomplete_reads"),_status(q.get("assessment"))] for q in row.get("cmv_read_codon_qc", [])], "Ingen kodon-QC tillgänglig.", 1150),
            "__DETAILTABLE__": _html_table(["Gen", "Mutation", "Referenskodon", "Observerat kodon", "Kvalitetsfiltrerat djup", "Variantläsningar", "Variantandel %", "Lokal indel", "Återställd läsram", "Oklart läsramsstöd", "Låg baskvalitet", "Ofullständiga läsningar", "Status"], [[a.get("gene"),a.get("mutation"),a.get("reference_codon"),a.get("alternate_codon"),a.get("depth"),a.get("alt_reads"),a.get("frequency"),a.get("local_indel_support"),a.get("restored_frame_support"),a.get("unresolved_frame_support"),a.get("low_baseq_reads"),a.get("incomplete_reads"),_status(a.get("status"))] for a in row.get("cmv_read_codon_diagnostics", [])], "Inget ytterligare diagnostiskt läsunderlag.", 1220),
            "__SNVTABLE__": _html_table(["Referensposition", "Vanligaste bas", "Alternativ bas", "Läsdjup", "Variantandel %"], minor_rows, "Inga mindre SNV-fynd att visa.", 650),
        }
        for key, value in substitutions.items():
            html = html.replace(key, value)
        filename = safe_report_filename(str(sample or "sample"))
        report_path = os.path.join(directory, f"{filename}_{run_tag}_CMV.html")
        with open(report_path, "w", encoding="utf-8") as handle:
            handle.write(html)
        row["cmv_html_report"] = report_path

    _style_workbook(workbook)
    results_ws.sheet_view.showGridLines = False
    results_ws.freeze_panes = "C2"
    results_ws.sheet_properties.pageSetUpPr.fitToPage = True
    results_ws.sheet_view.tabSelected = True
    workbook.active = 0
    results_ws.column_dimensions["A"].width = 25
    results_ws.column_dimensions["B"].width = 12
    results_ws.column_dimensions["C"].width = 17
    results_ws.column_dimensions["D"].width = 19
    results_ws.column_dimensions["E"].width = 23
    results_ws.column_dimensions["F"].width = 49
    results_ws.column_dimensions["G"].width = 19
    results_ws.column_dimensions["H"].width = 28
    results_ws.column_dimensions["I"].width = 35
    results_ws.column_dimensions["J"].width = 23
    for cells in results_ws.iter_rows(min_row=2):
        cells[3].number_format = '0.00"%"'
        if cells[2].value == "–":
            for cell in cells:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
    workbook.save(os.path.join(directory, f"CMV_variants_{run_tag}.xlsx"))
