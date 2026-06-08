import os
import sys
import subprocess
import re
import gzip
import statistics
from datetime import datetime
from gooey import Gooey, GooeyParser
from openpyxl import Workbook
from openpyxl.styles import PatternFill


def app_base_dir():
    return os.path.dirname(os.path.abspath(__file__))

def ensure_dir(dirpath, must_be_writable=False):
    os.makedirs(dirpath, exist_ok=True)
    if must_be_writable:
        try:
            testfile = os.path.join(dirpath, ".write_test.tmp")
            with open(testfile, "w", encoding="utf-8") as f:
                f.write("ok")
            os.remove(testfile)
        except Exception as e:
            raise PermissionError(f"Cannot write to '{dirpath}': {e}")

def which_or_none(exe_name):
    try:
        import shutil
        return shutil.which(exe_name)
    except Exception:
        return None

def check_required_tools():
    missing = []
    for tool in ["minimap2", "samtools", "medaka_consensus"]:
        if which_or_none(tool) is None:
            missing.append(tool)

    if missing:
        raise FileNotFoundError(
            "Missing required tools on PATH: " + ", ".join(missing) +
            "\nPlease install them and make sure they are available in your shell PATH."
        )

def path_exists_any(*paths):
    for p in paths:
        if p and os.path.exists(p):
            return p
    return None

def safe_remove(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass

# FILEfiltering

def is_real_fastq_gz(path):
    name = os.path.basename(path)
    if name.startswith("._") or name.startswith("."):
        return False
    return name.endswith(".fastq.gz")

def is_real_fasta(path):
    name = os.path.basename(path)
    if name.startswith("._") or name.startswith("."):
        return False
    return name.endswith(".fa") or name.endswith(".fasta")

def find_files_recursively(directory, predicate):
    matched = []
    try:
        for root, _, files in os.walk(directory):
            for file in files:
                full = os.path.join(root, file)
                try:
                    if predicate(full):
                        matched.append(full)
                except Exception:
                    continue
    except Exception as e:
        print(f"Error finding files in {directory}: {e}")
    return matched

#FASTQCount

def count_fastq_bases(fastq_gz_path):
    total_bases = 0
    try:
        with gzip.open(fastq_gz_path, "rt", encoding="utf-8", errors="replace") as f:
            for i, line in enumerate(f):
                if i % 4 == 1:
                    total_bases += len(line.strip())
    except Exception as e:
        print(f"Error counting FASTQ bases in {fastq_gz_path}: {e}")
    return total_bases

# SAM/BAMParse

LN_RE = re.compile(r"\bLN:(\d+)")
SN_RE = re.compile(r"\bSN:([^\t]+)")
CIGAR_RE = re.compile(r"(\d+)([MIDNSHP=X])")

def reference_consumed_by_cigar(cigar):
    ref_consumed = 0
    for length_str, op in CIGAR_RE.findall(cigar):
        length = int(length_str)
        if op in ["M", "=", "X", "D", "N"]:
            ref_consumed += length
    return ref_consumed

def query_aligned_consumed_by_cigar(cigar):
    query_consumed = 0
    for length_str, op in CIGAR_RE.findall(cigar):
        length = int(length_str)
        if op in ["M", "=", "X", "I"]:
            query_consumed += length
    return query_consumed

def evaluate_alignment_quality(output_sam):
    ref_lengths = {}
    ref_aligned_bases = {}
    ref_mapped_reads = {}
    ref_aligned_query_bases = {}

    try:
        with open(output_sam, "r", encoding="utf-8", errors="replace") as sam_file:
            for line in sam_file:
                if line.startswith("@SQ"):
                    ln_m = LN_RE.search(line)
                    sn_m = SN_RE.search(line)
                    if ln_m and sn_m:
                        ref_lengths[sn_m.group(1)] = int(ln_m.group(1))

                elif line and not line.startswith("@"):
                    fields = line.rstrip("\n").split("\t")
                    if len(fields) < 6:
                        continue

                    flag = int(fields[1])
                    ref_name = fields[2]
                    cigar = fields[5]

                    if (flag & 0x4) or ref_name == "*" or cigar == "*":
                        continue

                    ref_mapped_reads[ref_name] = ref_mapped_reads.get(ref_name, 0) + 1

                    ref_consumed = reference_consumed_by_cigar(cigar)
                    query_consumed_aligned = query_aligned_consumed_by_cigar(cigar)

                    ref_aligned_bases[ref_name] = ref_aligned_bases.get(ref_name, 0) + ref_consumed
                    ref_aligned_query_bases[ref_name] = ref_aligned_query_bases.get(ref_name, 0) + query_consumed_aligned

        best_ref = None
        best_score = -1.0

        for ref_name, aligned_bases in ref_aligned_bases.items():
            ref_len = ref_lengths.get(ref_name, 0)
            score = aligned_bases / ref_len if ref_len > 0 else 0.0

            if score > best_score:
                best_score = score
                best_ref = ref_name

        if not best_ref:
            return 0.0, 0, None, 0, 0

        return (
            best_score,
            ref_mapped_reads.get(best_ref, 0),
            best_ref,
            ref_aligned_query_bases.get(best_ref, 0),
            ref_lengths.get(best_ref, 0)
        )

    except Exception as e:
        print(f"Error evaluating alignment quality for {output_sam}: {e}")
        return 0.0, 0, None, 0, 0

def calculate_reference_breadth_coverage_percent(bam_file, ref_name, ref_length):
    if not bam_file or not os.path.exists(bam_file):
        return 0.0

    if not ref_name or not ref_length:
        return 0.0

    covered_positions = 0

    try:
        cmd = ["samtools", "depth", "-a", "-r", ref_name, bam_file]

        with subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True
        ) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue

                try:
                    depth = int(fields[2])
                    if depth > 0:
                        covered_positions += 1
                except ValueError:
                    continue

            _, stderr = proc.communicate()

            if proc.returncode != 0:
                print(f"samtools depth warning/error for {bam_file}: {stderr}")
                return 0.0

        return round((covered_positions / ref_length) * 100, 2) if ref_length > 0 else 0.0

    except Exception as e:
        print(f"Error calculating reference breadth coverage for {bam_file}: {e}")
        return 0.0

def empty_depth_stats():
    return {
        "mean_depth": 0.0,
        "median_depth": 0.0,
        "min_depth": 0,
        "positions_ge_10x_percent": 0.0,
        "positions_ge_30x_percent": 0.0,
    }

def calculate_depth_stats(bam_file, ref_name, region=None):
    depths = []

    if not bam_file or not os.path.exists(bam_file) or not ref_name:
        return empty_depth_stats()

    try:
        if region:
            region_string = f"{ref_name}:{region[0]}-{region[1]}"
            cmd = ["samtools", "depth", "-a", "-r", region_string, bam_file]
        else:
            cmd = ["samtools", "depth", "-a", "-r", ref_name, bam_file]

        with subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True
        ) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue

                try:
                    depths.append(int(fields[2]))
                except ValueError:
                    continue

            _, stderr = proc.communicate()

            if proc.returncode != 0:
                print(f"samtools depth warning/error for {bam_file}: {stderr}")

        if not depths:
            return empty_depth_stats()

        total = len(depths)

        return {
            "mean_depth": round(sum(depths) / total, 2),
            "median_depth": round(statistics.median(depths), 2),
            "min_depth": min(depths),
            "positions_ge_10x_percent": round((sum(1 for d in depths if d >= 10) / total) * 100, 2),
            "positions_ge_30x_percent": round((sum(1 for d in depths if d >= 30) / total) * 100, 2),
        }

    except Exception as e:
        print(f"Error calculating depth stats for {bam_file}: {e}")
        return empty_depth_stats()

def merge_depth_values(depth_values):
    if not depth_values:
        return empty_depth_stats()

    total = len(depth_values)

    return {
        "mean_depth": round(sum(depth_values) / total, 2),
        "median_depth": round(statistics.median(depth_values), 2),
        "min_depth": min(depth_values),
        "positions_ge_10x_percent": round((sum(1 for d in depth_values if d >= 10) / total) * 100, 2),
        "positions_ge_30x_percent": round((sum(1 for d in depth_values if d >= 30) / total) * 100, 2),
    }

def get_depth_values_for_region(bam_file, ref_name, region):
    depths = []

    if not bam_file or not os.path.exists(bam_file) or not ref_name or not region:
        return depths

    try:
        region_string = f"{ref_name}:{region[0]}-{region[1]}"
        cmd = ["samtools", "depth", "-a", "-r", region_string, bam_file]

        with subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True
        ) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue
                try:
                    depths.append(int(fields[2]))
                except ValueError:
                    continue

            _, stderr = proc.communicate()
            if proc.returncode != 0:
                print(f"samtools depth warning/error for {bam_file}: {stderr}")

    except Exception as e:
        print(f"Error getting depth values for region {region} in {bam_file}: {e}")

    return depths

def calculate_multi_region_depth_stats(bam_file, ref_name, regions):
    all_depths = []
    for region in regions:
        all_depths.extend(get_depth_values_for_region(bam_file, ref_name, region))
    return merge_depth_values(all_depths)

def detect_aligned_target_region_from_sam(sam_file, ref_name):
    starts = []
    ends = []

    try:
        with open(sam_file, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                if line.startswith("@"):
                    continue

                fields = line.rstrip("\n").split("\t")
                if len(fields) < 6:
                    continue

                try:
                    flag = int(fields[1])
                    rname = fields[2]
                    pos = int(fields[3])
                    cigar = fields[5]
                except ValueError:
                    continue

                if (flag & 0x4) or rname != ref_name or cigar == "*":
                    continue

                ref_consumed = reference_consumed_by_cigar(cigar)
                if ref_consumed <= 0:
                    continue

                starts.append(pos)
                ends.append(pos + ref_consumed - 1)

        if not starts or not ends:
            return None, None, 0

        start = min(starts)
        end = max(ends)
        length = end - start + 1

        return start, end, length

    except Exception as e:
        print(f"Error detecting aligned target region from {sam_file}: {e}")
        return None, None, 0

def coverage_fill(cell, coverage_percent, red_fill, yellow_fill, green_fill):
    if coverage_percent >= 90:
        cell.fill = green_fill
    elif coverage_percent >= 40:
        cell.fill = yellow_fill
    else:
        cell.fill = red_fill

def qc_fill(cell, qc_verdict, red_fill, yellow_fill, green_fill):
    verdict = str(qc_verdict).upper()
    if verdict == "HIGH":
        cell.fill = green_fill
    elif verdict == "MEDIUM":
        cell.fill = yellow_fill
    elif verdict == "LOW":
        cell.fill = red_fill

# FASTA

def read_single_fasta_sequence(fasta_path):
    seq_parts = []
    try:
        with open(fasta_path, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                if not line.startswith(">"):
                    seq_parts.append(line.strip())
        return "".join(seq_parts).upper()
    except Exception as e:
        print(f"Error reading FASTA sequence from {fasta_path}: {e}")
        return ""

def fasta_sequence_stats(fasta_path):
    seq = read_single_fasta_sequence(fasta_path)
    if not seq:
        return {
            "consensus_length": 0,
            "n_count": 0,
            "n_percent": 100.0,
            "non_n_length": 0,
        }

    seq = re.sub(r"[^ACGTN]", "", seq.upper())
    total = len(seq)
    n_count = seq.count("N")
    non_n = total - n_count

    return {
        "consensus_length": total,
        "n_count": n_count,
        "n_percent": round((n_count / total) * 100, 2) if total else 100.0,
        "non_n_length": non_n,
    }

CODON_TABLE = {
    "TTT":"F","TTC":"F","TTA":"L","TTG":"L",
    "TCT":"S","TCC":"S","TCA":"S","TCG":"S",
    "TAT":"Y","TAC":"Y","TAA":"*","TAG":"*",
    "TGT":"C","TGC":"C","TGA":"*","TGG":"W",

    "CTT":"L","CTC":"L","CTA":"L","CTG":"L",
    "CCT":"P","CCC":"P","CCA":"P","CCG":"P",
    "CAT":"H","CAC":"H","CAA":"Q","CAG":"Q",
    "CGT":"R","CGC":"R","CGA":"R","CGG":"R",

    "ATT":"I","ATC":"I","ATA":"I","ATG":"M",
    "ACT":"T","ACC":"T","ACA":"T","ACG":"T",
    "AAT":"N","AAC":"N","AAA":"K","AAG":"K",
    "AGT":"S","AGC":"S","AGA":"R","AGG":"R",

    "GTT":"V","GTC":"V","GTA":"V","GTG":"V",
    "GCT":"A","GCC":"A","GCA":"A","GCG":"A",
    "GAT":"D","GAC":"D","GAA":"E","GAG":"E",
    "GGT":"G","GGC":"G","GGA":"G","GGG":"G",
}

def translate_codon(codon):
    codon = codon.upper().replace("U", "T")
    if len(codon) != 3:
        return "X"
    if any(base not in "ACGT" for base in codon):
        return "X"
    return CODON_TABLE.get(codon, "X")

def circular_subseq(seq, start_1based, length):
    if not seq:
        return ""
    n = len(seq)
    start_idx = (start_1based - 1) % n
    out = []
    for i in range(length):
        out.append(seq[(start_idx + i) % n])
    return "".join(out)

def circular_codon(seq, start_1based):
    return circular_subseq(seq, start_1based, 3)

#QC

HBV_EXPECTED_GENOME_LENGTH = 3200

# Original primer windows:
# System 1 outer: 56F to 1827R
# System 1 nested: 251F to 1801R
# System 2: 1689F to 262R, circular wraparound

HBV_SYSTEM1_FULL_REGIONS = [(251, 1801)]
HBV_SYSTEM2_FULL_REGIONS = [(1689, HBV_EXPECTED_GENOME_LENGTH), (1, 262)]

HBV_SYSTEM1_UNIQUE_REGIONS = [(263, 1688)]
HBV_SYSTEM2_UNIQUE_REGIONS = [(1802, HBV_EXPECTED_GENOME_LENGTH), (1, 250)]
HBV_SYSTEM_OVERLAP_REGIONS = [(251, 262), (1689, 1801)]

def expected_target_length_for_qc(virus_type, detected_target_length):
    if virus_type == "Hepatitis B Virus":
        return HBV_EXPECTED_GENOME_LENGTH

    if virus_type == "Adenovirus":
        return detected_target_length or 0

    return detected_target_length or 0

def calculate_hbv_system_depth_stats(bam_file, ref_name):
    return {
        "system1_full": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM1_FULL_REGIONS),
        "system2_full": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM2_FULL_REGIONS),
        "system1_unique": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM1_UNIQUE_REGIONS),
        "system2_unique": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM2_UNIQUE_REGIONS),
        "overlap_only": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM_OVERLAP_REGIONS),
    }

def empty_hbv_system_stats():
    return {
        "system1_full": empty_depth_stats(),
        "system2_full": empty_depth_stats(),
        "system1_unique": empty_depth_stats(),
        "system2_unique": empty_depth_stats(),
        "overlap_only": empty_depth_stats(),
    }

def hbv_system_qc_flags(hbv_system_stats):
    s1 = hbv_system_stats.get("system1_unique", empty_depth_stats())
    s2 = hbv_system_stats.get("system2_unique", empty_depth_stats())
    overlap = hbv_system_stats.get("overlap_only", empty_depth_stats())

    s1_median = s1.get("median_depth", 0)
    s2_median = s2.get("median_depth", 0)
    s1_10x = s1.get("positions_ge_10x_percent", 0)
    s2_10x = s2.get("positions_ge_10x_percent", 0)
    s1_30x = s1.get("positions_ge_30x_percent", 0)
    s2_30x = s2.get("positions_ge_30x_percent", 0)

    high_ok = (
        s1_median >= 30 and
        s2_median >= 30 and
        s1_10x >= 90 and
        s2_10x >= 90 and
        s1_30x >= 70 and
        s2_30x >= 70
    )

    medium_ok = (
        s1_median >= 10 and
        s2_median >= 10 and
        s1_10x >= 70 and
        s2_10x >= 70
    )

    warnings = []

    if s1_median < 10:
        warnings.append("System 1 unique region has low median depth <10x")
    elif s1_median < 30:
        warnings.append("System 1 unique region has moderate median depth <30x")

    if s2_median < 10:
        warnings.append("System 2 unique region has low median depth <10x")
    elif s2_median < 30:
        warnings.append("System 2 unique region has moderate median depth <30x")

    if s1_10x < 70:
        warnings.append(f"System 1 unique region has poor >=10x breadth ({s1_10x}%)")
    elif s1_10x < 90:
        warnings.append(f"System 1 unique region has partial >=10x breadth ({s1_10x}%)")

    if s2_10x < 70:
        warnings.append(f"System 2 unique region has poor >=10x breadth ({s2_10x}%)")
    elif s2_10x < 90:
        warnings.append(f"System 2 unique region has partial >=10x breadth ({s2_10x}%)")

    min_median = max(min(s1_median, s2_median), 1)
    max_median = max(s1_median, s2_median)
    balance_ratio = round(max_median / min_median, 2) if max_median > 0 else 0.0

    if s1_median == 0 or s2_median == 0:
        warnings.append("One HBV system unique region has zero median depth")
    elif balance_ratio >= 10:
        warnings.append(f"Strong imbalance between HBV systems based on unique regions ({balance_ratio}x)")
    elif balance_ratio >= 5:
        warnings.append(f"Moderate imbalance between HBV systems based on unique regions ({balance_ratio}x)")

    overlap_median = overlap.get("median_depth", 0)
    if overlap_median >= 30 and (s1_median < 30 or s2_median < 30):
        warnings.append("Overlap has good depth but at least one unique system region is weaker; QC uses unique regions")

    return {
        "high_ok": high_ok,
        "medium_ok": medium_ok,
        "warning": "; ".join(warnings) if warnings else "Both HBV systems acceptable by unique-region QC",
        "balance_ratio": balance_ratio,
    }

def determine_qc_verdict(
    virus_type,
    mapped_reads,
    coverage_percent,
    median_depth,
    n_percent,
    consensus_length,
    expected_length,
    hbv_system_high_ok=True,
    hbv_system_medium_ok=True,
    hbv_system_warning=""
):
    reasons = []

    if expected_length and expected_length > 0:
        recovery_percent = round((consensus_length / expected_length) * 100, 2)
    else:
        recovery_percent = 0.0

    if virus_type == "Hepatitis B Virus":
        coverage_name = "HBV genome recovery"
    elif virus_type == "Adenovirus":
        coverage_name = "Adenovirus target/amplicon recovery"
    else:
        coverage_name = "target recovery"

    high = (
        mapped_reads >= 100 and
        coverage_percent >= 90 and
        median_depth >= 30 and
        n_percent <= 2 and
        recovery_percent >= 90 and
        hbv_system_high_ok
    )

    medium = (
        mapped_reads >= 50 and
        coverage_percent >= 70 and
        median_depth >= 10 and
        n_percent <= 10 and
        recovery_percent >= 70 and
        hbv_system_medium_ok
    )

    if high:
        verdict = "HIGH"
    elif medium:
        verdict = "MEDIUM"
    else:
        verdict = "LOW"

    if mapped_reads < 50:
        reasons.append("low mapped read count")
    elif mapped_reads < 100:
        reasons.append("moderate mapped read count")
    else:
        reasons.append("good mapped read count")

    if coverage_percent < 70:
        reasons.append(f"low {coverage_name}")
    elif coverage_percent < 90:
        reasons.append(f"partial {coverage_name}")
    else:
        reasons.append(f"strong {coverage_name}")

    if median_depth < 10:
        reasons.append("low median depth")
    elif median_depth < 30:
        reasons.append("moderate median depth")
    else:
        reasons.append("good median depth")

    if n_percent > 10:
        reasons.append("high N content")
    elif n_percent > 2:
        reasons.append("moderate N content")
    else:
        reasons.append("low N content")

    if expected_length and expected_length > 0:
        if recovery_percent < 70:
            reasons.append(f"short consensus vs expected length ({recovery_percent}%)")
        elif recovery_percent < 90:
            reasons.append(f"partial consensus length recovery ({recovery_percent}%)")
        else:
            reasons.append(f"good consensus length recovery ({recovery_percent}%)")

    if virus_type == "Hepatitis B Virus":
        if hbv_system_high_ok:
            reasons.append("HBV System 1 and System 2 unique-region depth supports HIGH QC")
        elif hbv_system_medium_ok:
            reasons.append("HBV System 1 and System 2 unique-region depth supports MEDIUM QC only")
        else:
            reasons.append("HBV System 1/System 2 unique-region depth fails medium threshold")

        if hbv_system_warning:
            reasons.append(hbv_system_warning)

    return verdict, "; ".join(reasons), recovery_percent

# Medaka

def generate_consensus(input_fastq, draft_fasta, medaka_model, output_folder, sample_name):
    try:
        medaka_output_folder = os.path.join(output_folder, f"{sample_name}_medaka")
        ensure_dir(medaka_output_folder, must_be_writable=True)

        cmd = [
            "medaka_consensus",
            "-i", input_fastq,
            "-d", draft_fasta,
            "-o", medaka_output_folder,
            "-m", medaka_model,
        ]

        env=os.environ.copy()
        env["PYTHONNOUSERSITE"] = "1"

        subprocess.run(cmd, check=True, env=env)

        medaka_out = os.path.join(medaka_output_folder, "consensus.fasta")
        if not (os.path.exists(medaka_out) and os.path.getsize(medaka_out) > 0):
            print(f"Medaka did not produce the expected output file: {medaka_out}")
            return None

        renamed = os.path.join(medaka_output_folder, f"{sample_name}.consensus.fasta")

        with open(medaka_out, "r", encoding="utf-8", errors="replace") as f:
            lines = f.read().splitlines()

        if not lines:
            print(f"Warning: consensus for {sample_name} is empty.")
            return None

        body = [line for line in lines if not line.startswith(">")]

        with open(renamed, "w", encoding="utf-8") as f:
            f.write(f">{sample_name}\n")
            f.write("\n".join(body) + ("\n" if body else ""))

        print(f"Medaka completed. Output: {renamed}")
        return renamed if os.path.getsize(renamed) > 0 else None

    except subprocess.CalledProcessError as e:
        print(f"Error in Medaka process for {input_fastq}: {e}")
        return None
    except Exception as e:
        print(f"Unexpected error in Medaka for {input_fastq}: {e}")
        return None

#Logging

def autosize_worksheet(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 60)

def log_reference_usage(result_rows, output_folder, virus_type):
    try:
        ensure_dir(output_folder, must_be_writable=True)
        excel_path = os.path.join(output_folder, "reference_usage.xlsx")

        wb = Workbook()

        ws = wb.active
        ws.title = "Summary"

        details_ws = wb.create_sheet("Details")

        summary_headers = [
            "Sample",
            "Virus Type",
            "Best Reference (SN)",
            "Mapped Reads",
            "Target/Genome Coverage (%)",
            "Mean Depth",
            "Median Depth",
            "Minimum Depth",
            "% Positions >=10x",
            "% Positions >=30x",
            "Consensus Length",
            "N Count",
            "N %",
            "Consensus Recovery vs Expected (%)",
            "QC Verdict",
            "QC Reason",
            "BAM File"
        ]

        if virus_type == "Hepatitis B Virus":
            summary_headers.extend([
                "HBV System 1 Full Median Depth",
                "HBV System 1 Full % >=10x",
                "HBV System 1 Full % >=30x",
                "HBV System 2 Full Median Depth",
                "HBV System 2 Full % >=10x",
                "HBV System 2 Full % >=30x",
                "HBV System 1 Unique Median Depth",
                "HBV System 1 Unique % >=10x",
                "HBV System 1 Unique % >=30x",
                "HBV System 2 Unique Median Depth",
                "HBV System 2 Unique % >=10x",
                "HBV System 2 Unique % >=30x",
                "HBV Overlap Median Depth",
                "HBV Overlap % >=10x",
                "HBV Overlap % >=30x",
                "HBV System Balance Ratio",
                "HBV System QC Warning"
            ])

        detail_headers = [
            "Sample",
            "Reference Length",
            "Detected Target Start",
            "Detected Target End",
            "Detected Target Length",
            "Expected Target/Genome Length",
            "Coverage Type"
        ]

        ws.append(summary_headers)
        details_ws.append(detail_headers)

        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

        for row in result_rows:
            summary_row = [
                row.get("sample", ""),
                row.get("virus_type", ""),
                row.get("best_reference", ""),
                row.get("mapped_reads", 0),
                row.get("target_or_genome_coverage_percent", 0),
                row.get("mean_depth", 0),
                row.get("median_depth", 0),
                row.get("min_depth", 0),
                row.get("positions_ge_10x_percent", 0),
                row.get("positions_ge_30x_percent", 0),
                row.get("consensus_length", 0),
                row.get("n_count", 0),
                row.get("n_percent", 0),
                row.get("consensus_recovery_percent", 0),
                row.get("qc_verdict", ""),
                row.get("qc_reason", ""),
                row.get("bam_file", ""),
            ]

            if virus_type == "Hepatitis B Virus":
                summary_row.extend([
                    row.get("hbv_system1_full_median_depth", 0),
                    row.get("hbv_system1_full_positions_ge_10x_percent", 0),
                    row.get("hbv_system1_full_positions_ge_30x_percent", 0),
                    row.get("hbv_system2_full_median_depth", 0),
                    row.get("hbv_system2_full_positions_ge_10x_percent", 0),
                    row.get("hbv_system2_full_positions_ge_30x_percent", 0),
                    row.get("hbv_system1_unique_median_depth", 0),
                    row.get("hbv_system1_unique_positions_ge_10x_percent", 0),
                    row.get("hbv_system1_unique_positions_ge_30x_percent", 0),
                    row.get("hbv_system2_unique_median_depth", 0),
                    row.get("hbv_system2_unique_positions_ge_10x_percent", 0),
                    row.get("hbv_system2_unique_positions_ge_30x_percent", 0),
                    row.get("hbv_overlap_median_depth", 0),
                    row.get("hbv_overlap_positions_ge_10x_percent", 0),
                    row.get("hbv_overlap_positions_ge_30x_percent", 0),
                    row.get("hbv_system_balance_ratio", 0),
                    row.get("hbv_system_qc_warning", ""),
                ])

            ws.append(summary_row)

            details_ws.append([
                row.get("sample", ""),
                row.get("reference_length", 0),
                row.get("detected_target_start", ""),
                row.get("detected_target_end", ""),
                row.get("detected_target_length", 0),
                row.get("expected_target_length", 0),
                row.get("coverage_type", ""),
            ])

            mapped_cell = ws.cell(row=ws.max_row, column=4)
            mapped_reads = row.get("mapped_reads", 0)
            mapped_cell.fill = red_fill if mapped_reads < 100 else green_fill

            coverage_cell = ws.cell(row=ws.max_row, column=5)
            coverage = row.get("target_or_genome_coverage_percent", 0)
            coverage_fill(coverage_cell, coverage, red_fill, yellow_fill, green_fill)

            qc_cell = ws.cell(row=ws.max_row, column=15)
            qc_fill(qc_cell, row.get("qc_verdict", ""), red_fill, yellow_fill, green_fill)

            if virus_type == "Hepatitis B Virus":
                # Highlight the unique-region system metrics because those drive QC.
                for col in [24, 25, 26, 27, 28, 29]:
                    cell = ws.cell(row=ws.max_row, column=col)
                    try:
                        value = float(cell.value)
                    except Exception:
                        value = 0.0

                    if "Median Depth" in str(ws.cell(row=1, column=col).value):
                        if value >= 30:
                            cell.fill = green_fill
                        elif value >= 10:
                            cell.fill = yellow_fill
                        else:
                            cell.fill = red_fill
                    else:
                        coverage_fill(cell, value, red_fill, yellow_fill, green_fill)

        autosize_worksheet(ws)
        autosize_worksheet(details_ws)

        wb.save(excel_path)
        print(f"Reference usage Excel saved to {excel_path}")

    except Exception as e:
        print(f"Error logging reference usage: {e}")

# Per-sample processing (loop)

def add_hbv_system_stats_to_row(row, hbv_system_stats, hbv_flags):
    s1_full = hbv_system_stats.get("system1_full", empty_depth_stats())
    s2_full = hbv_system_stats.get("system2_full", empty_depth_stats())
    s1_unique = hbv_system_stats.get("system1_unique", empty_depth_stats())
    s2_unique = hbv_system_stats.get("system2_unique", empty_depth_stats())
    overlap = hbv_system_stats.get("overlap_only", empty_depth_stats())

    row["hbv_system1_full_median_depth"] = s1_full.get("median_depth", 0)
    row["hbv_system1_full_positions_ge_10x_percent"] = s1_full.get("positions_ge_10x_percent", 0)
    row["hbv_system1_full_positions_ge_30x_percent"] = s1_full.get("positions_ge_30x_percent", 0)

    row["hbv_system2_full_median_depth"] = s2_full.get("median_depth", 0)
    row["hbv_system2_full_positions_ge_10x_percent"] = s2_full.get("positions_ge_10x_percent", 0)
    row["hbv_system2_full_positions_ge_30x_percent"] = s2_full.get("positions_ge_30x_percent", 0)

    row["hbv_system1_unique_median_depth"] = s1_unique.get("median_depth", 0)
    row["hbv_system1_unique_positions_ge_10x_percent"] = s1_unique.get("positions_ge_10x_percent", 0)
    row["hbv_system1_unique_positions_ge_30x_percent"] = s1_unique.get("positions_ge_30x_percent", 0)

    row["hbv_system2_unique_median_depth"] = s2_unique.get("median_depth", 0)
    row["hbv_system2_unique_positions_ge_10x_percent"] = s2_unique.get("positions_ge_10x_percent", 0)
    row["hbv_system2_unique_positions_ge_30x_percent"] = s2_unique.get("positions_ge_30x_percent", 0)

    row["hbv_overlap_median_depth"] = overlap.get("median_depth", 0)
    row["hbv_overlap_positions_ge_10x_percent"] = overlap.get("positions_ge_10x_percent", 0)
    row["hbv_overlap_positions_ge_30x_percent"] = overlap.get("positions_ge_30x_percent", 0)

    row["hbv_system_balance_ratio"] = hbv_flags.get("balance_ratio", 0)
    row["hbv_system_qc_warning"] = hbv_flags.get("warning", "")

    return row

def process_sample(
    input_file,
    reference_files,
    output_folder,
    medaka_model,
    virus_type,
    final_consensus_list,
    result_rows
):
    try:
        base_name = os.path.splitext(os.path.basename(input_file))[0].replace(".fastq", "").replace(".gz", "")
        sample_dir = os.path.join(output_folder, base_name)
        ensure_dir(sample_dir, must_be_writable=True)

        best = {
            "ref_path": None,
            "ref_sn": None,
            "approx_coverage": -1.0,
            "sam_path": None,
            "mapped": 0,
            "aligned_query_bases": 0,
            "ref_length": 0
        }

        for reference_file in reference_files:
            ref_name = os.path.splitext(os.path.basename(reference_file))[0]
            sam_tmp = os.path.join(sample_dir, f"{ref_name}.sam")

            cmd = ["minimap2", "-ax", "map-ont", reference_file, input_file]
            with open(sam_tmp, "w", encoding="utf-8") as sam_out:
                subprocess.run(cmd, check=True, stdout=sam_out)

            approx_cov, mapped, sn, aligned_query_bases, ref_length = evaluate_alignment_quality(sam_tmp)

            if approx_cov > best["approx_coverage"]:
                safe_remove(best["sam_path"])
                best.update({
                    "ref_path": reference_file,
                    "ref_sn": sn,
                    "approx_coverage": approx_cov,
                    "sam_path": sam_tmp,
                    "mapped": mapped,
                    "aligned_query_bases": aligned_query_bases,
                    "ref_length": ref_length
                })
            else:
                safe_remove(sam_tmp)

        if not (best["ref_path"] and best["sam_path"]):
            print(f"No suitable reference found for {input_file}. Skipping.")
            return

        detected_target_start, detected_target_end, detected_target_length = detect_aligned_target_region_from_sam(
            best["sam_path"],
            best["ref_sn"]
        )

        tmp_ref = os.path.join(sample_dir, f"{base_name}_best_reference.fasta")
        found = False

        with open(best["ref_path"], "r", encoding="utf-8", errors="replace") as ref_in, \
             open(tmp_ref, "w", encoding="utf-8") as ref_out:
            write = False
            for line in ref_in:
                if line.startswith(">"):
                    header = line[1:].strip()
                    name_only = header.split()[0]
                    write = (name_only == best["ref_sn"])
                if write:
                    ref_out.write(line)
                    found = True

        if not found:
            print(f"Could not find sequence '{best['ref_sn']}' in {best['ref_path']}. Using full reference instead.")
            tmp_ref = best["ref_path"]

        bam_sorted = os.path.join(sample_dir, f"{base_name}.sorted.bam")
        subprocess.run(["samtools", "sort", "-O", "bam", "-o", bam_sorted, best["sam_path"]], check=True)
        subprocess.run(["samtools", "index", bam_sorted], check=True)

        reference_coverage_percent = calculate_reference_breadth_coverage_percent(
            bam_file=bam_sorted,
            ref_name=best["ref_sn"],
            ref_length=best["ref_length"]
        )

        hbv_system_stats = empty_hbv_system_stats()
        hbv_flags = {
            "high_ok": True,
            "medium_ok": True,
            "warning": "",
            "balance_ratio": 0,
        }

        if virus_type == "Adenovirus" and detected_target_start and detected_target_end:
            depth_region = (detected_target_start, detected_target_end)
            coverage_type = "Detected Adenovirus Target/Amplicon Coverage"
            expected_target_length = expected_target_length_for_qc(virus_type, detected_target_length)

            depth_stats = calculate_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                region=depth_region
            )

            target_or_genome_coverage_percent = 100.0 if detected_target_length > 0 else 0.0

        elif virus_type == "Hepatitis B Virus":
            depth_region = None
            coverage_type = "Approximate HBV Whole-Genome Coverage"
            expected_target_length = HBV_EXPECTED_GENOME_LENGTH

            depth_stats = calculate_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                region=None
            )

            target_or_genome_coverage_percent = reference_coverage_percent

            hbv_system_stats = calculate_hbv_system_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"]
            )
            hbv_flags = hbv_system_qc_flags(hbv_system_stats)

        else:
            depth_region = None
            coverage_type = "Reference Coverage"
            expected_target_length = expected_target_length_for_qc(virus_type, detected_target_length)

            depth_stats = calculate_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                region=None
            )

            target_or_genome_coverage_percent = reference_coverage_percent

        print(
            f"Best reference for {input_file}: {best['ref_sn']} "
            f"with {best['mapped']} mapped reads."
        )

        if virus_type == "Adenovirus":
            print(
                f"Detected adenovirus target region: "
                f"{detected_target_start}-{detected_target_end} "
                f"({detected_target_length} bp)."
            )

        print(
            f"{coverage_type}: {target_or_genome_coverage_percent}% | "
            f"Median depth: {depth_stats.get('median_depth', 0)}x"
        )

        if virus_type == "Hepatitis B Virus":
            print(
                "HBV unique-region QC: "
                f"System 1 median {hbv_system_stats['system1_unique'].get('median_depth', 0)}x, "
                f"System 2 median {hbv_system_stats['system2_unique'].get('median_depth', 0)}x | "
                f"{hbv_flags.get('warning', '')}"
            )

        draft_consensus = os.path.join(sample_dir, f"{base_name}.draft.fasta")
        subprocess.run(["samtools", "consensus", "-f", "fasta", "-o", draft_consensus, bam_sorted], check=True)

        if not (os.path.exists(draft_consensus) and os.path.getsize(draft_consensus) > 0):
            print(f"Draft consensus missing/empty for {base_name}; skipping Medaka.")

            fasta_stats = {
                "consensus_length": 0,
                "n_count": 0,
                "n_percent": 100.0,
                "non_n_length": 0,
            }

            qc_verdict, qc_reason, consensus_recovery_percent = determine_qc_verdict(
                virus_type=virus_type,
                mapped_reads=best["mapped"],
                coverage_percent=target_or_genome_coverage_percent,
                median_depth=depth_stats.get("median_depth", 0),
                n_percent=fasta_stats["n_percent"],
                consensus_length=fasta_stats["consensus_length"],
                expected_length=expected_target_length,
                hbv_system_high_ok=hbv_flags.get("high_ok", True),
                hbv_system_medium_ok=hbv_flags.get("medium_ok", True),
                hbv_system_warning=hbv_flags.get("warning", "")
            )

            row = {
                "sample": base_name,
                "virus_type": virus_type,
                "best_reference": best["ref_sn"] or "NA",
                "mapped_reads": best["mapped"],
                "reference_length": best["ref_length"],
                "detected_target_start": detected_target_start or "",
                "detected_target_end": detected_target_end or "",
                "detected_target_length": detected_target_length,
                "expected_target_length": expected_target_length,
                "coverage_type": coverage_type,
                "target_or_genome_coverage_percent": target_or_genome_coverage_percent,
                "mean_depth": depth_stats.get("mean_depth", 0),
                "median_depth": depth_stats.get("median_depth", 0),
                "min_depth": depth_stats.get("min_depth", 0),
                "positions_ge_10x_percent": depth_stats.get("positions_ge_10x_percent", 0),
                "positions_ge_30x_percent": depth_stats.get("positions_ge_30x_percent", 0),
                "consensus_length": fasta_stats["consensus_length"],
                "n_count": fasta_stats["n_count"],
                "n_percent": fasta_stats["n_percent"],
                "consensus_recovery_percent": consensus_recovery_percent,
                "qc_verdict": qc_verdict,
                "qc_reason": qc_reason,
                "bam_file": bam_sorted,
            }

            if virus_type == "Hepatitis B Virus":
                row = add_hbv_system_stats_to_row(row, hbv_system_stats, hbv_flags)

            result_rows.append(row)

            safe_remove(best["sam_path"])
            if tmp_ref != best["ref_path"]:
                safe_remove(tmp_ref)
            return

        medaka_result = generate_consensus(input_file, draft_consensus, medaka_model, sample_dir, base_name)
        final_consensus_path_for_analysis = medaka_result if medaka_result else draft_consensus

        if medaka_result:
            final_consensus_list.append(medaka_result)

        fasta_stats = fasta_sequence_stats(final_consensus_path_for_analysis)

        qc_verdict, qc_reason, consensus_recovery_percent = determine_qc_verdict(
            virus_type=virus_type,
            mapped_reads=best["mapped"],
            coverage_percent=target_or_genome_coverage_percent,
            median_depth=depth_stats.get("median_depth", 0),
            n_percent=fasta_stats["n_percent"],
            consensus_length=fasta_stats["consensus_length"],
            expected_length=expected_target_length,
            hbv_system_high_ok=hbv_flags.get("high_ok", True),
            hbv_system_medium_ok=hbv_flags.get("medium_ok", True),
            hbv_system_warning=hbv_flags.get("warning", "")
        )

        row = {
            "sample": base_name,
            "virus_type": virus_type,
            "best_reference": best["ref_sn"] or "NA",
            "mapped_reads": best["mapped"],
            "reference_length": best["ref_length"],
            "detected_target_start": detected_target_start or "",
            "detected_target_end": detected_target_end or "",
            "detected_target_length": detected_target_length,
            "expected_target_length": expected_target_length,
            "coverage_type": coverage_type,
            "target_or_genome_coverage_percent": target_or_genome_coverage_percent,
            "mean_depth": depth_stats.get("mean_depth", 0),
            "median_depth": depth_stats.get("median_depth", 0),
            "min_depth": depth_stats.get("min_depth", 0),
            "positions_ge_10x_percent": depth_stats.get("positions_ge_10x_percent", 0),
            "positions_ge_30x_percent": depth_stats.get("positions_ge_30x_percent", 0),
            "consensus_length": fasta_stats["consensus_length"],
            "n_count": fasta_stats["n_count"],
            "n_percent": fasta_stats["n_percent"],
            "consensus_recovery_percent": consensus_recovery_percent,
            "qc_verdict": qc_verdict,
            "qc_reason": qc_reason,
            "bam_file": bam_sorted,
        }

        if virus_type == "Hepatitis B Virus":
            row = add_hbv_system_stats_to_row(row, hbv_system_stats, hbv_flags)

        result_rows.append(row)

        print(f"QC Verdict for {base_name}: {qc_verdict} ({qc_reason})")

        safe_remove(best["sam_path"])
        safe_remove(draft_consensus)

        if tmp_ref != best["ref_path"]:
            safe_remove(tmp_ref)

    except subprocess.CalledProcessError as e:
        print(f"Subprocess error while processing {input_file}: {e}")
    except Exception as e:
        print(f"Unexpected error in processing sample {input_file}: {e}")

def concatenate_fasta_files(fasta_files, output_file):
    if not fasta_files:
        print("No consensus FASTA files to concatenate.")
        return False

    try:
        ensure_dir(os.path.dirname(output_file), must_be_writable=True)
        with open(output_file, "w", encoding="utf-8") as outfile:
            for fasta in fasta_files:
                if os.path.isfile(fasta) and os.path.getsize(fasta) > 0:
                    with open(fasta, "r", encoding="utf-8", errors="replace") as infile:
                        outfile.write(infile.read())
                else:
                    print(f"Skipping empty or missing FASTA file: {fasta}")

        print(f"Concatenated FASTA file saved to {output_file}")
        return True

    except Exception as e:
        print(f"Error concatenating FASTA files: {e}")
        return False

def split_fasta(output_fasta, split_folder):
    ensure_dir(split_folder, must_be_writable=True)

    with open(output_fasta, "r", encoding="utf-8", errors="replace") as f:
        current = None
        for line in f:
            if line.startswith(">"):
                if current:
                    current.close()
                header = line.strip()[1:].replace(" ", "_")
                current = open(os.path.join(split_folder, f"{header}.fasta"), "w", encoding="utf-8")
            if current:
                current.write(line)

        if current:
            current.close()

    print(f"Split FASTA sequences saved to {split_folder}")

# NCBI pipeline (Not active, work in progress,decapritate)

def run_ncbi_database_script(input_folder):
    try:
        base = app_base_dir()

        ncbi_script_path = path_exists_any(
            os.path.join(base, "AdenoPipeline", "BLASTscript.py"),
            "/Volumes/Newt-Vault/AdenoPipeline/BLASTscript.py",
            "/Volumes/NewtVault/AdenoPipeline/BLASTscript.py"
        )

        if not ncbi_script_path:
            print("NCBI script not found. Put BLASTscript.py in ./AdenoPipeline/ next to this script or use the mounted volume path.")
            return

        output_excel = os.path.join(input_folder, "BLAST_results.xlsx")
        invalid_files_csv = os.path.join(input_folder, "invalid_files.csv")

        blastplus_path = path_exists_any(
            os.path.join(base, "blastplus"),
            "/Users/xsjoak/Documents/Entero/blastplus"
        )

        if not blastplus_path:
            print("BLAST+ folder not found. Put it in ./blastplus next to the script or adjust the path.")
            return

        subprocess.run([
            sys.executable,
            ncbi_script_path,
            input_folder,
            blastplus_path,
            output_excel,
            invalid_files_csv
        ], check=True)

        print(f"NCBI database script executed successfully. Results saved to {output_excel}")

    except subprocess.CalledProcessError as e:
        print(f"Error running the NCBI database script: {e}")
    except Exception as e:
        print(f"Unexpected error while running the NCBI script: {e}")

#Pipeline runners

def run_full_pipeline(input_folder, output_folder, reference_folder, medaka_model, virus_type, run_ncbi=False):
    print("\nRunning full pipeline...\n")
    print(f"Virus type: {virus_type}")
    print(f"Input: {input_folder}")
    print(f"Output: {output_folder}")
    print(f"Reference folder: {reference_folder}")
    print(f"Medaka model: {medaka_model}")
    print("Keep BAM files: True\n")

    try:
        check_required_tools()
    except Exception as e:
        print(f"[ERROR] {e}")
        return

    ensure_dir(output_folder, must_be_writable=True)

    input_files = find_files_recursively(input_folder, is_real_fastq_gz)
    reference_files = [
        os.path.join(reference_folder, f)
        for f in os.listdir(reference_folder)
        if is_real_fasta(os.path.join(reference_folder, f))
    ]

    if not input_files:
        print("No .fastq.gz input files found.")
        return

    if not reference_files:
        print("No reference .fa/.fasta files found.")
        return

    result_rows = []
    final_consensus_list = []

    for input_file in input_files:
        print(f"\n--- Processing sample: {input_file} ---")
        process_sample(
            input_file=input_file,
            reference_files=reference_files,
            output_folder=output_folder,
            medaka_model=medaka_model,
            virus_type=virus_type,
            final_consensus_list=final_consensus_list,
            result_rows=result_rows
        )

    if result_rows:
        log_reference_usage(result_rows, output_folder, virus_type)
    else:
        print("No reference usage data recorded.")

    today = datetime.now().strftime("%Y-%m-%d")
    safe_virus_name = virus_type.replace(" ", "_")
    final_fasta_path = os.path.join(output_folder, f"{safe_virus_name}_{today}_{len(final_consensus_list)}.fasta")

    if final_consensus_list:
        success = concatenate_fasta_files(final_consensus_list, final_fasta_path)
        if success and run_ncbi:
            split_folder = os.path.join(output_folder, "split_fasta")
            split_fasta(final_fasta_path, split_folder)
            run_ncbi_database_script(split_folder)
    else:
        print("No consensus files produced; skipping concatenation and optional NCBI step.")

# GUI Gooey

@Gooey(
    program_name="Aurora - Universal Virus Identification System",
    default_size=(820, 780),
    clear_before_run=True,
    label_text_color="#e8e8e8",
    header_bg_color="#78D5F1",
    body_bg_color="#AAAAAA",
    footer_bg_color="#e8e8e8",
    button_bg_color="#e9e9e96c",
    button_fg_color="#000000",
    terminal_font_color="#515151",
    terminal_font_size=10,
)
def main():
    parser = GooeyParser(description="Universal Virus Identification Pipeline")

    parser.add_argument(
        "virus_type",
        choices=["Adenovirus", "Hepatitis B Virus", "Other"],
        help="Select the virus being analysed"
    )

    parser.add_argument(
        "analysis_mode",
        choices=["Full Pipeline"],
        default="Full Pipeline",
        help="Select analysis mode"
    )

    parser.add_argument(
        "--input_folder",
        widget="DirChooser",
        help="Input directory containing .fastq.gz files (Full Pipeline only)"
    )

    parser.add_argument(
        "--output_folder",
        widget="DirChooser",
        required=True,
        help="Output directory"
    )

    parser.add_argument(
        "--reference_folder",
        widget="DirChooser",
        help="Folder containing .fa/.fasta references (Full Pipeline only)"
    )

    parser.add_argument(
        "--medaka_model",
        choices=[
            "r1041_e82_400bps_sup_v5.2.0",
            "r1041_e82_400bps_hac_v5.0.0",
            "r1041_e82_400bps_sup_v4.3.0",
            "r1041_e82_400bps_sup_v5.0.0",
            "r1041_e82_400bps_sup_v4.2.0",
            "r1041_e82_400bps_sup_v3.5.1",
            "r1041_e82_400bps_sup_variant_v4.2.0",
            "r1041_e82_400bps_hac_v4.2.0",
            "r1041_e82_400bps_hac_variant_v4.2.0",
            "r1041_e82_400bps_fast_g632",
            "r1041_e82_400bps_fast_variant_g632",
            "r1041_e82_260bps_sup_v4.1.0",
            "r1041_e82_260bps_sup_variant_v4.1.0",
            "r1041_e82_260bps_hac_v4.1.0",
            "r1041_e82_260bps_hac_variant_v4.1.0",
            "r1041_e82_260bps_fast_g632",
            "r1041_e82_260bps_fast_variant_g632",
            "r941_min_fast_g507",
            "r941_min_fast_variant_g507",
            "r941_min_hac_g507",
            "r941_min_hac_variant_g507",
            "r941_min_sup_g507",
            "r941_min_sup_variant_g507",
            "r941_prom_fast_g507",
            "r941_prom_fast_variant_g507",
            "r941_prom_hac_g507",
            "r941_prom_hac_variant_g507",
            "r941_prom_sup_g507",
            "r941_prom_sup_variant_g507",
        ],
        default="r1041_e82_400bps_sup_v4.2.0",
        metavar="Medaka Model",
        help="Select Medaka model"
    )

    parser.add_argument(
        "--run_ncbi",
        action="store_true",
        help="Run NCBI pipeline after consensus (Full Pipeline only)"
    )


    args = parser.parse_args()

    if args.analysis_mode == "Full Pipeline":
        if not args.input_folder:
            print("[ERROR] Please provide --input_folder for Full Pipeline mode.")
            return

        if not args.reference_folder:
            print("[ERROR] Please provide --reference_folder for Full Pipeline mode.")
            return

        run_full_pipeline(
            input_folder=args.input_folder,
            output_folder=args.output_folder,
            reference_folder=args.reference_folder,
            medaka_model=args.medaka_model,
            virus_type=args.virus_type,
            run_ncbi=args.run_ncbi
        )
        return

    print("[ERROR] Unknown analysis mode.")

if __name__ == "__main__":
    main()
