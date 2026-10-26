"""Excel workbook generation for run summaries and mutation reports."""
import os
from html import escape as html_escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .config import (
    HBV_MEDICATIONS,
    HBV_MIN_COVERAGE,
    HBV_MIN_VARIANT_FREQUENCY,
    HBV_RT_PANEL,
    HBV_S_ESCAPE_PANEL,
    HBV_S_ESCAPE_ASSOCIATIONS,
    HCV_MIN_COVERAGE,
    HCV_MIN_VARIANT_FREQUENCY,
    HCV_SCRIPT_VERSION,
)
from .filesystem import ensure_dir, run_samples_folder, safe_report_filename
from .hbv import (
    hbv_combination_interpretations,
    hbv_combination_summary_text,
    hbv_detected_calls,
    hbv_status_for_medication,
    hbv_swedish_evidence_label,
    hbv_swedish_interpretation,
    hbv_translate_free_text,
)
from .reporting_html import generate_hcv_html_reports, write_hbv_html_report


def coverage_fill(cell, coverage_percent, red_fill, yellow_fill, green_fill):
    if coverage_percent >= 90:
        cell.fill = green_fill
    elif coverage_percent >= 40:
        cell.fill = yellow_fill
    else:
        cell.fill = red_fill

def qc_fill(cell, qc_verdict, red_fill, yellow_fill, green_fill):
    verdict = str(qc_verdict).upper()
    if verdict == "HIGH":
        cell.fill = green_fill
    elif verdict == "MEDIUM":
        cell.fill = yellow_fill
    elif verdict == "LOW":
        cell.fill = red_fill

def autosize_worksheet(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 60)

def log_reference_usage(result_rows, output_folder, virus_type, run_tag):
    try:
        ensure_dir(output_folder, must_be_writable=True)
        excel_path = os.path.join(output_folder, f"reference_usage_{run_tag}.xlsx")

        wb = Workbook()

        ws = wb.active
        ws.title = "Summary"

        details_ws = wb.create_sheet("Details")

        summary_headers = [
            "Sample",
            "Virus Type",
            "Best Reference (SN)",
            "Mapped Reads",
            "Target/Genome Coverage (%)",
            "Mean Depth",
            "Median Depth",
            "Minimum Depth",
            "% Positions >=10x" if virus_type != "Adenovirus" else "Read depth >=10x (%)",
            "% Positions >=30x" if virus_type != "Adenovirus" else "Read depth >=30x (%)",
            "Consensus Length",
            "N Count",
            "N %",
            (
                "Callable Consensus Recovery vs Expected (%)"
                if virus_type == "Adenovirus"
                else "Consensus Recovery vs Expected (%)"
            ),
            "QC Verdict",
            "QC Reason",
            "BAM File"
        ]

        if virus_type == "Hepatitis B Virus":
            summary_headers.extend([
                "HBV System 1 Full Median Depth",
                "HBV System 1 Full % >=10x",
                "HBV System 1 Full % >=30x",
                "HBV System 2 Full Median Depth",
                "HBV System 2 Full % >=10x",
                "HBV System 2 Full % >=30x",
                "HBV System 1 Unique Median Depth",
                "HBV System 1 Unique % >=10x",
                "HBV System 1 Unique % >=30x",
                "HBV System 2 Unique Median Depth",
                "HBV System 2 Unique % >=10x",
                "HBV System 2 Unique % >=30x",
                "HBV Overlap Median Depth",
                "HBV Overlap % >=10x",
                "HBV Overlap % >=30x",
                "HBV System Balance Ratio",
                "HBV System QC Warning"
            ])

        detail_headers = [
            "Sample",
            "Reference Length",
            "Detected Target Start",
            "Detected Target End",
            "Detected Target Length",
            "Expected Target/Genome Length",
            "Coverage Type",
            "Reference Selection Mode",
            "Adenovirus Target Min Depth",
            "Consensus Source",
            "Consensus Fallback Reason",
        ]

        ws.append(summary_headers)
        details_ws.append(detail_headers)

        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

        for row in result_rows:
            summary_row = [
                row.get("sample", ""),
                row.get("virus_type", ""),
                row.get("best_reference", ""),
                row.get("mapped_reads", 0),
                row.get("target_or_genome_coverage_percent", 0),
                row.get("mean_depth", 0),
                row.get("median_depth", 0),
                row.get("min_depth", 0),
                row.get("positions_ge_10x_percent", 0),
                row.get("positions_ge_30x_percent", 0),
                row.get("consensus_length", 0),
                row.get("n_count", 0),
                row.get("n_percent", 0),
                row.get("consensus_recovery_percent", 0),
                row.get("qc_verdict", ""),
                row.get("qc_reason", ""),
                row.get("bam_file", ""),
            ]

            if virus_type == "Hepatitis B Virus":
                summary_row.extend([
                    row.get("hbv_system1_full_median_depth", 0),
                    row.get("hbv_system1_full_positions_ge_10x_percent", 0),
                    row.get("hbv_system1_full_positions_ge_30x_percent", 0),
                    row.get("hbv_system2_full_median_depth", 0),
                    row.get("hbv_system2_full_positions_ge_10x_percent", 0),
                    row.get("hbv_system2_full_positions_ge_30x_percent", 0),
                    row.get("hbv_system1_unique_median_depth", 0),
                    row.get("hbv_system1_unique_positions_ge_10x_percent", 0),
                    row.get("hbv_system1_unique_positions_ge_30x_percent", 0),
                    row.get("hbv_system2_unique_median_depth", 0),
                    row.get("hbv_system2_unique_positions_ge_10x_percent", 0),
                    row.get("hbv_system2_unique_positions_ge_30x_percent", 0),
                    row.get("hbv_overlap_median_depth", 0),
                    row.get("hbv_overlap_positions_ge_10x_percent", 0),
                    row.get("hbv_overlap_positions_ge_30x_percent", 0),
                    row.get("hbv_system_balance_ratio", 0),
                    row.get("hbv_system_qc_warning", ""),
                ])

            ws.append(summary_row)

            details_ws.append([
                row.get("sample", ""),
                row.get("reference_length", 0),
                row.get("detected_target_start", ""),
                row.get("detected_target_end", ""),
                row.get("detected_target_length", 0),
                row.get("expected_target_length", 0),
                row.get("coverage_type", ""),
                row.get("reference_selection_mode", ""),
                row.get("adeno_target_min_depth", ""),
                row.get("consensus_source", ""),
                row.get("consensus_fallback_reason", ""),
            ])

            mapped_cell = ws.cell(row=ws.max_row, column=4)
            mapped_reads = row.get("mapped_reads", 0)
            mapped_cell.fill = red_fill if mapped_reads < 100 else green_fill

            coverage_cell = ws.cell(row=ws.max_row, column=5)
            coverage = row.get("target_or_genome_coverage_percent", 0)
            coverage_fill(coverage_cell, coverage, red_fill, yellow_fill, green_fill)

            qc_cell = ws.cell(row=ws.max_row, column=15)
            qc_verdict = str(row.get("qc_verdict", "") or "").upper()
            qc_fill(qc_cell, qc_verdict, red_fill, yellow_fill, green_fill)

            # Adenovirus summary coloring: a LOW final QC row must not look
            # reassuringly green simply because one isolated metric (for
            # example mapped reads or the detected target span) passed. Keep
            # genuinely failing metrics red and downgrade otherwise-passing
            # key metrics to yellow when the overall Adenovirus QC is LOW.
            if virus_type == "Adenovirus":
                median_depth_cell = ws.cell(row=ws.max_row, column=7)
                n_percent_cell = ws.cell(row=ws.max_row, column=13)
                recovery_cell = ws.cell(row=ws.max_row, column=14)

                median_depth = float(row.get("median_depth", 0) or 0)
                n_percent_raw = row.get("n_percent", 100)
                n_percent = float(100 if n_percent_raw is None else n_percent_raw)
                recovery = float(row.get("consensus_recovery_percent", 0) or 0)

                if median_depth < 10:
                    median_depth_cell.fill = red_fill
                elif median_depth < 30:
                    median_depth_cell.fill = yellow_fill
                else:
                    median_depth_cell.fill = green_fill

                if n_percent > 10:
                    n_percent_cell.fill = red_fill
                elif n_percent > 2:
                    n_percent_cell.fill = yellow_fill
                else:
                    n_percent_cell.fill = green_fill

                coverage_fill(recovery_cell, recovery, red_fill, yellow_fill, green_fill)

                if qc_verdict == "LOW":
                    # Passing individual values remain visible as yellow, not
                    # green, so the row cannot be mistaken for a good sample.
                    if mapped_reads >= 100:
                        mapped_cell.fill = yellow_fill
                    if coverage >= 90:
                        coverage_cell.fill = yellow_fill
                    if median_depth >= 30:
                        median_depth_cell.fill = yellow_fill
                    if n_percent <= 2:
                        n_percent_cell.fill = yellow_fill
                    if recovery >= 90:
                        recovery_cell.fill = yellow_fill

            if virus_type == "Hepatitis B Virus":
                # Highlight the unique-region system metrics because those drive QC.
                for col in [24, 25, 26, 27, 28, 29]:
                    cell = ws.cell(row=ws.max_row, column=col)
                    try:
                        value = float(cell.value)
                    except Exception:
                        value = 0.0

                    if "Median Depth" in str(ws.cell(row=1, column=col).value):
                        if value >= 30:
                            cell.fill = green_fill
                        elif value >= 10:
                            cell.fill = yellow_fill
                        else:
                            cell.fill = red_fill
                    else:
                        coverage_fill(cell, value, red_fill, yellow_fill, green_fill)

        autosize_worksheet(ws)
        autosize_worksheet(details_ws)

        wb.save(excel_path)
        print(f"Reference usage Excel saved to {excel_path}")

    except Exception as e:
        print(f"Error logging reference usage: {e}")

def log_hcv_resistance_results(hcv_rows, output_folder, run_tag):
    try:
        ensure_dir(output_folder, must_be_writable=True)
        with open(os.path.join(output_folder, f"HCV_VERSION_{run_tag}.txt"), "w", encoding="utf-8") as vf:
            vf.write(f"{HCV_SCRIPT_VERSION}\nscript={os.path.abspath(__file__)}\n")
        excel_path = os.path.join(output_folder, f"hcv_resistance_results_{run_tag}.xlsx")
        generate_hcv_html_reports(hcv_rows, output_folder, run_tag)

        wb = Workbook()
        ws = wb.active
        ws.title = "HCV Resistance"

        ws.append(["Hepatit C virus antiviralresistensbedömning"])
        ws.append([f"Script version: {HCV_SCRIPT_VERSION}"])
        ws.append([f"Run: {run_tag}"])
        ws.append(["Resultat"])
        ws.append([])

        headers = [
            "Prov-ID",
            "Kvalitet",
            "Bedömning",
            "M28",
            "Q30",
            "L31",
            "H58",
            "Y93",
            "Mutationer < Täckning",
            "Post-trimning",
            "Readmappningsdistribution",
            "Median nollskild täckning",
            "Subtype/Genotype Reference",
            "Best Reference",
            "BAM File",
            "NS5A start used",
            "Script version",
        ]

        ws.append(headers)

        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        for row in hcv_rows:
            ws.append([
                row.get("sample", ""),
                row.get("hcv_quality", ""),
                row.get("hcv_interpretation", ""),
                row.get("hcv_M28", " "),
                row.get("hcv_Q30", " "),
                row.get("hcv_L31", " "),
                row.get("hcv_H58", " "),
                row.get("hcv_Y93", " "),
                row.get("hcv_low_coverage", " "),
                row.get("hcv_post_trim", ""),
                row.get("hcv_mapping_distribution", ""),
                row.get("hcv_median_nonzero_coverage", ""),
                row.get("hcv_subtype", ""),
                row.get("best_reference", ""),
                row.get("bam_file", ""),
                row.get("hcv_ns5a_start_used", ""),
                row.get("hcv_script_version", ""),
            ])

            quality_cell = ws.cell(row=ws.max_row, column=2)
            interp_cell = ws.cell(row=ws.max_row, column=3)

            if quality_cell.value == "Not Included":
                quality_cell.fill = yellow_fill
            elif quality_cell.value == "✓":
                quality_cell.fill = green_fill
            else:
                quality_cell.fill = red_fill

            if interp_cell.value == "Not Included":
                interp_cell.fill = yellow_fill
            elif interp_cell.value == "Resistent":
                interp_cell.fill = red_fill
            elif interp_cell.value == "RAS minoritetsvariant":
                interp_cell.fill = yellow_fill
            elif interp_cell.value == "Känslig":
                interp_cell.fill = green_fill
            else:
                interp_cell.fill = yellow_fill

        debug_ws = wb.create_sheet("HCV Debug")
        debug_ws.append([
            "Sample", "Best Reference", "Subtype", "NS5A start used",
            "Position", "AA position", "Reference codon", "Reference AA", "Reportable AAs",
            "Metadata NS5A start", "Selected offset", "Coordinate method", "Motif score",
            "Codon range", "Codon depth", "AA distribution", "Reported >=15% calls",
            "Min depth cutoff", "Report threshold"
        ])
        for row in hcv_rows:
            for detail in row.get("hcv_debug_details", []) or []:
                debug_ws.append([
                    row.get("sample", ""),
                    row.get("best_reference", ""),
                    row.get("hcv_subtype", ""),
                    row.get("hcv_ns5a_start_used", ""),
                    detail.get("position", ""),
                    detail.get("aa_position", ""),
                    detail.get("reference_codon", ""),
                    detail.get("reference_aa", ""),
                    detail.get("reportable_aas", ""),
                    detail.get("metadata_ns5a_start", ""),
                    detail.get("selected_offset", ""),
                    detail.get("coordinate_method", ""),
                    detail.get("motif_score", ""),
                    detail.get("codon_range", ""),
                    detail.get("depth", ""),
                    detail.get("aa_distribution", ""),
                    detail.get("reported_calls", ""),
                    HCV_MIN_COVERAGE,
                    f">={round(HCV_MIN_VARIANT_FREQUENCY * 100, 1)}%",
                ])

        autosize_worksheet(ws)
        autosize_worksheet(debug_ws)
        wb.save(excel_path)
        mirror_path = os.path.join(output_folder, f"hcv_resistance_results_PRODUCTION_DEBUG_COPY_{run_tag}.xlsx")
        wb.save(mirror_path)
        with open(os.path.join(output_folder, f"HCV_PRODUCTION_SENTINEL_{run_tag}.txt"), "w", encoding="utf-8") as sf:
            sf.write(f"{HCV_SCRIPT_VERSION}\nscript={os.path.abspath(__file__)}\nexcel={excel_path}\nmirror={mirror_path}\n")
        print(f"HCV resistance Excel saved to {excel_path}")
        print(f"HCV debug mirror Excel saved to {mirror_path}")

    except Exception as e:
        print(f"Error logging HCV resistance results: {e}")

def log_hbv_mutation_results(rows, output_folder, run_tag):
    """Write a streamlined Swedish Excel workbook and one Swedish HTML report per sample."""
    ensure_dir(output_folder, must_be_writable=True)
    excel_path = os.path.join(output_folder, f"hbv_resistensresultat_{run_tag}.xlsx")
    html_dir = os.path.join(run_samples_folder(output_folder, run_tag), "HBV_HTML_reports")
    ensure_dir(html_dir, must_be_writable=True)

    wb = Workbook()
    overview = wb.active
    overview.title = "Översikt"
    medications = wb.create_sheet("Läkemedel")
    mutations = wb.create_sheet("Mutationer")
    combinations = wb.create_sheet("Kombinationsbedömning")
    technical = wb.create_sheet("Teknisk QC")
    panel = wb.create_sheet("Paneldefinition")

    dark_blue = "1F4E78"
    light_blue = "D9EAF7"
    pale_blue = "EAF3F8"
    green = "C6EFCE"
    yellow = "FFEB9C"
    orange = "FCE4D6"
    red = "FFC7CE"
    grey = "E7E6E6"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    overview.merge_cells("A1:M1")
    overview["A1"] = "Hepatit B-virus – resistens och immunologisk escape"
    overview["A1"].font = Font(size=18, bold=True, color=white)
    overview["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    overview["A1"].alignment = Alignment(horizontal="left", vertical="center")
    overview.row_dimensions[1].height = 30
    overview.merge_cells("A2:M2")
    overview["A2"] = (
        f"Läsnivåtröskel ≥{HBV_MIN_VARIANT_FREQUENCY * 100:.1f}% | "
        f"Minsta kodontäckning {HBV_MIN_COVERAGE} | Körning {run_tag}"
    )
    overview["A2"].font = Font(italic=True, color="4F5B62")

    overview_headers = [
        "Prov-ID", "Kvalitet", "Samlad bedömning", "Bästa referens",
        "Prioriterade RT-fynd", "Övriga RT-fynd", "Escape-fynd", "Terapiassocierade S-fynd", "Övriga S-fynd",
        "Kombinationsbedömning", "Läkemedelsöversikt", "Otillräcklig täckning", "HTML-rapport",
    ]
    overview.append([])
    overview.append(overview_headers)

    medications.append(["Prov-ID", "Läkemedel", "Bedömning", "Påvisade mutationer", "Högsta andel (%)"])
    combinations.append(["Prov-ID", "Läkemedel/mönster", "Bedömning", "Mutationer i provet", "Tolkning", "Evidensnivå"])
    mutations.append([
        "Prov-ID", "Gen", "Mutation", "Andel (%)", "Population", "Mutantläsningar",
        "Kodondjup", "Fyndtyp", "Kategori", "Läkemedelsassociationer", "Bästa referens",
    ])
    technical.append([
        "Prov-ID", "Kvalitet", "Samlad bedömning", "Bästa referens", "Otillräcklig täckning",
        "BAM-fil", "Minsta kodontäckning", "Rapporteringströskel", "Scriptversion", "Körning",
    ])
    panel.append(["Gen", "Position", "Rapporterbara substitutioner", "Fyndtyp", "Kategori/läkemedel"])

    html_links = []
    swedish_drugs = {"Lamivudine": "Lamivudin", "Adefovir": "Adefovir", "Entecavir": "Entekavir", "Tenofovir": "Tenofovir"}

    for row in rows:
        sample = row.get("sample", "")
        calls = hbv_detected_calls(row)
        rt_calls = [x for x in calls if x["gene"] == "RT"]
        s_calls = [x for x in calls if x["gene"] == "S"]
        escape_calls = []
        therapy_surface_calls = []
        other_surface_calls = []
        for item in s_calls:
            labels = set(HBV_S_ESCAPE_ASSOCIATIONS.get(item.get("mutation", ""), {}).get("associations", ()) or ())
            if labels - {"Terapiassocierad"}:
                escape_calls.append(item)
            elif "Terapiassocierad" in labels or item.get("flag") == "THERAPY_SURFACE":
                therapy_surface_calls.append(item)
            else:
                other_surface_calls.append(item)
        priority = [x for x in rt_calls if x.get("flag") == "PRIORITY"]
        other_rt = [x for x in rt_calls if x.get("flag") != "PRIORITY"]

        def compact(items):
            return "; ".join(f"{x['mutation']} ({x['frequency']} %)" for x in items) or "—"

        medication_parts = []
        for drug in HBV_MEDICATIONS:
            _, mutation_text, _ = hbv_status_for_medication(row, drug)
            if mutation_text:
                medication_parts.append(f"{swedish_drugs.get(drug, drug)}: {mutation_text}")
        medication_overview = " | ".join(medication_parts) or "Inga läkemedelsassocierade mutationer"

        filename = f"{safe_report_filename(sample)}_{run_tag}_HBV.html"
        html_path = os.path.join(html_dir, filename)
        write_hbv_html_report(row, html_path, run_tag)
        html_links.append((sample, filename))

        overview.append([
            sample,
            row.get("hbv_quality", ""),
            hbv_swedish_interpretation(row.get("hbv_interpretation", "")),
            row.get("best_reference", ""),
            compact(priority), compact(other_rt), compact(escape_calls), compact(therapy_surface_calls), compact(other_surface_calls), hbv_combination_summary_text(row), medication_overview,
            row.get("hbv_low_coverage_positions", " ") or "—",
            filename,
        ])
        link_cell = overview.cell(overview.max_row, 13)
        link_cell.hyperlink = os.path.relpath(html_path, output_folder)
        link_cell.style = "Hyperlink"

        for pattern in hbv_combination_interpretations(row):
            combinations.append([
                sample, pattern.get("category", ""), pattern.get("title", ""),
                pattern.get("mutations", "") or "—", pattern.get("interpretation", ""),
                {"major": "Dominerande etablerat mönster", "minor": "Etablerat mönster i minoritetspopulation", "caution": "Försiktighet/kontextberoende", "none": "Inget mönster", "unknown": "Ej bedömbar"}.get(pattern.get("level"), pattern.get("level", "")),
            ])

        for drug in HBV_MEDICATIONS:
            status, mutation_text, _ = hbv_status_for_medication(row, drug)
            frequencies = [x["frequency"] for x in calls if drug in x.get("medications", {})]
            medications.append([
                sample, swedish_drugs.get(drug, drug), status, mutation_text or "—",
                max(frequencies) if frequencies else "",
            ])

        if calls:
            for call in calls:
                meds = "; ".join(
                    f"{swedish_drugs.get(drug, drug)}: {hbv_swedish_evidence_label(evidence)}"
                    for drug, evidence in call.get("medications", {}).items()
                ) or "—"
                flag_sv = {
                    "PRIORITY": "Prioriterat",
                    "ESTABLISHED": "Etablerat",
                    "CAUTION": "Kontextberoende/försiktighet",
                    "ESCAPE": "Escape-associerat",
                    "HIGH_IMPACT": "Escape-associerat – högre påverkan",
                    "THERAPY_SURFACE": "Terapiassocierad ytantigenförändring",
                }.get(call.get("flag"), call.get("flag") or "—")
                mutations.append([
                    sample, call["gene"], call["mutation"], call["frequency"], call["population"],
                    call["count"], call["depth"], flag_sv, hbv_translate_free_text(call.get("category", "")), meds,
                    row.get("best_reference", ""),
                ])
        else:
            mutations.append([sample, "—", "Inga panelmutationer påvisade", "", "", "", "", "", "", "", row.get("best_reference", "")])

        technical.append([
            sample, row.get("hbv_quality", ""), hbv_swedish_interpretation(row.get("hbv_interpretation", "")),
            row.get("best_reference", ""), row.get("hbv_low_coverage_positions", " ") or "—",
            row.get("bam_file", ""), HBV_MIN_COVERAGE, f">={HBV_MIN_VARIANT_FREQUENCY * 100:.1f}%",
            row.get("hbv_mutation_script_version", ""), run_tag,
        ])

    for aa_pos, rule in HBV_RT_PANEL.items():
        panel.append(["RT", rule.get("label", f"rt{aa_pos}"), ", ".join(sorted(rule.get("aas", set()))), rule.get("flag", ""), hbv_translate_free_text(rule.get("drugs", ""))])
    for aa_pos, rule in HBV_S_ESCAPE_PANEL.items():
        finding_type = "Terapiassocierad ytantigenförändring" if rule.get("flag") == "THERAPY_SURFACE" else "Escape-associerat"
        panel.append(["S", rule.get("label", f"s{aa_pos}"), ", ".join(sorted(rule.get("aas", set()))), finding_type, hbv_translate_free_text(rule.get("category", ""))])

    for ws in [overview, medications, combinations, mutations, technical, panel]:
        header_row = 4 if ws is overview else 1
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=dark_blue)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row_cells in ws.iter_rows(min_row=header_row + 1):
            for cell in row_cells:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.sheet_view.showGridLines = False

    for row_idx in range(5, overview.max_row + 1):
        quality = overview.cell(row_idx, 2)
        assessment = overview.cell(row_idx, 3)
        quality.fill = PatternFill("solid", fgColor=green if quality.value == "✓" else red)
        text_value = str(assessment.value or "").lower()
        if "inga panelmutationer" in text_value:
            assessment.fill = PatternFill("solid", fgColor=green)
        elif "ej bedömbar" in text_value:
            assessment.fill = PatternFill("solid", fgColor=grey)
        elif "minoritet" in text_value:
            assessment.fill = PatternFill("solid", fgColor=yellow)
        else:
            assessment.fill = PatternFill("solid", fgColor=orange)

    overview.column_dimensions["A"].width = 36
    overview.column_dimensions["B"].width = 11
    overview.column_dimensions["C"].width = 35
    overview.column_dimensions["D"].width = 16
    for col in ["E", "F", "G", "H", "I"]:
        overview.column_dimensions[col].width = 30
    overview.column_dimensions["J"].width = 48
    overview.column_dimensions["K"].width = 58
    overview.column_dimensions["L"].width = 35
    overview.column_dimensions["M"].width = 34

    for ws in [medications, combinations, mutations, technical, panel]:
        autosize_worksheet(ws)
    medications.column_dimensions["D"].width = 65
    combinations.column_dimensions["A"].width = 28
    combinations.column_dimensions["B"].width = 26
    combinations.column_dimensions["C"].width = 38
    combinations.column_dimensions["D"].width = 55
    combinations.column_dimensions["E"].width = 85
    combinations.column_dimensions["F"].width = 28
    mutations.column_dimensions["I"].width = 38
    mutations.column_dimensions["J"].width = 52
    technical.column_dimensions["F"].width = 75
    panel.column_dimensions["E"].width = 60

    wb.save(excel_path)

    index_path = os.path.join(html_dir, f"index_{run_tag}.html")
    links = "".join(
        f"<li><a href='{html_escape(filename)}'>{html_escape(str(sample))}</a></li>"
        for sample, filename in html_links
    )
    with open(index_path, "w", encoding="utf-8") as handle:
        handle.write(
            "<!DOCTYPE html><html lang='sv'><head><meta charset='UTF-8'>"
            "<title>HBV-rapporter</title><style>body{font-family:Arial;margin:40px;max-width:900px}"
            "li{margin:8px 0}a{color:#185b8d}</style></head><body>"
            f"<h1>HBV-rapporter – {html_escape(run_tag)}</h1><ul>{links}</ul></body></html>"
        )

    print(f"Svensk HBV-resistensrapport sparad: {excel_path}")
    print(f"HTML-rapporter sparade i: {html_dir}")
