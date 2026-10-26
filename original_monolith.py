import os
import sys
import subprocess
import re
import gzip
import statistics
from html import escape as html_escape
import base64
import mimetypes
import shutil
import threading
from datetime import datetime
from PySide6.QtCore import QObject, QPropertyAnimation, QEasingCurve, QThread, Qt, Signal, Slot, QRectF
from PySide6.QtGui import (
    QColor, QFont, QTextCursor, QPainter, QPainterPath,
    QLinearGradient, QPen, QPixmap,
)
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QFileDialog, QFormLayout,
    QFrame, QGraphicsDropShadowEffect, QHBoxLayout, QLabel, QLineEdit,
    QMainWindow, QMessageBox, QPlainTextEdit, QProgressBar, QPushButton,
    QSizePolicy, QVBoxLayout, QWidget,
)
# Generated from aurora_mainwindow_fixed(1).ui with pyside6-uic.
# Compile resources.qrc with pyside6-rcc so the banner image is registered.
try:
    import resources_rc  # noqa: F401
    from ui_aurora_mainwindow import Ui_MainWindow
except ImportError as exc:
    raise RuntimeError(
        'Missing generated Qt files. Run:\n'
        '  pyside6-rcc resources.qrc -o resources_rc.py\n'
        '  pyside6-uic \"aurora_mainwindow_fixed(1).ui\" -o ui_aurora_mainwindow.py'
    ) from exc

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

try:
    import pysam
except ImportError:
    pysam = None


def app_base_dir():
    return os.path.dirname(os.path.abspath(__file__))

def ensure_dir(dirpath, must_be_writable=False):
    os.makedirs(dirpath, exist_ok=True)
    if must_be_writable:
        try:
            testfile = os.path.join(dirpath, ".write_test.tmp")
            with open(testfile, "w", encoding="utf-8") as f:
                f.write("ok")
            os.remove(testfile)
        except Exception as e:
            raise PermissionError(f"Cannot write to '{dirpath}': {e}")

def which_or_none(exe_name):
    try:
        import shutil
        return shutil.which(exe_name)
    except Exception:
        return None

def check_required_tools():
    missing = []
    for tool in ["minimap2", "samtools", "medaka_consensus"]:
        if which_or_none(tool) is None:
            missing.append(tool)

    if missing:
        raise FileNotFoundError(
            "Missing required tools on PATH: " + ", ".join(missing) +
            "\nPlease install them and make sure they are available in your shell PATH."
        )

def path_exists_any(*paths):
    for p in paths:
        if p and os.path.exists(p):
            return p
    return None

def safe_remove(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass

# FILEfiltering

def is_real_fastq_gz(path):
    name = os.path.basename(path)
    if name.startswith("._") or name.startswith("."):
        return False
    return name.endswith(".fastq.gz")

def is_real_fasta(path):
    name = os.path.basename(path)
    if name.startswith("._") or name.startswith("."):
        return False
    return name.endswith(".fa") or name.endswith(".fasta")

def find_files_recursively(directory, predicate):
    matched = []
    try:
        for root, _, files in os.walk(directory):
            for file in files:
                full = os.path.join(root, file)
                try:
                    if predicate(full):
                        matched.append(full)
                except Exception:
                    continue
    except Exception as e:
        print(f"Error finding files in {directory}: {e}")
    return matched

#FASTQCount

def count_fastq_bases(fastq_gz_path):
    total_bases = 0
    try:
        with gzip.open(fastq_gz_path, "rt", encoding="utf-8", errors="replace") as f:
            for i, line in enumerate(f):
                if i % 4 == 1:
                    total_bases += len(line.strip())
    except Exception as e:
        print(f"Error counting FASTQ bases in {fastq_gz_path}: {e}")
    return total_bases

# SAM/BAMParse

LN_RE = re.compile(r"\bLN:(\d+)")
SN_RE = re.compile(r"\bSN:([^\t]+)")
CIGAR_RE = re.compile(r"(\d+)([MIDNSHP=X])")

def reference_consumed_by_cigar(cigar):
    ref_consumed = 0
    for length_str, op in CIGAR_RE.findall(cigar):
        length = int(length_str)
        if op in ["M", "=", "X", "D", "N"]:
            ref_consumed += length
    return ref_consumed

def query_aligned_consumed_by_cigar(cigar):
    query_consumed = 0
    for length_str, op in CIGAR_RE.findall(cigar):
        length = int(length_str)
        if op in ["M", "=", "X", "I"]:
            query_consumed += length
    return query_consumed

def evaluate_alignment_quality(output_sam):
    """Score a reference using primary-alignment identity, not aligned length alone.

    The previous score summed reference-consuming CIGAR operations and divided by
    reference length. For full-length HBV reads, several genotype references can
    therefore receive nearly identical scores and whichever file is encountered
    first may win. Here we use the SAM NM edit-distance tag to estimate matching
    query bases, while ignoring secondary/supplementary alignments.
    """
    ref_lengths = {}
    ref_mapped_reads = {}
    ref_aligned_query_bases = {}
    ref_matching_query_bases = {}
    ref_mapq_sum = {}

    try:
        with open(output_sam, "r", encoding="utf-8", errors="replace") as sam_file:
            for line in sam_file:
                if line.startswith("@SQ"):
                    ln_m = LN_RE.search(line)
                    sn_m = SN_RE.search(line)
                    if ln_m and sn_m:
                        ref_lengths[sn_m.group(1)] = int(ln_m.group(1))
                    continue

                if not line or line.startswith("@"):
                    continue

                fields = line.rstrip("\n").split("\t")
                if len(fields) < 11:
                    continue

                try:
                    flag = int(fields[1])
                    mapq = int(fields[4])
                except ValueError:
                    continue

                ref_name = fields[2]
                cigar = fields[5]

                # Skip unmapped, secondary, supplementary and QC-failed records.
                if flag & (0x4 | 0x100 | 0x200 | 0x800):
                    continue
                if ref_name == "*" or cigar == "*":
                    continue

                query_aligned = query_aligned_consumed_by_cigar(cigar)
                if query_aligned <= 0:
                    continue

                nm = 0
                for tag in fields[11:]:
                    if tag.startswith("NM:i:"):
                        try:
                            nm = int(tag[5:])
                        except ValueError:
                            nm = 0
                        break

                matching_bases = max(query_aligned - nm, 0)

                ref_mapped_reads[ref_name] = ref_mapped_reads.get(ref_name, 0) + 1
                ref_aligned_query_bases[ref_name] = ref_aligned_query_bases.get(ref_name, 0) + query_aligned
                ref_matching_query_bases[ref_name] = ref_matching_query_bases.get(ref_name, 0) + matching_bases
                ref_mapq_sum[ref_name] = ref_mapq_sum.get(ref_name, 0) + mapq

        best_ref = None
        best_rank = None

        for ref_name, matching_bases in ref_matching_query_bases.items():
            aligned_bases = ref_aligned_query_bases.get(ref_name, 0)
            identity = matching_bases / aligned_bases if aligned_bases else 0.0
            mapped_reads = ref_mapped_reads.get(ref_name, 0)
            mean_mapq = ref_mapq_sum.get(ref_name, 0) / mapped_reads if mapped_reads else 0.0

            # Reference choice must primarily reflect sequence identity. Using
            # total matching bases first can favour a slightly poorer reference
            # when alignment length differs by only a few bases.
            rank = (identity, matching_bases, mapped_reads, mean_mapq, ref_name)
            if best_rank is None or rank > best_rank:
                best_rank = rank
                best_ref = ref_name

        if not best_ref:
            return 0.0, 0, None, 0, 0

        aligned_bases = ref_aligned_query_bases.get(best_ref, 0)
        matching_bases = ref_matching_query_bases.get(best_ref, 0)
        identity_percent = round((matching_bases / aligned_bases) * 100, 6) if aligned_bases else 0.0

        # The caller compares this scalar across one-reference SAM files. Put
        # identity in the dominant component and use matching support only as a
        # fractional tie-breaker. Exact ref-D reads must therefore select ref_d.
        selection_score = (identity_percent * 1_000_000_000.0) + float(matching_bases)

        return (
            selection_score,
            ref_mapped_reads.get(best_ref, 0),
            best_ref,
            aligned_bases,
            ref_lengths.get(best_ref, 0)
        )

    except Exception as e:
        print(f"Error evaluating alignment quality for {output_sam}: {e}")
        return 0.0, 0, None, 0, 0

def calculate_reference_breadth_coverage_percent(bam_file, ref_name, ref_length):
    if not bam_file or not os.path.exists(bam_file):
        return 0.0

    if not ref_name or not ref_length:
        return 0.0

    covered_positions = 0

    try:
        cmd = ["samtools", "depth", "-a", "-r", ref_name, bam_file]

        with subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True
        ) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue

                try:
                    depth = int(fields[2])
                    if depth > 0:
                        covered_positions += 1
                except ValueError:
                    continue

            _, stderr = proc.communicate()

            if proc.returncode != 0:
                print(f"samtools depth warning/error for {bam_file}: {stderr}")
                return 0.0

        return round((covered_positions / ref_length) * 100, 2) if ref_length > 0 else 0.0

    except Exception as e:
        print(f"Error calculating reference breadth coverage for {bam_file}: {e}")
        return 0.0

def empty_depth_stats():
    return {
        "mean_depth": 0.0,
        "median_depth": 0.0,
        "min_depth": 0,
        "positions_ge_10x_percent": 0.0,
        "positions_ge_30x_percent": 0.0,
    }

def calculate_depth_stats(bam_file, ref_name, region=None):
    depths = []

    if not bam_file or not os.path.exists(bam_file) or not ref_name:
        return empty_depth_stats()

    try:
        if region:
            region_string = f"{ref_name}:{region[0]}-{region[1]}"
            cmd = ["samtools", "depth", "-a", "-r", region_string, bam_file]
        else:
            cmd = ["samtools", "depth", "-a", "-r", ref_name, bam_file]

        with subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True
        ) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue

                try:
                    depths.append(int(fields[2]))
                except ValueError:
                    continue

            _, stderr = proc.communicate()

            if proc.returncode != 0:
                print(f"samtools depth warning/error for {bam_file}: {stderr}")

        if not depths:
            return empty_depth_stats()

        total = len(depths)

        return {
            "mean_depth": round(sum(depths) / total, 2),
            "median_depth": round(statistics.median(depths), 2),
            "min_depth": min(depths),
            "positions_ge_10x_percent": round((sum(1 for d in depths if d >= 10) / total) * 100, 2),
            "positions_ge_30x_percent": round((sum(1 for d in depths if d >= 30) / total) * 100, 2),
        }

    except Exception as e:
        print(f"Error calculating depth stats for {bam_file}: {e}")
        return empty_depth_stats()

def merge_depth_values(depth_values):
    if not depth_values:
        return empty_depth_stats()

    total = len(depth_values)

    return {
        "mean_depth": round(sum(depth_values) / total, 2),
        "median_depth": round(statistics.median(depth_values), 2),
        "min_depth": min(depth_values),
        "positions_ge_10x_percent": round((sum(1 for d in depth_values if d >= 10) / total) * 100, 2),
        "positions_ge_30x_percent": round((sum(1 for d in depth_values if d >= 30) / total) * 100, 2),
    }

def get_depth_values_for_region(bam_file, ref_name, region):
    depths = []

    if not bam_file or not os.path.exists(bam_file) or not ref_name or not region:
        return depths

    try:
        region_string = f"{ref_name}:{region[0]}-{region[1]}"
        cmd = ["samtools", "depth", "-a", "-r", region_string, bam_file]

        with subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True
        ) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue
                try:
                    depths.append(int(fields[2]))
                except ValueError:
                    continue

            _, stderr = proc.communicate()
            if proc.returncode != 0:
                print(f"samtools depth warning/error for {bam_file}: {stderr}")

    except Exception as e:
        print(f"Error getting depth values for region {region} in {bam_file}: {e}")

    return depths

def calculate_multi_region_depth_stats(bam_file, ref_name, regions):
    all_depths = []
    for region in regions:
        all_depths.extend(get_depth_values_for_region(bam_file, ref_name, region))
    return merge_depth_values(all_depths)

def detect_aligned_target_region_from_sam(sam_file, ref_name):
    starts = []
    ends = []

    try:
        with open(sam_file, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                if line.startswith("@"):
                    continue

                fields = line.rstrip("\n").split("\t")
                if len(fields) < 6:
                    continue

                try:
                    flag = int(fields[1])
                    rname = fields[2]
                    pos = int(fields[3])
                    cigar = fields[5]
                except ValueError:
                    continue

                if (flag & 0x4) or rname != ref_name or cigar == "*":
                    continue

                ref_consumed = reference_consumed_by_cigar(cigar)
                if ref_consumed <= 0:
                    continue

                starts.append(pos)
                ends.append(pos + ref_consumed - 1)

        if not starts or not ends:
            return None, None, 0

        start = min(starts)
        end = max(ends)
        length = end - start + 1

        return start, end, length

    except Exception as e:
        print(f"Error detecting aligned target region from {sam_file}: {e}")
        return None, None, 0

def coverage_fill(cell, coverage_percent, red_fill, yellow_fill, green_fill):
    if coverage_percent >= 90:
        cell.fill = green_fill
    elif coverage_percent >= 40:
        cell.fill = yellow_fill
    else:
        cell.fill = red_fill

def qc_fill(cell, qc_verdict, red_fill, yellow_fill, green_fill):
    verdict = str(qc_verdict).upper()
    if verdict == "HIGH":
        cell.fill = green_fill
    elif verdict == "MEDIUM":
        cell.fill = yellow_fill
    elif verdict == "LOW":
        cell.fill = red_fill

# FASTA

def read_single_fasta_sequence(fasta_path):
    seq_parts = []
    try:
        with open(fasta_path, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                if not line.startswith(">"):
                    seq_parts.append(line.strip())
        return "".join(seq_parts).upper()
    except Exception as e:
        print(f"Error reading FASTA sequence from {fasta_path}: {e}")
        return ""

def read_named_fasta_sequence(fasta_path, ref_name):
    """Read one named sequence from a possibly multi-FASTA reference file."""
    if not fasta_path or not os.path.exists(fasta_path) or not ref_name:
        return ""

    wanted = str(ref_name).split()[0]
    seq_parts = []
    collecting = False
    try:
        with open(fasta_path, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                if line.startswith(">"):
                    current = line[1:].strip().split()[0]
                    if collecting:
                        break
                    collecting = current == wanted
                    continue
                if collecting:
                    seq_parts.append(line.strip())
    except Exception as e:
        print(f"Error reading reference sequence {ref_name} from {fasta_path}: {e}")
        return ""

    return re.sub(r"[^ACGTN]", "", "".join(seq_parts).upper())


def fasta_sequence_stats(fasta_path):
    seq = read_single_fasta_sequence(fasta_path)
    if not seq:
        return {
            "consensus_length": 0,
            "n_count": 0,
            "n_percent": 100.0,
            "non_n_length": 0,
        }

    seq = re.sub(r"[^ACGTN]", "", seq.upper())
    total = len(seq)
    n_count = seq.count("N")
    non_n = total - n_count

    return {
        "consensus_length": total,
        "n_count": n_count,
        "n_percent": round((n_count / total) * 100, 2) if total else 100.0,
        "non_n_length": non_n,
    }

CODON_TABLE = {
    "TTT":"F","TTC":"F","TTA":"L","TTG":"L",
    "TCT":"S","TCC":"S","TCA":"S","TCG":"S",
    "TAT":"Y","TAC":"Y","TAA":"*","TAG":"*",
    "TGT":"C","TGC":"C","TGA":"*","TGG":"W",

    "CTT":"L","CTC":"L","CTA":"L","CTG":"L",
    "CCT":"P","CCC":"P","CCA":"P","CCG":"P",
    "CAT":"H","CAC":"H","CAA":"Q","CAG":"Q",
    "CGT":"R","CGC":"R","CGA":"R","CGG":"R",

    "ATT":"I","ATC":"I","ATA":"I","ATG":"M",
    "ACT":"T","ACC":"T","ACA":"T","ACG":"T",
    "AAT":"N","AAC":"N","AAA":"K","AAG":"K",
    "AGT":"S","AGC":"S","AGA":"R","AGG":"R",

    "GTT":"V","GTC":"V","GTA":"V","GTG":"V",
    "GCT":"A","GCC":"A","GCA":"A","GCG":"A",
    "GAT":"D","GAC":"D","GAA":"E","GAG":"E",
    "GGT":"G","GGC":"G","GGA":"G","GGG":"G",
}

def translate_codon(codon):
    codon = codon.upper().replace("U", "T")
    if len(codon) != 3:
        return "X"
    if any(base not in "ACGT" for base in codon):
        return "X"
    return CODON_TABLE.get(codon, "X")

def circular_subseq(seq, start_1based, length):
    if not seq:
        return ""
    n = len(seq)
    start_idx = (start_1based - 1) % n
    out = []
    for i in range(length):
        out.append(seq[(start_idx + i) % n])
    return "".join(out)

def circular_codon(seq, start_1based):
    return circular_subseq(seq, start_1based, 3)

#QC

HBV_EXPECTED_GENOME_LENGTH = 3200

# Original primer windows:
# System 1 outer: 56F to 1827R
# System 1 nested: 251F to 1801R
# System 2: 1689F to 262R, circular wraparound

HBV_SYSTEM1_FULL_REGIONS = [(251, 1801)]
HBV_SYSTEM2_FULL_REGIONS = [(1689, HBV_EXPECTED_GENOME_LENGTH), (1, 262)]

HBV_SYSTEM1_UNIQUE_REGIONS = [(263, 1688)]
HBV_SYSTEM2_UNIQUE_REGIONS = [(1802, HBV_EXPECTED_GENOME_LENGTH), (1, 250)]
HBV_SYSTEM_OVERLAP_REGIONS = [(251, 262), (1689, 1801)]

def expected_target_length_for_qc(virus_type, detected_target_length):
    if virus_type == "Hepatitis B Virus":
        return HBV_EXPECTED_GENOME_LENGTH

    if virus_type == "Adenovirus":
        return detected_target_length or 0

    return detected_target_length or 0

def calculate_hbv_system_depth_stats(bam_file, ref_name):
    return {
        "system1_full": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM1_FULL_REGIONS),
        "system2_full": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM2_FULL_REGIONS),
        "system1_unique": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM1_UNIQUE_REGIONS),
        "system2_unique": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM2_UNIQUE_REGIONS),
        "overlap_only": calculate_multi_region_depth_stats(bam_file, ref_name, HBV_SYSTEM_OVERLAP_REGIONS),
    }

def empty_hbv_system_stats():
    return {
        "system1_full": empty_depth_stats(),
        "system2_full": empty_depth_stats(),
        "system1_unique": empty_depth_stats(),
        "system2_unique": empty_depth_stats(),
        "overlap_only": empty_depth_stats(),
    }

def hbv_system_qc_flags(hbv_system_stats):
    s1 = hbv_system_stats.get("system1_unique", empty_depth_stats())
    s2 = hbv_system_stats.get("system2_unique", empty_depth_stats())
    overlap = hbv_system_stats.get("overlap_only", empty_depth_stats())

    s1_median = s1.get("median_depth", 0)
    s2_median = s2.get("median_depth", 0)
    s1_10x = s1.get("positions_ge_10x_percent", 0)
    s2_10x = s2.get("positions_ge_10x_percent", 0)
    s1_30x = s1.get("positions_ge_30x_percent", 0)
    s2_30x = s2.get("positions_ge_30x_percent", 0)

    high_ok = (
        s1_median >= 30 and
        s2_median >= 30 and
        s1_10x >= 90 and
        s2_10x >= 90 and
        s1_30x >= 70 and
        s2_30x >= 70
    )

    medium_ok = (
        s1_median >= 10 and
        s2_median >= 10 and
        s1_10x >= 70 and
        s2_10x >= 70
    )

    warnings = []

    if s1_median < 10:
        warnings.append("System 1 unique region has low median depth <10x")
    elif s1_median < 30:
        warnings.append("System 1 unique region has moderate median depth <30x")

    if s2_median < 10:
        warnings.append("System 2 unique region has low median depth <10x")
    elif s2_median < 30:
        warnings.append("System 2 unique region has moderate median depth <30x")

    if s1_10x < 70:
        warnings.append(f"System 1 unique region has poor >=10x breadth ({s1_10x}%)")
    elif s1_10x < 90:
        warnings.append(f"System 1 unique region has partial >=10x breadth ({s1_10x}%)")

    if s2_10x < 70:
        warnings.append(f"System 2 unique region has poor >=10x breadth ({s2_10x}%)")
    elif s2_10x < 90:
        warnings.append(f"System 2 unique region has partial >=10x breadth ({s2_10x}%)")

    min_median = max(min(s1_median, s2_median), 1)
    max_median = max(s1_median, s2_median)
    balance_ratio = round(max_median / min_median, 2) if max_median > 0 else 0.0

    if s1_median == 0 or s2_median == 0:
        warnings.append("One HBV system unique region has zero median depth")
    elif balance_ratio >= 10:
        warnings.append(f"Strong imbalance between HBV systems based on unique regions ({balance_ratio}x)")
    elif balance_ratio >= 5:
        warnings.append(f"Moderate imbalance between HBV systems based on unique regions ({balance_ratio}x)")

    overlap_median = overlap.get("median_depth", 0)
    if overlap_median >= 30 and (s1_median < 30 or s2_median < 30):
        warnings.append("Overlap has good depth but at least one unique system region is weaker; QC uses unique regions")

    return {
        "high_ok": high_ok,
        "medium_ok": medium_ok,
        "warning": "; ".join(warnings) if warnings else "Both HBV systems acceptable by unique-region QC",
        "balance_ratio": balance_ratio,
    }

def determine_qc_verdict(
    virus_type,
    mapped_reads,
    coverage_percent,
    median_depth,
    n_percent,
    consensus_length,
    expected_length,
    hbv_system_high_ok=True,
    hbv_system_medium_ok=True,
    hbv_system_warning=""
):
    reasons = []

    if expected_length and expected_length > 0:
        recovery_percent = round((consensus_length / expected_length) * 100, 2)
    else:
        recovery_percent = 0.0

    if virus_type == "Hepatitis B Virus":
        coverage_name = "HBV genome recovery"
    elif virus_type == "Adenovirus":
        coverage_name = "Adenovirus target/amplicon recovery"
    else:
        coverage_name = "target recovery"

    high = (
        mapped_reads >= 100 and
        coverage_percent >= 90 and
        median_depth >= 30 and
        n_percent <= 2 and
        recovery_percent >= 90 and
        hbv_system_high_ok
    )

    medium = (
        mapped_reads >= 50 and
        coverage_percent >= 70 and
        median_depth >= 10 and
        n_percent <= 10 and
        recovery_percent >= 70 and
        hbv_system_medium_ok
    )

    if high:
        verdict = "HIGH"
    elif medium:
        verdict = "MEDIUM"
    else:
        verdict = "LOW"

    if mapped_reads < 50:
        reasons.append("low mapped read count")
    elif mapped_reads < 100:
        reasons.append("moderate mapped read count")
    else:
        reasons.append("good mapped read count")

    if coverage_percent < 70:
        reasons.append(f"low {coverage_name}")
    elif coverage_percent < 90:
        reasons.append(f"partial {coverage_name}")
    else:
        reasons.append(f"strong {coverage_name}")

    if median_depth < 10:
        reasons.append("low median depth")
    elif median_depth < 30:
        reasons.append("moderate median depth")
    else:
        reasons.append("good median depth")

    if n_percent > 10:
        reasons.append("high N content")
    elif n_percent > 2:
        reasons.append("moderate N content")
    else:
        reasons.append("low N content")

    if expected_length and expected_length > 0:
        if recovery_percent < 70:
            reasons.append(f"short consensus vs expected length ({recovery_percent}%)")
        elif recovery_percent < 90:
            reasons.append(f"partial consensus length recovery ({recovery_percent}%)")
        else:
            reasons.append(f"good consensus length recovery ({recovery_percent}%)")

    if virus_type == "Hepatitis B Virus":
        if hbv_system_high_ok:
            reasons.append("HBV System 1 and System 2 unique-region depth supports HIGH QC")
        elif hbv_system_medium_ok:
            reasons.append("HBV System 1 and System 2 unique-region depth supports MEDIUM QC only")
        else:
            reasons.append("HBV System 1/System 2 unique-region depth fails medium threshold")

        if hbv_system_warning:
            reasons.append(hbv_system_warning)

    return verdict, "; ".join(reasons), recovery_percent

# Medaka


# Medaka model auto-detection
MEDAKA_MODEL_HEADER_KEYS = (
    "model_version_id", "basecall_model", "basecaller_model", "model",
)

def normalise_dorado_model_to_medaka(raw_model):
    """Convert common Dorado model identifiers to Medaka model names."""
    if not raw_model:
        return None
    model = str(raw_model).strip().strip('"').strip("'")
    model = model.split(',')[0].split(';')[0]
    model = re.sub(r"^(?:dna|rna)_", "", model, flags=re.IGNORECASE)
    model = model.replace("@", "_")
    model = re.sub(r"r10\.4\.1", "r1041", model, flags=re.IGNORECASE)
    model = re.sub(r"e8\.2", "e82", model, flags=re.IGNORECASE)
    model = re.sub(r"r9\.4\.1", "r941", model, flags=re.IGNORECASE)
    model = re.sub(r"[^A-Za-z0-9_.-]+", "_", model)
    return model.strip("_") or None


def detect_medaka_model_from_fastq(fastq_gz_path, max_records=50):
    """Inspect early FASTQ headers and return the most frequent detected model."""
    from collections import Counter
    detected = Counter()
    patterns = [
        re.compile(r"(?:model_version_id|basecall_model|basecaller_model|model)=([^\s]+)", re.I),
        re.compile(r"(?:model_version_id|basecall_model|basecaller_model|model):([^\s]+)", re.I),
    ]
    try:
        with gzip.open(fastq_gz_path, "rt", encoding="utf-8", errors="replace") as handle:
            for record_index in range(max_records):
                header = handle.readline()
                if not header:
                    break
                handle.readline(); handle.readline(); handle.readline()
                if not header.startswith("@"):
                    continue
                for pattern in patterns:
                    match = pattern.search(header)
                    if match:
                        normalised = normalise_dorado_model_to_medaka(match.group(1))
                        if normalised:
                            detected[normalised] += 1
                        break
    except Exception as exc:
        print(f"[MODEL WARNING] Could not inspect {fastq_gz_path}: {exc}")
        return None
    return detected.most_common(1)[0][0] if detected else None


def choose_medaka_model(input_fastq, fallback_model, auto_detect_models=False):
    if not auto_detect_models:
        return fallback_model
    detected = detect_medaka_model_from_fastq(input_fastq)
    if detected:
        print(f"[MODEL] Auto-detected Medaka model for {os.path.basename(input_fastq)}: {detected}")
        return detected
    print(
        f"[MODEL WARNING] No model metadata detected in {os.path.basename(input_fastq)}; "
        f"using fallback model: {fallback_model}"
    )
    return fallback_model

IUPAC_BASES = {
    frozenset({"A"}): "A",
    frozenset({"C"}): "C",
    frozenset({"G"}): "G",
    frozenset({"T"}): "T",
    frozenset({"A", "G"}): "R",
    frozenset({"C", "T"}): "Y",
    frozenset({"G", "C"}): "S",
    frozenset({"A", "T"}): "W",
    frozenset({"G", "T"}): "K",
    frozenset({"A", "C"}): "M",
    frozenset({"C", "G", "T"}): "B",
    frozenset({"A", "G", "T"}): "D",
    frozenset({"A", "C", "T"}): "H",
    frozenset({"A", "C", "G"}): "V",
    frozenset({"A", "C", "G", "T"}): "N",
}


def generate_iub_consensus_from_bam(
    bam_file,
    ref_name,
    ref_path,
    output_folder,
    sample_name,
    variant_frequency=0.15,
    min_base_quality=20,
    min_mapping_quality=20,
):
    """Create a standalone IUPAC/IUB consensus FASTA from aligned reads.

    Bases present at or above variant_frequency are represented together with
    the standard IUPAC ambiguity code. Positions without usable read support
    are written as N. Insertions are not represented in this reference-length
    consensus.
    """
    if pysam is None:
        print("[IUB WARNING] pysam is unavailable; IUB consensus was not created.")
        return None

    reference_sequence = read_named_fasta_sequence(ref_path, ref_name)
    if not reference_sequence:
        print(f"[IUB WARNING] Could not read reference '{ref_name}' from {ref_path}.")
        return None

    reference_sequence = reference_sequence.upper().replace("U", "T")
    calls = ["N"] * len(reference_sequence)

    try:
        with pysam.AlignmentFile(bam_file, "rb") as bam:
            for column in bam.pileup(
                ref_name,
                0,
                len(reference_sequence),
                truncate=True,
                stepper="samtools",
                min_base_quality=min_base_quality,
                min_mapping_quality=min_mapping_quality,
                max_depth=1000000,
            ):
                pos = column.reference_pos
                if pos < 0 or pos >= len(calls):
                    continue

                counts = {"A": 0, "C": 0, "G": 0, "T": 0}
                for pileup_read in column.pileups:
                    if pileup_read.is_del or pileup_read.is_refskip:
                        continue
                    alignment = pileup_read.alignment
                    if alignment.is_unmapped or alignment.is_secondary or alignment.is_supplementary:
                        continue
                    query_pos = pileup_read.query_position
                    if query_pos is None or not alignment.query_sequence:
                        continue
                    base = alignment.query_sequence[query_pos].upper().replace("U", "T")
                    if base in counts:
                        counts[base] += 1

                depth = sum(counts.values())
                if depth == 0:
                    continue

                represented = {
                    base for base, count in counts.items()
                    if count > 0 and (count / depth) >= variant_frequency
                }

                # Always retain the most abundant observed base if rounding or
                # an unusually mixed position leaves no base above threshold.
                if not represented:
                    represented = {max(counts, key=counts.get)}

                calls[pos] = IUPAC_BASES.get(frozenset(represented), "N")

        ensure_dir(output_folder, must_be_writable=True)
        output_path = os.path.join(output_folder, f"{sample_name}.consensus.iub.fasta")
        sequence = "".join(calls)
        with open(output_path, "w", encoding="utf-8") as handle:
            handle.write(
                f">{sample_name} IUB_consensus ref={ref_name} "
                f"threshold={variant_frequency:.3f}\n"
            )
            for start in range(0, len(sequence), 80):
                handle.write(sequence[start:start + 80] + "\n")

        print(f"IUB consensus completed. Output: {output_path}")
        return output_path
    except Exception as exc:
        print(f"[IUB WARNING] Could not create IUB consensus for {sample_name}: {exc}")
        return None


def generate_consensus(input_fastq, draft_fasta, medaka_model, output_folder, sample_name):
    try:
        medaka_output_folder = os.path.join(output_folder, f"{sample_name}_medaka")
        ensure_dir(medaka_output_folder, must_be_writable=True)

        cmd = [
            "medaka_consensus",
            "-i", input_fastq,
            "-d", draft_fasta,
            "-o", medaka_output_folder,
            "-m", medaka_model,
        ]

        env=os.environ.copy()
        env["PYTHONNOUSERSITE"] = "1"

        subprocess.run(cmd, check=True, env=env)

        medaka_out = os.path.join(medaka_output_folder, "consensus.fasta")
        if not (os.path.exists(medaka_out) and os.path.getsize(medaka_out) > 0):
            print(f"Medaka did not produce the expected output file: {medaka_out}")
            return None

        renamed = os.path.join(medaka_output_folder, f"{sample_name}.consensus.fasta")

        with open(medaka_out, "r", encoding="utf-8", errors="replace") as f:
            lines = f.read().splitlines()

        if not lines:
            print(f"Warning: consensus for {sample_name} is empty.")
            return None

        body = [line for line in lines if not line.startswith(">")]

        with open(renamed, "w", encoding="utf-8") as f:
            f.write(f">{sample_name}\n")
            f.write("\n".join(body) + ("\n" if body else ""))

        print(f"Medaka completed. Output: {renamed}")
        return renamed if os.path.getsize(renamed) > 0 else None

    except subprocess.CalledProcessError as e:
        print(f"Error in Medaka process for {input_fastq}: {e}")
        return None
    except Exception as e:
        print(f"Unexpected error in Medaka for {input_fastq}: {e}")
        return None


# HCV resistance / minor variant analysis
#
# This section calls HCV NS5A resistance-associated substitutions from the
# aligned reads in the BAM file. Do not use consensus alone for this, because
# low-frequency resistant variants can be below consensus level.
#
# Coordinate requirement:
# For each HCV reference sequence, the pipeline must know the 1-based reference
# coordinate of NS5A amino-acid/codon 1. You can provide this in either way:
#
# 1) Put tokens in the FASTA header, for example:
#    >HCV_1a_ref subtype=1A ns5a_codon1_start=1234
#
# 2) Put a file named hcv_reference_metadata.tsv next to the reference FASTA
#    files or next to this script. Required columns:
#    ref_name<TAB>subtype<TAB>ns5a_codon1_start
#
# The ref_name must match the SAM/FASTA sequence name exactly.

HCV_MIN_COVERAGE = 100
HCV_MIN_VARIANT_FREQUENCY = 0.15  # 15 percent reporting threshold; lower-frequency variants are hidden from the main report.
HCV_SCRIPT_VERSION = "Aurora3.1.0.2026_07_16"


def calculate_hcv_amplicon_breadth_coverage_percent(bam_file, ref_name, start_1based, end_1based):
    """Calculate HCV breadth only across the detected fixed-primer amplicon."""
    if not bam_file or not os.path.exists(bam_file) or not ref_name:
        return 0.0
    if not start_1based or not end_1based or end_1based < start_1based:
        return 0.0

    target_length = end_1based - start_1based + 1
    covered_positions = 0
    region_string = f"{ref_name}:{start_1based}-{end_1based}"

    try:
        cmd = ["samtools", "depth", "-a", "-r", region_string, bam_file]
        with subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue
                try:
                    if int(fields[2]) > 0:
                        covered_positions += 1
                except ValueError:
                    continue
            _, stderr = proc.communicate()
            if proc.returncode != 0:
                print(f"samtools depth warning/error for HCV target {bam_file}: {stderr}")
                return 0.0
    except Exception as exc:
        print(f"Error calculating HCV target breadth for {bam_file}: {exc}")
        return 0.0

    return round((covered_positions / target_length) * 100, 2) if target_length else 0.0


def hcv_amplicon_consensus_stats(fasta_path, start_1based, end_1based):
    """Calculate consensus QC statistics for the detected HCV amplicon only."""
    seq = read_single_fasta_sequence(fasta_path)
    if not seq or not start_1based or not end_1based or end_1based < start_1based:
        return {
            "consensus_length": 0,
            "n_count": 0,
            "n_percent": 100.0,
            "non_n_length": 0,
        }

    seq = re.sub(r"[^ACGTN]", "", seq.upper())
    target_length = end_1based - start_1based + 1

    # samtools/Medaka commonly retain reference coordinates and fill uncovered
    # sequence with N. If a tool emits only the amplicon, use that sequence as-is.
    if len(seq) >= end_1based:
        target_seq = seq[start_1based - 1:end_1based]
    elif len(seq) <= target_length + 20:
        target_seq = seq
    else:
        target_seq = seq[:target_length]

    total = len(target_seq)
    n_count = target_seq.count("N")
    non_n = total - n_count
    return {
        "consensus_length": non_n,
        "n_count": n_count,
        "n_percent": round((n_count / total) * 100, 2) if total else 100.0,
        "non_n_length": non_n,
    }


# Subtype-specific NS5A resistance panels.
# Each subtype only evaluates positions/substitutions considered reportable for
# that subtype in this simplified assay. Do not add a subtype here unless the
# positions and amino-acid substitutions have been clinically/locally validated.
#
# The Excel output still keeps the historical M28/Q30/L31/H58/Y93 columns for
# compatibility. For non-1A subtypes, the position labels in the debug sheet show
# the true reference residue, e.g. genotype 3a position 30 is A30, not Q30.
HCV_NS5A_PANEL_BY_SUBTYPE = {
    "1A": {
        28: {"label": "M28", "reportable_aas": {"A", "G", "T", "V"}},
        30: {"label": "Q30", "reportable_aas": {"H", "R"}},
        31: {"label": "L31", "reportable_aas": {"M", "V"}},
        58: {"label": "H58", "reportable_aas": {"D"}},
        93: {"label": "Y93", "reportable_aas": {"H", "N"}},
    },
    "1B": {
        31: {"label": "L31", "reportable_aas": {"F", "I", "M", "V"}},
        93: {"label": "Y93", "reportable_aas": {"H", "N"}},
    },
    "3A": {
        30: {"label": "A30", "reportable_aas": {"K"}},
        93: {"label": "Y93", "reportable_aas": {"H"}},
    },
}

# Compatibility tuple used for fixed Excel columns and error fallback messages.
HCV_NS5A_POSITIONS = tuple(sorted({
    aa_pos
    for subtype_panel in HCV_NS5A_PANEL_BY_SUBTYPE.values()
    for aa_pos in subtype_panel.keys()
}))


# Expected wild-type NS5A residues used only to validate/locate the reading frame
# in the selected reference. They are not resistance rules.
HCV_NS5A_REFERENCE_MOTIF_BY_SUBTYPE = {
    "1A": {28: "M", 30: "Q", 31: "L", 58: "H", 93: "Y"},
    "1B": {28: "M", 30: "R", 31: "L", 58: "P", 93: "Y"},
    "3A": {28: "M", 30: "A", 31: "L", 58: "P", 93: "Y"},
}

# Interpretation tiers for read-level RAS frequencies.
HCV_MINOR_RAS_MAX_FREQUENCY = 0.50

HCV_REFERENCE_METADATA = {
    # Built-in fallbacks for common HCV references seen in this project.
    # Local hcv_reference_metadata.tsv and FASTA-header metadata are still supported,
    # but these fallbacks prevent common accession rows from being missed.
    "NC_004102.1": {"subtype": "1A", "ns5a_codon1_start": 6258},
    "NC_004102": {"subtype": "1A", "ns5a_codon1_start": 6258},
    "AJ238799.1": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "AJ238799": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "M58335.1": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "M58335": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "AF011753.1": {"subtype": "1A", "ns5a_codon1_start": 6258},
    "AF011753": {"subtype": "1A", "ns5a_codon1_start": 6258},
    "D90208.1": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "D90208": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "NC_009824.1": {"subtype": "3A", "ns5a_codon1_start": 6258},
    "NC_009824": {"subtype": "3A", "ns5a_codon1_start": 6258},
    "AB030907.1": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "AB030907": {"subtype": "1B", "ns5a_codon1_start": 6258},
    "JQ065709.1": {"subtype": "3B", "ns5a_codon1_start": 6258},
    "JQ065709": {"subtype": "3B", "ns5a_codon1_start": 6258},
    "FJ462435.1": {"subtype": "4B", "ns5a_codon1_start": 6258},
    "FJ462435": {"subtype": "4B", "ns5a_codon1_start": 6258},
    "DQ988074.1": {"subtype": "4A", "ns5a_codon1_start": 6258},
    "DQ988074": {"subtype": "4A", "ns5a_codon1_start": 6258},
    "HQ639944.1": {"subtype": "2A", "ns5a_codon1_start": 6258},
    "HQ639944": {"subtype": "2A", "ns5a_codon1_start": 6258},
}


def hcv_subtype_from_name(name):
    text = str(name or "").lower().replace("_", " ").replace("-", " ")
    compact = text.replace(" ", "")
    subtype_patterns = [
        ("1A", ["1a", "gt1a", "genotype1a", "subtype=1a"]),
        ("1B", ["1b", "gt1b", "genotype1b", "subtype=1b"]),
        ("2A", ["2a", "gt2a", "genotype2a", "subtype=2a"]),
        ("3A", ["3a", "gt3a", "genotype3a", "subtype=3a"]),
        ("3B", ["3b", "gt3b", "genotype3b", "subtype=3b"]),
        ("4A", ["4a", "gt4a", "genotype4a", "subtype=4a"]),
        ("4B", ["4b", "gt4b", "genotype4b", "subtype=4b"]),
    ]
    for subtype, patterns in subtype_patterns:
        if any(pattern in compact for pattern in patterns):
            return subtype
    return "Unknown"


def read_fastq_stats(fastq_gz_path):
    read_count = 0
    total_bases = 0
    try:
        with gzip.open(fastq_gz_path, "rt", encoding="utf-8", errors="replace") as f:
            for i, line in enumerate(f):
                if i % 4 == 1:
                    seq_len = len(line.strip())
                    read_count += 1
                    total_bases += seq_len
    except Exception as e:
        print(f"Error reading FASTQ stats in {fastq_gz_path}: {e}")

    mean_len = round(total_bases / read_count, 1) if read_count else 0
    return read_count, mean_len


def calculate_median_nonzero_depth(bam_file, ref_name, region=None):
    depths = []
    if not bam_file or not os.path.exists(bam_file) or not ref_name:
        return 0

    try:
        region_string = (
            f"{ref_name}:{region[0]}-{region[1]}"
            if region and region[0] and region[1]
            else ref_name
        )
        cmd = ["samtools", "depth", "-r", region_string, bam_file]
        with subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True) as proc:
            for line in proc.stdout:
                fields = line.rstrip("\n").split("\t")
                if len(fields) < 3:
                    continue
                try:
                    depth = int(fields[2])
                    if depth > 0:
                        depths.append(depth)
                except ValueError:
                    continue
            _, stderr = proc.communicate()
            if proc.returncode != 0:
                print(f"samtools depth warning/error for {bam_file}: {stderr}")
    except Exception as e:
        print(f"Error calculating non-zero median depth for {bam_file}: {e}")

    return round(statistics.median(depths), 1) if depths else 0


def parse_hcv_metadata_tsv(tsv_path):
    metadata = {}
    if not tsv_path or not os.path.exists(tsv_path):
        return metadata

    try:
        with open(tsv_path, "r", encoding="utf-8", errors="replace") as f:
            header = None
            for line in f:
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                parts = line.split("\t")
                if header is None and any(x.lower() in ["ref_name", "reference", "subtype"] for x in parts):
                    header = [x.strip().lower() for x in parts]
                    continue

                if header:
                    row = dict(zip(header, parts))
                    ref_name = row.get("ref_name") or row.get("reference") or row.get("name")
                    subtype = row.get("subtype") or row.get("genotype") or "Unknown"
                    start = row.get("ns5a_codon1_start") or row.get("ns5a_start") or row.get("codon1_start")
                else:
                    if len(parts) < 3:
                        continue
                    ref_name, subtype, start = parts[0], parts[1], parts[2]

                if not ref_name or not start:
                    continue
                try:
                    metadata[ref_name] = {
                        "subtype": str(subtype).upper(),
                        "ns5a_codon1_start": int(start),
                    }
                except ValueError:
                    continue
    except Exception as e:
        print(f"Error reading HCV metadata TSV {tsv_path}: {e}")

    return metadata


def parse_hcv_metadata_from_fasta_header(fasta_path, ref_name):
    if not fasta_path or not os.path.exists(fasta_path) or not ref_name:
        return None

    try:
        with open(fasta_path, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                if not line.startswith(">"):
                    continue
                header = line[1:].strip()
                name_only = header.split()[0]
                if name_only != ref_name:
                    continue

                subtype_match = re.search(r"(?:subtype|genotype|gt)=([A-Za-z0-9]+)", header, re.IGNORECASE)
                start_match = re.search(r"(?:ns5a_codon1_start|ns5a_start|codon1_start)=(\d+)", header, re.IGNORECASE)
                if start_match:
                    subtype = subtype_match.group(1).upper() if subtype_match else hcv_subtype_from_name(header)
                    return {
                        "subtype": subtype,
                        "ns5a_codon1_start": int(start_match.group(1)),
                    }
    except Exception as e:
        print(f"Error parsing HCV FASTA header metadata from {fasta_path}: {e}")

    return None


def get_hcv_reference_metadata(ref_path, ref_name):
    """Return HCV reference metadata for the selected SAM reference name.

    Lookup is deliberately tolerant: exact accession, accession without version,
    local TSV next to the reference folder, FASTA header tokens, then built-ins.
    """
    names_to_try = []
    if ref_name:
        names_to_try.append(ref_name)
        names_to_try.append(str(ref_name).split()[0])
        names_to_try.append(str(ref_name).split('.')[0])

    candidates = []
    if ref_path:
        candidates.append(os.path.join(os.path.dirname(ref_path), "hcv_reference_metadata.tsv"))
    candidates.append(os.path.join(app_base_dir(), "hcv_reference_metadata.tsv"))

    for candidate in candidates:
        metadata = parse_hcv_metadata_tsv(candidate)
        for name in names_to_try:
            if name in metadata:
                return metadata[name]

    from_header = parse_hcv_metadata_from_fasta_header(ref_path, ref_name)
    if from_header:
        return from_header

    for name in names_to_try:
        if name in HCV_REFERENCE_METADATA:
            return HCV_REFERENCE_METADATA[name]

    return None


def read_codon_from_aligned_segment(read, ref_positions_1based):
    pairs = read.get_aligned_pairs(matches_only=False)
    ref_to_query = {}

    for query_pos, ref_pos in pairs:
        if ref_pos is None:
            continue
        ref_to_query[ref_pos + 1] = query_pos

    bases = []
    seq = read.query_sequence or ""

    for ref_pos in ref_positions_1based:
        qpos = ref_to_query.get(ref_pos)
        if qpos is None:
            return None
        if qpos >= len(seq):
            return None
        base = seq[qpos].upper()
        if base not in "ACGT":
            return None
        bases.append(base)

    return "".join(bases)


def call_hcv_ns5a_position_from_bam(
    bam_file,
    ref_name,
    codon_start_1based,
    reference_aa,
    resistance_aas,
    min_variant_frequency=HCV_MIN_VARIANT_FREQUENCY
):
    if pysam is None:
        raise ImportError("pysam is required for HCV minor-variant calling. Install with: pip install pysam")

    from collections import Counter

    codon_positions = [codon_start_1based, codon_start_1based + 1, codon_start_1based + 2]
    aa_counts = Counter()

    with pysam.AlignmentFile(bam_file, "rb") as bam:
        for read in bam.fetch(ref_name, codon_start_1based - 1, codon_start_1based + 2):
            if read.is_unmapped or read.is_secondary or read.is_supplementary:
                continue
            if read.is_duplicate or read.is_qcfail:
                continue

            codon = read_codon_from_aligned_segment(read, codon_positions)
            if not codon:
                continue

            aa = translate_codon(codon)
            if aa != "X":
                aa_counts[aa] += 1

    depth = sum(aa_counts.values())
    calls = []

    for aa, count in aa_counts.most_common():
        freq = count / depth if depth else 0
        if aa != reference_aa and aa in resistance_aas and freq >= min_variant_frequency:
            calls.append({
                "aa": aa,
                "count": count,
                "frequency": round(freq * 100, 1),
            })

    return {
        "depth": depth,
        "calls": calls,
        "aa_distribution": dict(aa_counts),
    }


def format_hcv_position_call(call_result):
    calls = call_result.get("calls", []) if call_result else []
    if not calls:
        return " "
    return " | ".join(f"{call['aa']} : {call['frequency']}%" for call in calls)


def determine_hcv_interpretation(position_results, low_coverage_positions):
    if low_coverage_positions:
        return "Ej bedömbar"

    frequencies = []
    for result in position_results.values():
        for call in result.get("calls", []) or []:
            try:
                frequencies.append(float(call.get("frequency", 0)) / 100.0)
            except (TypeError, ValueError):
                continue

    if not frequencies:
        return "Känslig"
    if max(frequencies) < HCV_MINOR_RAS_MAX_FREQUENCY:
        return "RAS minoritetsvariant"
    return "Resistent"


def hcv_candidate_ns5a_starts(meta):
    starts = []
    try:
        if meta and meta.get("ns5a_codon1_start"):
            starts.append(int(meta.get("ns5a_codon1_start")))
    except Exception:
        pass

    # Common coordinate systems in this project:
    # 6258 = complete-genome-style HCV references with 5'UTR included.
    # 5917 = CDS/polyprotein-style references starting at the polyprotein ATG.
    for candidate in [6258, 5917]:
        if candidate not in starts:
            starts.append(candidate)
    return starts


def call_hcv_panel_for_start(
    bam_file,
    ref_name,
    ref_path,
    subtype,
    ns5a_codon1_start
):
    position_results = {}
    low_coverage = []
    details = []

    subtype = str(subtype or "Unknown").upper().strip()
    subtype_panel = HCV_NS5A_PANEL_BY_SUBTYPE.get(subtype, {})
    if not subtype_panel:
        raise ValueError(f"Subtype {subtype} does not have a validated NS5A resistance panel")

    reference_sequence = read_named_fasta_sequence(ref_path, ref_name)
    if not reference_sequence:
        raise ValueError(f"Could not read selected HCV reference sequence '{ref_name}' from {ref_path}")

    for aa_pos, rule in subtype_panel.items():
        codon_start = int(ns5a_codon1_start) + ((aa_pos - 1) * 3)
        reference_codon = circular_codon(reference_sequence, codon_start)
        reference_aa = translate_codon(reference_codon)
        label = rule.get("label") or (f"{reference_aa}{aa_pos}" if reference_aa != "X" else f"AA{aa_pos}")

        result = call_hcv_ns5a_position_from_bam(
            bam_file=bam_file,
            ref_name=ref_name,
            codon_start_1based=codon_start,
            reference_aa=reference_aa,
            resistance_aas=rule.get("reportable_aas", set()),
            min_variant_frequency=HCV_MIN_VARIANT_FREQUENCY,
        )

        # Keep numeric keys internally so output columns remain stable even when
        # the subtype reference residue is A28/D58 rather than M28/H58.
        position_results[aa_pos] = result
        depth = int(result.get("depth", 0) or 0)
        if depth < HCV_MIN_COVERAGE:
            low_coverage.append(label)

        aa_dist = result.get("aa_distribution", {}) or {}
        dist_parts = []
        for aa, count in sorted(aa_dist.items(), key=lambda kv: (-kv[1], kv[0])):
            freq = round((count / depth) * 100, 1) if depth else 0
            dist_parts.append(f"{aa}:{count} ({freq}%)")

        details.append({
            "position": label,
            "aa_position": aa_pos,
            "reference_codon": reference_codon,
            "reference_aa": reference_aa,
            "reportable_aas": ",".join(sorted(rule.get("reportable_aas", set()))),
            "codon_start": codon_start,
            "codon_range": f"{codon_start}-{codon_start + 2}",
            "depth": depth,
            "aa_distribution": "; ".join(dist_parts),
            "reported_calls": format_hcv_position_call(result).strip(),
        })

    positions_passing = sum(1 for d in details if d["depth"] >= HCV_MIN_COVERAGE)
    total_depth = sum(d["depth"] for d in details)
    min_depth = min([d["depth"] for d in details], default=0)

    return {
        "ns5a_start": int(ns5a_codon1_start),
        "position_results": position_results,
        "low_coverage": low_coverage,
        "details": details,
        "positions_passing": positions_passing,
        "total_depth": total_depth,
        "min_depth": min_depth,
    }

def score_hcv_ns5a_start(reference_sequence, subtype, candidate_start):
    motif = HCV_NS5A_REFERENCE_MOTIF_BY_SUBTYPE.get(str(subtype or "").upper(), {})
    residues = {}
    score = 0
    for aa_pos, expected_aa in motif.items():
        codon_start = int(candidate_start) + ((aa_pos - 1) * 3)
        aa = translate_codon(circular_codon(reference_sequence, codon_start))
        residues[aa_pos] = aa
        if aa == expected_aa:
            score += 1

    # Y93 is essential for a valid NS5A resistance coordinate system.
    y93_ok = residues.get(93) == "Y"
    return score, y93_ok, residues


def locate_hcv_ns5a_start(ref_path, ref_name, subtype, metadata_start):
    reference_sequence = read_named_fasta_sequence(ref_path, ref_name)
    if not reference_sequence:
        raise ValueError(f"Could not read selected HCV reference sequence '{ref_name}' from {ref_path}")

    try:
        metadata_start = int(metadata_start)
    except (TypeError, ValueError):
        metadata_start = 0
    if metadata_start <= 0:
        raise ValueError("Missing or invalid NS5A start coordinate")

    motif = HCV_NS5A_REFERENCE_MOTIF_BY_SUBTYPE.get(str(subtype or "").upper(), {})
    if not motif:
        return metadata_start, {"score": 0, "residues": {}, "offset": 0, "method": "metadata-only"}

    candidates = []
    # Search around the supplied coordinate. This resolves accession-specific
    # offsets and frame errors while retaining the expected genomic neighbourhood.
    for candidate in range(max(1, metadata_start - 120), metadata_start + 121):
        score, y93_ok, residues = score_hcv_ns5a_start(reference_sequence, subtype, candidate)
        candidates.append((y93_ok, score, -abs(candidate - metadata_start), candidate, residues))

    candidates.sort(reverse=True)
    y93_ok, score, _, selected_start, residues = candidates[0]
    minimum_score = max(3, len(motif) - 1)
    if not y93_ok or score < minimum_score:
        observed = ", ".join(f"{p}:{aa}" for p, aa in sorted(residues.items()))
        raise ValueError(
            f"Could not validate NS5A frame for subtype {subtype} near {metadata_start}; "
            f"best start {selected_start} matched {score}/{len(motif)} residues ({observed})"
        )

    return selected_start, {
        "score": score,
        "residues": residues,
        "offset": selected_start - metadata_start,
        "method": "reference-motif-autoframe",
    }


def hcv_panel_labels_for_subtype(subtype):
    """Return the clinically relevant report labels for the selected subtype panel."""
    subtype = str(subtype or "Unknown").upper().strip()
    subtype_panel = HCV_NS5A_PANEL_BY_SUBTYPE.get(subtype, {})
    return [rule.get("label", f"AA{aa_pos}") for aa_pos, rule in subtype_panel.items()]


def choose_hcv_panel_result(bam_file, ref_name, ref_path, meta):
    try:
        metadata_start = int(meta.get("ns5a_codon1_start", 0) or 0)
    except Exception:
        metadata_start = 0

    if metadata_start <= 0:
        return None

    subtype = str(meta.get("subtype", "Unknown")).upper().strip()

    try:
        selected_start, validation = locate_hcv_ns5a_start(
            ref_path, ref_name, subtype, metadata_start
        )
        result = call_hcv_panel_for_start(
            bam_file, ref_name, ref_path, subtype, selected_start
        )
        result["metadata_ns5a_start"] = metadata_start
        result["coordinate_validation"] = validation
        for detail in result.get("details", []):
            detail["metadata_ns5a_start"] = metadata_start
            detail["selected_offset"] = validation.get("offset", 0)
            detail["coordinate_method"] = validation.get("method", "")
            detail["motif_score"] = validation.get("score", "")
        return result
    except Exception as validation_error:
        # If reference-motif auto-frame validation fails, still try the curated
        # metadata coordinate. This prevents validated subtype panels such as 1B
        # from being discarded only because the motif check is too strict for a
        # particular reference accession. The debug sheet records the fallback.
        try:
            validation = {
                "score": "",
                "residues": {},
                "offset": 0,
                "method": "metadata-coordinate-after-autoframe-failed",
                "autoframe_error": str(validation_error),
            }
            result = call_hcv_panel_for_start(
                bam_file, ref_name, ref_path, subtype, metadata_start
            )
            result["metadata_ns5a_start"] = metadata_start
            result["coordinate_validation"] = validation
            for detail in result.get("details", []):
                detail["metadata_ns5a_start"] = metadata_start
                detail["selected_offset"] = 0
                detail["coordinate_method"] = validation.get("method", "")
                detail["motif_score"] = ""
                detail["autoframe_error"] = str(validation_error)
            return result
        except Exception as panel_error:
            failed_positions = hcv_panel_labels_for_subtype(subtype)
            error_text = (
                f"Auto-frame failed: {validation_error}; "
                f"metadata-coordinate panel failed: {panel_error}"
            )
            return {
                "ns5a_start": metadata_start,
                "metadata_ns5a_start": metadata_start,
                "position_results": {},
                "low_coverage": failed_positions,
                "details": [{
                    "position": "panel",
                    "aa_position": "",
                    "reference_codon": "",
                    "reference_aa": "",
                    "reportable_aas": "",
                    "metadata_ns5a_start": metadata_start,
                    "selected_offset": "",
                    "coordinate_method": "failed",
                    "motif_score": "",
                    "codon_range": "",
                    "depth": "",
                    "aa_distribution": "",
                    "reported_calls": error_text,
                }],
                "positions_passing": 0,
                "total_depth": 0,
                "min_depth": 0,
                "error": error_text,
            }


def analyse_hcv_resistance(
    sample,
    input_fastq,
    bam_file,
    ref_name,
    ref_path,
    mapped_reads_1a,
    mapped_reads_1b,
    median_nonzero_coverage
):
    post_trim_reads, post_trim_mean_len = read_fastq_stats(input_fastq)
    meta = get_hcv_reference_metadata(ref_path, ref_name)

    base_row = {
        "hcv_quality": "✕",
        "hcv_interpretation": "Null",
        "hcv_M28": " ",
        "hcv_Q30": " ",
        "hcv_L31": " ",
        "hcv_H58": " ",
        "hcv_Y93": " ",
        "hcv_low_coverage": " ",
        "hcv_post_trim": f"{post_trim_reads} | {post_trim_mean_len}",
        "hcv_mapping_distribution": f"{mapped_reads_1a} | {mapped_reads_1b}",
        "hcv_median_nonzero_coverage": median_nonzero_coverage,
        "hcv_subtype": hcv_subtype_from_name(ref_name),
        "hcv_ns5a_start_used": "",
        "hcv_debug_details": [],
        "hcv_script_version": HCV_SCRIPT_VERSION,
    }

    if not ref_name or str(ref_name).upper() == "NA":
        base_row.update({
            "hcv_quality": "✕",
            "hcv_interpretation": "Null",
            "hcv_subtype": "Unknown",
            "hcv_low_coverage": "No valid HCV reference mapping found; resistance analysis not possible.",
            "hcv_ns5a_start_used": "not scanned",
            "hcv_debug_details": [],
        })
        return base_row

    if not meta:
        base_row["hcv_low_coverage"] = (
            "Missing HCV metadata: add subtype and ns5a_codon1_start "
            "to FASTA header or hcv_reference_metadata.tsv"
        )
        return base_row

    base_row["hcv_subtype"] = str(meta.get("subtype", base_row["hcv_subtype"])).upper().strip()

    if base_row["hcv_subtype"] not in HCV_NS5A_PANEL_BY_SUBTYPE:
        base_row.update({
            "hcv_quality": "Not Included",
            "hcv_interpretation": "Not Included",
            "hcv_M28": " ",
            "hcv_Q30": " ",
            "hcv_L31": " ",
            "hcv_H58": " ",
            "hcv_Y93": " ",
            "hcv_low_coverage": (
                f"Resistance scanning disabled: subtype {base_row['hcv_subtype']} "
                "does not have a validated NS5A resistance panel in this script. "
                "Validated panels: "
                f"{', '.join(sorted(HCV_NS5A_PANEL_BY_SUBTYPE.keys()))}."
            ),
            "hcv_ns5a_start_used": "not scanned",
            "hcv_debug_details": [],
        })
        return base_row

    try:
        selected = choose_hcv_panel_result(bam_file, ref_name, ref_path, meta)
        if not selected:
            base_row["hcv_low_coverage"] = "HCV resistance analysis failed: no coordinate candidates"
            return base_row

        position_results = selected.get("position_results", {})
        low_coverage = selected.get("low_coverage", [])
        interpretation = determine_hcv_interpretation(position_results, low_coverage)

        low_coverage_text = ", ".join(low_coverage) if low_coverage else " "
        if selected.get("error"):
            low_coverage_text = f"HCV resistance analysis failed: {selected.get('error')}"

        base_row.update({
            "hcv_quality": "✕" if low_coverage or selected.get("error") else "✓",
            "hcv_interpretation": interpretation,
            "hcv_M28": format_hcv_position_call(position_results.get(28)),
            "hcv_Q30": format_hcv_position_call(position_results.get(30)),
            "hcv_L31": format_hcv_position_call(position_results.get(31)),
            "hcv_H58": format_hcv_position_call(position_results.get(58)),
            "hcv_Y93": format_hcv_position_call(position_results.get(93)),
            "hcv_low_coverage": low_coverage_text,
            "hcv_ns5a_start_used": selected.get("ns5a_start", ""),
            "hcv_debug_details": selected.get("details", []),
        })

    except Exception as e:
        base_row["hcv_quality"] = "✕"
        base_row["hcv_interpretation"] = "Null"
        base_row["hcv_low_coverage"] = f"HCV resistance analysis failed: {e}"

    return base_row


def build_run_tag(run_date, run_number):
    return f"{run_date}_run{int(run_number):02d}"


def allocate_run_identity(output_folder):
    """Allocate one date/run identifier shared by every output from this launch."""
    ensure_dir(output_folder, must_be_writable=True)
    run_date = datetime.now().strftime("%Y-%m-%d")
    pattern = re.compile(rf"^Aurora_{re.escape(run_date)}_run(\d+)\.run$")
    highest = 0
    for name in os.listdir(output_folder):
        match = pattern.match(name)
        if match:
            highest = max(highest, int(match.group(1)))
    run_number = highest + 1
    run_tag = build_run_tag(run_date, run_number)
    marker_path = os.path.join(output_folder, f"Aurora_{run_tag}.run")
    with open(marker_path, "x", encoding="utf-8") as marker_file:
        marker_file.write(
            f"date={run_date}\nrun_number={run_number}\n"
            f"created={datetime.now().isoformat(timespec='seconds')}\n"
            f"script={os.path.abspath(__file__)}\n"
        )
    return run_date, run_number, run_tag, marker_path



HCV_HTML_SCRIPT_VERSION = "HCV_NS5A_HTML_REPORT_2026_07_20"
HCV_NS5A_REPRESENTATIVE_MEDICATIONS = (
    "Daclatasvir",
    "Ledipasvir",
    "Elbasvir",
    "Ombitasvir",
    "Velpatasvir",
    "Pibrentasvir",
)


def hcv_detected_calls(row):
    """Flatten reportable HCV NS5A calls already produced by the existing caller."""
    calls = []
    for detail in row.get("hcv_debug_details", []) or []:
        label = str(detail.get("position", "") or "")
        reference_aa = str(detail.get("reference_aa", "") or "")
        depth = int(detail.get("depth", 0) or 0) if str(detail.get("depth", "")).strip() else 0
        reported = str(detail.get("reported_calls", "") or "").strip()
        if not label or not reported or reported.lower().startswith(("auto-frame failed", "hcv resistance analysis failed")):
            continue
        for part in reported.split("|"):
            match = re.search(r"([A-Z*])\s*:\s*([0-9.]+)%", part.strip())
            if not match:
                continue
            alt_aa = match.group(1)
            frequency = float(match.group(2))
            if alt_aa == reference_aa:
                continue
            calls.append({
                "position": label,
                "mutation": f"{label}{alt_aa}",
                "reference_aa": reference_aa,
                "alt_aa": alt_aa,
                "frequency": frequency,
                "depth": depth,
                "count": int(round(depth * frequency / 100.0)) if depth else 0,
                "population": "Majoritetspopulation" if frequency >= 50.0 else "Minoritetspopulation",
            })
    return calls


def hcv_html_summary(row, calls):
    quality = str(row.get("hcv_quality", "") or "")
    interpretation = str(row.get("hcv_interpretation", "") or "")
    if quality == "Not Included":
        return "Subtypen ingår inte i den validerade NS5A-panelen i denna analysversion."
    if quality != "✓":
        return "Resultatet är inte fullständigt bedömbart på grund av otillräcklig paneltäckning eller analysfel."
    if not calls:
        return "Inga NS5A-resistensassocierade substitutioner påvisades över rapporteringströskeln."
    findings = ", ".join(f"{item['mutation']} ({item['frequency']:.1f} %)" for item in calls)
    return f"Påvisade HCV NS5A-RAS: {findings}. Samlad bedömning: {interpretation}."


def hcv_status_class(row, calls):
    quality = str(row.get("hcv_quality", "") or "")
    if quality not in ("✓",):
        return "status-unknown"
    if any(item["frequency"] >= 50.0 for item in calls):
        return "status-high"
    if calls:
        return "status-medium"
    return "status-good"


def hcv_panel_coverage_html(row):
    rows = []
    for detail in row.get("hcv_debug_details", []) or []:
        label = str(detail.get("position", "") or "")
        if not label or label == "panel":
            continue
        depth = int(detail.get("depth", 0) or 0) if str(detail.get("depth", "")).strip() else 0
        reported = str(detail.get("reported_calls", "") or "").strip() or "Inget panelrelevant fynd"
        qc = "Godkänd" if depth >= HCV_MIN_COVERAGE else "Otillräcklig"
        rows.append(
            "<tr>"
            f"<td>{html_escape(label)}</td>"
            f"<td>{html_escape(str(detail.get('reference_aa', '') or ''))}</td>"
            f"<td>{html_escape(str(detail.get('reference_codon', '') or ''))}</td>"
            f"<td>{depth}</td>"
            f"<td>{html_escape(reported)}</td>"
            f"<td>{qc}</td>"
            "</tr>"
        )
    return "".join(rows) or "<tr><td colspan='6' class='empty'>Paneldata saknas.</td></tr>"


def hcv_medication_rows_html(row, calls):
    """GLUE-inspired medication view without inventing drug-specific scores."""
    if str(row.get("hcv_quality", "") or "") != "✓":
        assessment = "Ej bedömbar på grund av otillräcklig paneltäckning"
        findings = "—"
        css = "drug-unknown"
    elif calls:
        assessment = "NS5A-RAS påvisad – läkemedelsspecifik effekt ej beräknad"
        findings = "; ".join(f"{x['mutation']} ({x['frequency']:.1f} %)" for x in calls)
        css = "drug-ras"
    else:
        assessment = "Ingen panel-RAS påvisad"
        findings = "—"
        css = "drug-none"
    return "".join(
        "<tr>"
        f"<td><strong>{html_escape(drug)}</strong></td>"
        f"<td><span class='drug-status {css}'>{html_escape(assessment)}</span></td>"
        f"<td>{html_escape(findings)}</td>"
        "</tr>"
        for drug in HCV_NS5A_REPRESENTATIVE_MEDICATIONS
    )


def write_hcv_html_report(row, report_path, run_tag):
    calls = hcv_detected_calls(row)
    logo_uri = image_file_to_data_uri(find_hbv_report_logo(row, report_path))
    logo_html = f"<img class='report-logo' src='{logo_uri}' alt='Logotyp'>" if logo_uri else ""
    summary = hcv_html_summary(row, calls)
    mutation_rows = []
    for item in calls:
        mutation_rows.append(
            "<tr>"
            f"<td><strong>{html_escape(item['mutation'])}</strong></td>"
            f"<td>{item['frequency']:.1f} %</td>"
            f"<td>{html_escape(item['population'])}</td>"
            f"<td>{item['count']} / {item['depth']}</td>"
            f"<td>{html_escape(str(row.get('hcv_subtype', '') or ''))}</td>"
            "</tr>"
        )
    if not mutation_rows:
        mutation_rows.append("<tr><td colspan='5' class='empty'>Inga rapporterbara NS5A-RAS påvisades.</td></tr>")

    html = f"""<!DOCTYPE html>
<html lang='sv'>
<head>
<meta charset='UTF-8'>
<meta name='viewport' content='width=device-width, initial-scale=1'>
<title>HCV-resistensrapport – {html_escape(str(row.get('sample','')))}</title>
<style>
:root {{ --blue:#dbeeff; --blue-dark:#185b8d; --line:#c9d3dc; --text:#1e2a32; --muted:#5d6b75; }}
* {{ box-sizing:border-box; }}
body {{ font-family:Arial,Helvetica,sans-serif; color:var(--text); margin:0; background:#eef2f5; }}
.report {{ max-width:1050px; margin:24px auto; background:#fff; padding:34px 42px; box-shadow:0 3px 18px rgba(0,0,0,.12); }}
h1 {{ margin:0 0 5px; font-size:28px; color:#123e5f; }}
.header {{ display:flex; justify-content:space-between; align-items:flex-start; gap:24px; margin-bottom:20px; }}
.title-block {{ flex:1; }} .report-logo {{ max-width:220px; max-height:105px; object-fit:contain; }}
.subtitle {{ color:var(--muted); }} .summary {{ background:#f4f8fb; border-left:6px solid var(--blue-dark); padding:12px 15px; margin:14px 0 18px; }}
h2 {{ font-size:18px; margin:28px 0 10px; border-bottom:2px solid var(--blue-dark); padding-bottom:6px; }}
table {{ width:100%; border-collapse:collapse; font-size:14px; }} th {{ background:var(--blue); text-align:left; }}
th,td {{ border:1px solid var(--line); padding:9px 10px; vertical-align:top; }} .meta td:first-child {{ width:28%; background:#f4f8fb; font-weight:bold; }}
.status {{ padding:14px 16px; border-radius:6px; font-weight:bold; margin:16px 0; }}
.status-good {{ background:#dff3e3; border-left:7px solid #3e9d54; }} .status-medium {{ background:#fff4c7; border-left:7px solid #d0a600; }}
.status-high {{ background:#ffe1d6; border-left:7px solid #d86032; }} .status-unknown {{ background:#e3e6e8; border-left:7px solid #6f777d; }}
.drug-status {{ font-weight:bold; }} .drug-none {{ color:#2c7b3e; }} .drug-ras {{ color:#a17900; }} .drug-unknown {{ color:#6f777d; }}
.empty {{ color:var(--muted); font-style:italic; }} .notice {{ background:#f5f5f5; border-left:5px solid #777; padding:12px 14px; margin-top:26px; font-size:12px; }}
.small {{ font-size:12px; color:var(--muted); word-break:break-all; }}
.report-footer {{
    margin-top: 24px;
    padding-top: 10px;
    border-top: 1px solid var(--line);
    color: var(--muted);
    font-size: 10px;
    line-height: 1.35;
    text-align: center;
}}
@media print {{ body {{ background:#fff; }} .report {{ box-shadow:none; margin:0; max-width:none; }} .report-logo {{ max-height:85px; }} }}
</style>
</head>
<body><main class='report'>
<div class='header'><div class='title-block'>
<h1>Hepatit C-virus – NS5A-resistensrapport</h1>
<div class='subtitle'>Läsnivåbaserad RAS-analys · Endast för forskningsbruk</div>
</div>{logo_html}</div>
<table class='meta'>
<tr><td>Prov-ID</td><td>{html_escape(str(row.get('sample','')))}</td></tr>
<tr><td>Subtyp/genotyp</td><td>{html_escape(str(row.get('hcv_subtype','')))}</td></tr>
<tr><td>Bästa referens</td><td>{html_escape(str(row.get('best_reference','')))}</td></tr>
<tr><td>Rapportdatum</td><td>{datetime.now().strftime('%Y-%m-%d %H:%M')}</td></tr>
<tr><td>Körning</td><td>{html_escape(run_tag)}</td></tr>
<tr><td>Rapporteringströskel</td><td>≥ {HCV_MIN_VARIANT_FREQUENCY * 100:.1f} %</td></tr>
<tr><td>Minsta kodontäckning</td><td>{HCV_MIN_COVERAGE} läsningar</td></tr>
<tr><td>Median nollskild täckning</td><td>{html_escape(str(row.get('hcv_median_nonzero_coverage','')))}×</td></tr>
<tr><td>Antal mappade läsningar</td><td>{html_escape(str(row.get('mapped_reads','')))}</td></tr>
</table>
<div class='summary'><strong>Sammanfattning:</strong> {html_escape(summary)}</div>
<div class='status {hcv_status_class(row, calls)}'>{html_escape(str(row.get('hcv_interpretation','')))}</div>

<h2>NS5A-läkemedelsöversikt</h2>
<table><thead><tr><th>Läkemedel</th><th>Panelbedömning</th><th>Påvisade RAS</th></tr></thead>
<tbody>{hcv_medication_rows_html(row, calls)}</tbody></table>
<p class='small'>Tabellen är en klassbaserad screeningvy. Den aktuella panelen beräknar inte GLUE-liknande läkemedelsspecifika resistenspoäng eller behandlingsrekommendationer.</p>

<h2>Påvisade NS5A-resistensassocierade substitutioner</h2>
<table><thead><tr><th>Mutation</th><th>Andel</th><th>Population</th><th>Mutantläsningar / djup</th><th>Subtyp</th></tr></thead>
<tbody>{''.join(mutation_rows)}</tbody></table>

<h2>Paneltäckning och positionskontroll</h2>
<table><thead><tr><th>Position</th><th>Referens-AA</th><th>Referenskodon</th><th>Kodontäckning</th><th>Rapporterat resultat</th><th>QC</th></tr></thead>
<tbody>{hcv_panel_coverage_html(row)}</tbody></table>

<h2>Teknisk kvalitet</h2>
<table class='meta'>
<tr><td>Kvalitetsstatus</td><td>{html_escape(str(row.get('hcv_quality','')))}</td></tr>
<tr><td>Övergripande QC</td><td>{html_escape(str(row.get('qc_verdict','')))}</td></tr>
<tr><td>QC-kommentar</td><td>{html_escape(str(row.get('qc_reason','')))}</td></tr>
<tr><td>Positioner med otillräcklig täckning</td><td>{html_escape(str(row.get('hcv_low_coverage','')))}</td></tr>
<tr><td>NS5A-koordinat</td><td>start {html_escape(str(row.get('hcv_ns5a_start_used','')))}</td></tr>
<tr><td>Genom-/måltäckning</td><td>{html_escape(str(row.get('target_or_genome_coverage_percent','')))} %</td></tr>
<tr><td>Medeldjup</td><td>{html_escape(str(row.get('mean_depth','')))}×</td></tr>
<tr><td>Mediandjup</td><td>{html_escape(str(row.get('median_depth','')))}×</td></tr>
<tr><td>Positioner ≥10×</td><td>{html_escape(str(row.get('positions_ge_10x_percent','')))} %</td></tr>
<tr><td>Positioner ≥30×</td><td>{html_escape(str(row.get('positions_ge_30x_percent','')))} %</td></tr>
<tr><td>Konsensuslängd</td><td>{html_escape(str(row.get('consensus_length','')))} nt</td></tr>
<tr><td>N-baser</td><td>{html_escape(str(row.get('n_count','')))} ({html_escape(str(row.get('n_percent','')))} %)</td></tr>
<tr><td>Medaka-modell</td><td>{html_escape(str(row.get('medaka_model','')))}</td></tr>
<tr><td>FASTQ-fil</td><td class='small'>{html_escape(str(row.get('input_fastq','')))}</td></tr>
<tr><td>Konsensusfil</td><td class='small'>{html_escape(str(row.get('consensus_file','')))}</td></tr>
<tr><td>BAM-fil</td><td class='small'>{html_escape(str(row.get('bam_file','')))}</td></tr>
<tr><td>Referensfil</td><td class='small'>{html_escape(str(row.get('reference_file','')))}</td></tr>
<tr><td>Scriptversion</td><td>{html_escape(str(row.get('hcv_script_version','')))} / {HCV_HTML_SCRIPT_VERSION}</td></tr>
</table>
<div class='notice'><strong>Viktig begränsning:</strong> Rapporten visar endast de subtypspecifika NS5A-positioner och substitutioner som ingår i Aurora-panelen. Den ersätter inte en fullständig HCV-GLUE-analys, läkemedelsspecifik resistensalgoritm, klinisk behandlingsbedömning eller expertgranskning.</div>
<footer class='report-footer'>
Vid frågor, tekniska problem eller oväntade analysresultat, kontakta skriptets författare Joakim Bedner Stenbäck (joakim.stenback@vgregion.se).
</footer>

</main></body></html>"""
    ensure_dir(os.path.dirname(report_path), must_be_writable=True)
    with open(report_path, "w", encoding="utf-8") as handle:
        handle.write(html)


def generate_hcv_html_reports(hcv_rows, output_folder, run_tag):
    html_dir = os.path.join(run_samples_folder(output_folder, run_tag), "HCV_HTML_reports")
    ensure_dir(html_dir, must_be_writable=True)
    index_rows = []
    for row in hcv_rows:
        sample = str(row.get("sample", "") or "prov")
        filename = f"{safe_report_filename(sample)}_{run_tag}_HCV.html"
        report_path = os.path.join(html_dir, filename)
        write_hcv_html_report(row, report_path, run_tag)
        row["hcv_html_report"] = report_path
        index_rows.append(
            "<tr>"
            f"<td><a href='{html_escape(filename)}'>{html_escape(sample)}</a></td>"
            f"<td>{html_escape(str(row.get('hcv_subtype','')))}</td>"
            f"<td>{html_escape(str(row.get('hcv_interpretation','')))}</td>"
            f"<td>{html_escape(str(row.get('hcv_quality','')))}</td>"
            "</tr>"
        )
    index_path = os.path.join(html_dir, f"index_{run_tag}.html")
    index_html = f"""<!DOCTYPE html><html lang='sv'><head><meta charset='UTF-8'><title>HCV-rapporter {html_escape(run_tag)}</title>
<style>body{{font-family:Arial;margin:35px;color:#1e2a32}}table{{border-collapse:collapse;width:100%}}th{{background:#dbeeff}}th,td{{border:1px solid #c9d3dc;padding:9px;text-align:left}}a{{color:#185b8d}}</style></head>
<body><h1>HCV NS5A-resistensrapporter</h1><p>Körning: {html_escape(run_tag)}</p>
<table><thead><tr><th>Prov</th><th>Subtyp</th><th>Bedömning</th><th>Kvalitet</th></tr></thead><tbody>{''.join(index_rows)}</tbody></table></body></html>"""
    with open(index_path, "w", encoding="utf-8") as handle:
        handle.write(index_html)
    print(f"HCV HTML reports saved to {html_dir}")
    return html_dir

def log_hcv_resistance_results(hcv_rows, output_folder, run_tag):
    try:
        ensure_dir(output_folder, must_be_writable=True)
        with open(os.path.join(output_folder, f"HCV_VERSION_{run_tag}.txt"), "w", encoding="utf-8") as vf:
            vf.write(f"{HCV_SCRIPT_VERSION}\nscript={os.path.abspath(__file__)}\n")
        excel_path = os.path.join(output_folder, f"hcv_resistance_results_{run_tag}.xlsx")
        generate_hcv_html_reports(hcv_rows, output_folder, run_tag)

        wb = Workbook()
        ws = wb.active
        ws.title = "HCV Resistance"

        ws.append(["Hepatit C virus antiviralresistensbedömning"])
        ws.append([f"Script version: {HCV_SCRIPT_VERSION}"])
        ws.append([f"Run: {run_tag}"])
        ws.append(["Resultat"])
        ws.append([])

        headers = [
            "Prov-ID",
            "Kvalitet",
            "Bedömning",
            "M28",
            "Q30",
            "L31",
            "H58",
            "Y93",
            "Mutationer < Täckning",
            "Post-trimning",
            "Readmappningsdistribution",
            "Median nollskild täckning",
            "Subtype/Genotype Reference",
            "Best Reference",
            "BAM File",
            "NS5A start used",
            "Script version",
        ]

        ws.append(headers)

        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        for row in hcv_rows:
            ws.append([
                row.get("sample", ""),
                row.get("hcv_quality", ""),
                row.get("hcv_interpretation", ""),
                row.get("hcv_M28", " "),
                row.get("hcv_Q30", " "),
                row.get("hcv_L31", " "),
                row.get("hcv_H58", " "),
                row.get("hcv_Y93", " "),
                row.get("hcv_low_coverage", " "),
                row.get("hcv_post_trim", ""),
                row.get("hcv_mapping_distribution", ""),
                row.get("hcv_median_nonzero_coverage", ""),
                row.get("hcv_subtype", ""),
                row.get("best_reference", ""),
                row.get("bam_file", ""),
                row.get("hcv_ns5a_start_used", ""),
                row.get("hcv_script_version", ""),
            ])

            quality_cell = ws.cell(row=ws.max_row, column=2)
            interp_cell = ws.cell(row=ws.max_row, column=3)

            if quality_cell.value == "Not Included":
                quality_cell.fill = yellow_fill
            elif quality_cell.value == "✓":
                quality_cell.fill = green_fill
            else:
                quality_cell.fill = red_fill

            if interp_cell.value == "Not Included":
                interp_cell.fill = yellow_fill
            elif interp_cell.value == "Resistent":
                interp_cell.fill = red_fill
            elif interp_cell.value == "RAS minoritetsvariant":
                interp_cell.fill = yellow_fill
            elif interp_cell.value == "Känslig":
                interp_cell.fill = green_fill
            else:
                interp_cell.fill = yellow_fill

        debug_ws = wb.create_sheet("HCV Debug")
        debug_ws.append([
            "Sample", "Best Reference", "Subtype", "NS5A start used",
            "Position", "AA position", "Reference codon", "Reference AA", "Reportable AAs",
            "Metadata NS5A start", "Selected offset", "Coordinate method", "Motif score",
            "Codon range", "Codon depth", "AA distribution", "Reported >=15% calls",
            "Min depth cutoff", "Report threshold"
        ])
        for row in hcv_rows:
            for detail in row.get("hcv_debug_details", []) or []:
                debug_ws.append([
                    row.get("sample", ""),
                    row.get("best_reference", ""),
                    row.get("hcv_subtype", ""),
                    row.get("hcv_ns5a_start_used", ""),
                    detail.get("position", ""),
                    detail.get("aa_position", ""),
                    detail.get("reference_codon", ""),
                    detail.get("reference_aa", ""),
                    detail.get("reportable_aas", ""),
                    detail.get("metadata_ns5a_start", ""),
                    detail.get("selected_offset", ""),
                    detail.get("coordinate_method", ""),
                    detail.get("motif_score", ""),
                    detail.get("codon_range", ""),
                    detail.get("depth", ""),
                    detail.get("aa_distribution", ""),
                    detail.get("reported_calls", ""),
                    HCV_MIN_COVERAGE,
                    f">={round(HCV_MIN_VARIANT_FREQUENCY * 100, 1)}%",
                ])

        autosize_worksheet(ws)
        autosize_worksheet(debug_ws)
        wb.save(excel_path)
        mirror_path = os.path.join(output_folder, f"hcv_resistance_results_PRODUCTION_DEBUG_COPY_{run_tag}.xlsx")
        wb.save(mirror_path)
        with open(os.path.join(output_folder, f"HCV_PRODUCTION_SENTINEL_{run_tag}.txt"), "w", encoding="utf-8") as sf:
            sf.write(f"{HCV_SCRIPT_VERSION}\nscript={os.path.abspath(__file__)}\nexcel={excel_path}\nmirror={mirror_path}\n")
        print(f"HCV resistance Excel saved to {excel_path}")
        print(f"HCV debug mirror Excel saved to {mirror_path}")

    except Exception as e:
        print(f"Error logging HCV resistance results: {e}")


# HBV antiviral resistance and surface immune-escape analysis
# Research-use-only screening panel. Results require expert review and should not
# be used alone for treatment or infection-control decisions.
HBV_MUTATION_SCRIPT_VERSION = "Aurora3.1.0.HBV_COMBINATION_INTERPRETATION_2026_07_24"
HBV_MIN_COVERAGE = 100
HBV_MIN_VARIANT_FREQUENCY = 0.15

# Conservative, widely reported reverse-transcriptase substitutions.
# Tenofovir resistance is deliberately not assigned from a single substitution,
# because current evidence does not support a universal single-mutation rule.
HBV_RT_PANEL = {
    # Local priority markers: prominently flagged in the report rather than
    # downgraded to a generic compensatory warning.
    80:  {"label": "rtL80",  "aas": {"I", "V"}, "category": "priority resistance-associated marker", "flag": "PRIORITY", "drugs": "lamivudine/telbivudine-associated pattern"},
    173: {"label": "rtV173", "aas": {"L"}, "category": "priority resistance-associated marker", "flag": "PRIORITY", "drugs": "lamivudine-associated pattern"},

    # Well-established primary or key resistance-pattern substitutions.
    180: {"label": "rtL180", "aas": {"M"}, "category": "key resistance-pattern substitution", "flag": "ESTABLISHED", "drugs": "lamivudine; entecavir resistance background"},
    181: {"label": "rtA181", "aas": {"T", "V"}, "category": "primary resistance-associated", "flag": "ESTABLISHED", "drugs": "adefovir; reduced-susceptibility patterns"},
    204: {"label": "rtM204", "aas": {"I", "S", "V"}, "category": "primary resistance-associated", "flag": "ESTABLISHED", "drugs": "lamivudine/telbivudine"},
    236: {"label": "rtN236", "aas": {"T"}, "category": "primary resistance-associated", "flag": "ESTABLISHED", "drugs": "adefovir"},

    # These substitutions are biologically relevant but should not be interpreted
    # as standalone entecavir resistance without the appropriate background.
    169: {"label": "rtI169", "aas": {"T"}, "category": "context-dependent entecavir-associated", "flag": "CAUTION", "drugs": "interpret with rtL180M + rtM204V/I"},
    184: {"label": "rtT184", "aas": set("ACFGILMS"), "category": "context-dependent entecavir-associated", "flag": "CAUTION", "drugs": "interpret with rtL180M + rtM204V/I"},
    202: {"label": "rtS202", "aas": {"C", "G", "I"}, "category": "context-dependent entecavir-associated", "flag": "CAUTION", "drugs": "interpret with rtL180M + rtM204V/I"},
    250: {"label": "rtM250", "aas": {"I", "L", "V"}, "category": "context-dependent entecavir-associated", "flag": "CAUTION", "drugs": "interpret with rtL180M + rtM204V/I"},
}

# Curated HBsAg major-hydrophilic-region substitutions repeatedly associated
# with vaccine, antibody, or diagnostic escape. Association is not deterministic.
# Exact substitution-to-medication associations used for per-sample reporting.
# These are screening associations, not treatment recommendations. Entecavir
# substitutions often require a lamivudine-resistance background, and reported
# tenofovir associations are shown with an explicit caution label.
HBV_MEDICATIONS = ("Lamivudine", "Adefovir", "Entecavir", "Tenofovir")
HBV_MUTATION_MEDICATION_MAP = {
    "L80I": {"Lamivudine": "associated", "Entecavir": "associated/background"},
    "L80V": {"Lamivudine": "associated", "Entecavir": "associated/background"},
    "I169T": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "V173L": {"Lamivudine": "associated", "Entecavir": "associated/background"},
    "L180M": {"Lamivudine": "established pattern", "Entecavir": "resistance background"},
    "A181T": {"Lamivudine": "associated", "Adefovir": "associated"},
    "A181V": {"Lamivudine": "associated", "Adefovir": "associated"},
    "T184A": {"Entecavir": "context-dependent"},
    "T184G": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "T184I": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "T184C": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "T184M": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "T184S": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "T184L": {"Lamivudine": "associated", "Tenofovir": "candidate association - caution"},
    "S202G": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "S202I": {"Lamivudine": "associated", "Entecavir": "context-dependent"},
    "M204I": {"Lamivudine": "established", "Entecavir": "resistance background"},
    "M204V": {"Lamivudine": "established", "Entecavir": "resistance background"},
    "N236T": {"Adefovir": "established", "Tenofovir": "reported association - caution"},
    "M250V": {"Entecavir": "context-dependent"},
    "M250I": {"Entecavir": "context-dependent"},
    "M250L": {"Entecavir": "context-dependent"},
}


def hbv_medication_summary(details):
    by_medication = {drug: [] for drug in HBV_MEDICATIONS}
    unmapped = []

    for detail in details or []:
        if detail.get("gene") != "RT":
            continue
        ref_aa = hbv_panel_reference_aa(detail)
        observed_ref_aa = str(detail.get("reference_aa", ""))
        aa_pos = detail.get("aa_position", "")
        label = detail.get("position", "")
        for call in detail.get("calls", []) or []:
            alt_aa = str(call.get("aa", ""))
            if not alt_aa or alt_aa in {ref_aa, observed_ref_aa}:
                continue
            mutation = f"{ref_aa}{aa_pos}{alt_aa}"
            frequency = call.get("frequency", 0)
            display = f"rt{mutation} ({frequency}%)"
            mappings = HBV_MUTATION_MEDICATION_MAP.get(mutation, {})
            if not mappings:
                unmapped.append(f"{label}{alt_aa} ({frequency}%)")
                continue
            for drug, evidence in mappings.items():
                by_medication.setdefault(drug, []).append(f"{display} [{evidence}]")

    result = {
        f"hbv_drug_{drug.lower()}": "; ".join(by_medication.get(drug, [])) or " "
        for drug in HBV_MEDICATIONS
    }
    result["hbv_drug_unmapped"] = "; ".join(unmapped) or " "
    return result


# HBsAg immune-, vaccine-, HBIG- and diagnostic-escape screening panel.
#
# The core positions overlap the ten-position clinically relevant panel used in
# large HBV mutation surveys (120, 126, 129, 130, 133, 141-145). Additional
# substitutions at 123, 124, 127, 128, 131 and 134 are retained because they
# have been reported in escape/detection-failure literature. These calls are
# research-use screening findings and are not standalone proof of phenotypic
# escape. The caller still reports only nonsynonymous read-level populations at
# or above HBV_MIN_VARIANT_FREQUENCY.
HBV_S_ESCAPE_PANEL = {
    120: {"label": "sP120", "aas": {"L", "Q", "S", "T"}, "category": "immune/vaccine/diagnostic escape-associated"},
    123: {"label": "sT123", "aas": {"A", "N"}, "category": "immune/vaccine/diagnostic escape-associated"},
    124: {"label": "sC124", "aas": {"R", "Y"}, "category": "immune/vaccine/diagnostic escape-associated"},
    126: {"label": "sT126", "aas": {"A", "I", "N", "S"}, "category": "immune/vaccine/diagnostic escape-associated"},
    127: {"label": "sP127", "aas": {"L", "S"}, "category": "immune/vaccine/diagnostic escape-associated"},
    128: {"label": "sI128", "aas": {"V"}, "category": "secondary escape-associated marker"},
    129: {"label": "sQ129", "aas": {"H", "N", "R"}, "category": "immune/vaccine/diagnostic escape-associated"},
    130: {"label": "sG130", "aas": {"N", "R"}, "category": "immune/vaccine/diagnostic escape-associated"},
    131: {"label": "sN131", "aas": {"I"}, "category": "secondary escape-associated marker"},
    133: {"label": "sM133", "aas": {"I", "L", "T"}, "category": "immune/vaccine/diagnostic escape-associated"},
    134: {"label": "sF134", "aas": {"H", "L", "S", "Y"}, "category": "immune/vaccine/diagnostic escape-associated"},
    141: {"label": "sK141", "aas": {"E", "I"}, "category": "immune/vaccine/diagnostic escape-associated"},
    142: {"label": "sP142", "aas": {"L", "S"}, "category": "immune/vaccine/diagnostic escape-associated"},
    143: {"label": "sS143", "aas": {"L", "T"}, "category": "higher-impact vaccine/diagnostic escape-associated", "flag": "HIGH_IMPACT"},
    144: {"label": "sD144", "aas": {"A", "E", "G"}, "category": "higher-impact vaccine/HBIG/diagnostic escape-associated", "flag": "HIGH_IMPACT"},
    145: {"label": "sG145", "aas": {"A", "E", "K", "R"}, "category": "higher-impact vaccine/HBIG/diagnostic escape-associated", "flag": "HIGH_IMPACT"},
}


def parse_hbv_metadata_tsv(tsv_path):
    metadata = {}
    if not tsv_path or not os.path.exists(tsv_path):
        return metadata
    try:
        with open(tsv_path, "r", encoding="utf-8", errors="replace") as handle:
            header = None
            for raw in handle:
                line = raw.strip()
                if not line or line.startswith("#"):
                    continue
                parts = line.split("\t")
                if header is None and any(x.lower() in {"ref_name", "reference", "rt_codon1_start", "s_codon1_start"} for x in parts):
                    header = [x.strip().lower() for x in parts]
                    continue
                if header:
                    row = dict(zip(header, parts))
                    name = row.get("ref_name") or row.get("reference") or row.get("name")
                    rt_start = row.get("rt_codon1_start") or row.get("rt_start")
                    s_start = row.get("s_codon1_start") or row.get("s_start")
                elif len(parts) >= 3:
                    name, rt_start, s_start = parts[:3]
                else:
                    continue
                try:
                    metadata[name] = {"rt_codon1_start": int(rt_start), "s_codon1_start": int(s_start)}
                except (TypeError, ValueError):
                    continue
    except Exception as exc:
        print(f"Error reading HBV metadata TSV {tsv_path}: {exc}")
    return metadata


def parse_hbv_metadata_from_fasta_header(fasta_path, ref_name):
    if not fasta_path or not ref_name:
        return None
    try:
        with open(fasta_path, "r", encoding="utf-8", errors="replace") as handle:
            for line in handle:
                if not line.startswith(">"):
                    continue
                header = line[1:].strip()
                if header.split()[0] != str(ref_name).split()[0]:
                    continue
                rt = re.search(r"(?:rt_codon1_start|rt_start)=(\d+)", header, re.I)
                s = re.search(r"(?:s_codon1_start|s_start)=(\d+)", header, re.I)
                if rt and s:
                    return {"rt_codon1_start": int(rt.group(1)), "s_codon1_start": int(s.group(1))}
    except Exception as exc:
        print(f"Error parsing HBV FASTA metadata: {exc}")
    return None


def get_hbv_reference_metadata(ref_path, ref_name):
    names = [str(ref_name or ""), str(ref_name or "").split()[0], str(ref_name or "").split('.')[0]]
    for candidate in [
        os.path.join(os.path.dirname(ref_path), "hbv_reference_metadata.tsv") if ref_path else None,
        os.path.join(app_base_dir(), "hbv_reference_metadata.tsv"),
    ]:
        data = parse_hbv_metadata_tsv(candidate)
        for name in names:
            if name in data:
                return data[name]
    return parse_hbv_metadata_from_fasta_header(ref_path, ref_name)


def call_hbv_aa_position_from_bam(bam_file, ref_name, codon_start_1based, reference_aa, reportable_aas):
    if pysam is None:
        raise ImportError("pysam is required for HBV mutation calling. Install with: pip install pysam")
    from collections import Counter
    counts = Counter()
    positions = [codon_start_1based, codon_start_1based + 1, codon_start_1based + 2]
    with pysam.AlignmentFile(bam_file, "rb") as bam:
        for read in bam.fetch(ref_name, codon_start_1based - 1, codon_start_1based + 2):
            if read.is_unmapped or read.is_secondary or read.is_supplementary or read.is_duplicate or read.is_qcfail:
                continue
            codon = read_codon_from_aligned_segment(read, positions)
            if not codon:
                continue
            aa = translate_codon(codon)
            if aa != "X":
                counts[aa] += 1
    depth = sum(counts.values())
    calls = []
    for aa, count in counts.most_common():
        frequency = count / depth if depth else 0.0
        if aa != reference_aa and aa in reportable_aas and frequency >= HBV_MIN_VARIANT_FREQUENCY:
            calls.append({"aa": aa, "count": count, "frequency": round(frequency * 100, 1)})
    return {"depth": depth, "calls": calls, "aa_distribution": dict(counts)}


def format_hbv_position_call(result):
    calls = result.get("calls", []) if result else []
    if not calls:
        return " "
    return " | ".join(f"{call['aa']} : {call['frequency']}%" for call in calls)


def determine_hbv_interpretation(position_results, low_coverage_positions):
    """Return an HCV-style high-level result while retaining separate flags."""
    if low_coverage_positions:
        return "Null"

    frequencies = []
    has_rt = False
    has_escape = False
    for key, result in position_results.items():
        gene = key[0]
        for call in result.get("calls", []) or []:
            try:
                frequencies.append(float(call.get("frequency", 0)) / 100.0)
            except (TypeError, ValueError):
                continue
            if gene == "RT":
                has_rt = True
            elif gene == "S":
                has_escape = True

    if not frequencies:
        return "Inga panelmutationer påvisade"

    population = "minoritetspopulation" if max(frequencies) < 0.50 else "majoritetspopulation"
    if has_rt and has_escape:
        return f"RT-resistens och ytantigen-escape associerad {population}"
    if has_rt:
        return f"RT-resistensassocierad {population}"
    return f"Ytantigen-escape associerad {population}"


def analyse_hbv_mutations(sample, bam_file, ref_name, ref_path):
    base = {
        "hbv_quality": "✕",
        "hbv_interpretation": "Ej bedömbar",
        "hbv_low_coverage_positions": " ",
        "hbv_rt_priority_flags": " ",
        "hbv_rt_caution_flags": " ",
        "hbv_drug_lamivudine": " ",
        "hbv_drug_adefovir": " ",
        "hbv_drug_entecavir": " ",
        "hbv_drug_tenofovir": " ",
        "hbv_drug_unmapped": " ",
        "hbv_mutation_details": [],
        "hbv_mutation_script_version": HBV_MUTATION_SCRIPT_VERSION,
        "hbv_rt_codon1_start": "",
        "hbv_s_codon1_start": "",
        "hbv_coordinate_method": "",
    }
    for pos in HBV_RT_PANEL:
        base[f"hbv_rt_{pos}"] = " "
    for pos in HBV_S_ESCAPE_PANEL:
        base[f"hbv_s_{pos}"] = " "

    metadata = get_hbv_reference_metadata(ref_path, ref_name)
    if not metadata:
        base["hbv_low_coverage_positions"] = (
            "HBV-koordinatmetadata saknas: lägg till rt_codon1_start och "
            "s_codon1_start i FASTA-huvudet eller hbv_reference_metadata.tsv"
        )
        return base

    base["hbv_rt_codon1_start"] = metadata.get("rt_codon1_start", "")
    base["hbv_s_codon1_start"] = metadata.get("s_codon1_start", "")
    base["hbv_coordinate_method"] = "metadata"

    ref_seq = read_named_fasta_sequence(ref_path, ref_name)
    if not ref_seq:
        base["hbv_low_coverage_positions"] = "Den valda HBV-referenssekvensen kunde inte läsas"
        return base

    low_cov = []
    priority_flags = []
    caution_flags = []
    position_results = {}
    details = []

    try:
        for gene, start_key, panel in [
            ("RT", "rt_codon1_start", HBV_RT_PANEL),
            ("S", "s_codon1_start", HBV_S_ESCAPE_PANEL),
        ]:
            gene_start = int(metadata[start_key])
            for aa_pos, rule in panel.items():
                codon_start = gene_start + ((aa_pos - 1) * 3)
                ref_codon = circular_codon(ref_seq, codon_start)
                ref_aa = translate_codon(ref_codon)
                result = call_hbv_aa_position_from_bam(
                    bam_file, ref_name, codon_start, ref_aa, rule["aas"]
                )
                key = (gene, aa_pos)
                position_results[key] = result
                label = rule["label"]
                depth = int(result.get("depth", 0) or 0)
                if depth < HBV_MIN_COVERAGE:
                    low_cov.append(label)

                formatted = format_hbv_position_call(result)
                base[f"hbv_{gene.lower()}_{aa_pos}"] = formatted
                flag = rule.get("flag", "ESCAPE" if gene == "S" else "")
                if formatted.strip():
                    labelled = f"[{flag}] {label}: {formatted.strip()}" if flag else f"{label}: {formatted.strip()}"
                    if flag == "PRIORITY":
                        priority_flags.append(labelled)
                    elif flag == "CAUTION":
                        caution_flags.append(labelled)

                aa_dist = result.get("aa_distribution", {}) or {}
                dist_parts = []
                for aa, count in sorted(aa_dist.items(), key=lambda kv: (-kv[1], kv[0])):
                    freq = round((count / depth) * 100, 1) if depth else 0
                    dist_parts.append(f"{aa}:{count} ({freq}%)")

                details.append({
                    "gene": gene,
                    "position": label,
                    "aa_position": aa_pos,
                    "reference_codon": ref_codon,
                    "reference_aa": ref_aa,
                    "panel_reference_aa": re.sub(r"^(?:rt|s)([A-Z*])\d+$", r"\1", label),
                    "reportable_aas": ",".join(sorted(rule.get("aas", set()))),
                    "codon_start": codon_start,
                    "codon_range": f"{codon_start}-{codon_start + 2}",
                    "depth": depth,
                    "aa_distribution": "; ".join(dist_parts),
                    "reported_calls": formatted.strip(),
                    "calls": list(result.get("calls", []) or []),
                    "flag": flag,
                    "category": rule.get("category", "immune/vaccine/diagnostic escape-associated"),
                    "drugs": rule.get("drugs", ""),
                })

        medication_summary = hbv_medication_summary(details)
        base.update(medication_summary)
        base.update({
            "hbv_quality": "✕" if low_cov else "✓",
            "hbv_interpretation": determine_hbv_interpretation(position_results, low_cov),
            "hbv_low_coverage_positions": ", ".join(low_cov) if low_cov else " ",
            "hbv_rt_priority_flags": "; ".join(priority_flags) if priority_flags else " ",
            "hbv_rt_caution_flags": "; ".join(caution_flags) if caution_flags else " ",
            "hbv_mutation_details": details,
        })
    except Exception as exc:
        base["hbv_quality"] = "✕"
        base["hbv_interpretation"] = "Ej bedömbar"
        base["hbv_low_coverage_positions"] = f"HBV-mutationsanalysen misslyckades: {exc}"
    return base


def hbv_swedish_evidence_label(evidence):
    labels = {
        "established": "etablerad",
        "associated": "associerad",
        "associated/background": "associerad/bakgrund",
        "resistance background": "resistensbakgrund",
        "context-dependent": "kontextberoende",
        "candidate association - caution": "kandidatassociation – försiktighet",
        "reported association - caution": "rapporterad association – försiktighet",
    }
    return labels.get(str(evidence or ""), str(evidence or ""))




def hbv_translate_free_text(value):
    text = str(value or "")
    replacements = {
        "immune/vaccine/diagnostic escape-associated": "immun-, vaccin- eller diagnostikescape-associerad",
        "lamivudine": "lamivudin",
        "adefovir": "adefovir",
        "entecavir": "entekavir",
        "tenofovir": "tenofovir",
        "compensatory/background": "kompensatorisk/bakgrund",
        "primary resistance": "primär resistens",
        "resistance background": "resistensbakgrund",
        "context-dependent": "kontextberoende",
        "associated": "associerad",
    }
    for old, new in replacements.items():
        text = re.sub(re.escape(old), new, text, flags=re.I)
    return text


def hbv_population_label(frequency):
    try:
        return "Minoritetspopulation" if float(frequency) < 50.0 else "Majoritetspopulation"
    except (TypeError, ValueError):
        return "Okänd population"


def hbv_panel_reference_aa(detail):
    """Return the canonical wild-type amino acid encoded in labels such as rtV173."""
    label = str(detail.get("position", "") or "")
    match = re.match(r"^(?:rt|s)([A-Z*])\d+$", label)
    return match.group(1) if match else str(detail.get("reference_aa", "") or "")


def hbv_detected_calls(row):
    calls = []
    for detail in row.get("hbv_mutation_details", []) or []:
        canonical_ref_aa = hbv_panel_reference_aa(detail)
        observed_ref_aa = str(detail.get("reference_aa", "") or "")
        for call in detail.get("calls", []) or []:
            alt_aa = str(call.get("aa", "") or "")
            # Never report synonymous codon changes as amino-acid mutations.
            if not alt_aa or alt_aa in {canonical_ref_aa, observed_ref_aa}:
                continue
            mutation = f"{detail.get('position', '')}{alt_aa}"
            medication_map = {}
            if detail.get("gene") == "RT":
                key = f"{canonical_ref_aa}{detail.get('aa_position', '')}{alt_aa}"
                medication_map = HBV_MUTATION_MEDICATION_MAP.get(key, {})
            calls.append({
                "gene": detail.get("gene", ""),
                "mutation": mutation,
                "frequency": call.get("frequency", 0),
                "count": call.get("count", 0),
                "depth": detail.get("depth", 0),
                "population": hbv_population_label(call.get("frequency", 0)),
                "flag": detail.get("flag", ""),
                "category": detail.get("category", ""),
                "medications": medication_map,
            })
    return calls


def hbv_swedish_interpretation(value):
    translations = {
        "No panel mutation detected": "Inga panelmutationer påvisade",
        "Null": "Ej bedömbar",
        "RT resistance-associated minor population": "RT-resistensassocierad minoritetspopulation",
        "RT resistance-associated majority population": "RT-resistensassocierad majoritetspopulation",
        "Surface escape-associated minor population": "Ytantigen-escape associerad minoritetspopulation",
        "Surface escape-associated majority population": "Ytantigen-escape associerad majoritetspopulation",
        "RT resistance + surface escape-associated minor population": "RT-resistens och ytantigen-escape associerad minoritetspopulation",
        "RT resistance + surface escape-associated majority population": "RT-resistens och ytantigen-escape associerad majoritetspopulation",
    }
    return translations.get(str(value or ""), str(value or ""))


def hbv_rt_call_index(row):
    """Return detected RT calls keyed by canonical mutation name (for example rtM204V)."""
    return {
        str(item.get("mutation", "")): item
        for item in hbv_detected_calls(row)
        if item.get("gene") == "RT" and item.get("mutation")
    }


def hbv_pattern_mutation_text(index, mutations):
    parts = []
    for mutation in mutations:
        item = index.get(mutation)
        if item:
            parts.append(f"{mutation} ({float(item.get('frequency', 0) or 0):.1f} %)")
    return "; ".join(parts)


def hbv_pattern_max_frequency(pattern):
    """Return the highest reported component frequency in a combination result."""
    values = []
    for match in re.finditer(r"\(([0-9]+(?:\.[0-9]+)?)\s*%\)", str(pattern.get("mutations", "") or "")):
        try:
            values.append(float(match.group(1)))
        except ValueError:
            pass
    return max(values) if values else 0.0


def hbv_apply_population_context(patterns):
    """Separate biological rule strength from population dominance.

    An established resistance-associated pattern below 50% remains a real finding,
    but is presented as a yellow minority-population warning rather than as a
    dominant resistant population.
    """
    adjusted = []
    for original in patterns:
        pattern = dict(original)
        maximum = hbv_pattern_max_frequency(pattern)
        pattern["max_frequency"] = maximum
        if pattern.get("level") == "high":
            if maximum and maximum < 50.0:
                original_title = pattern.get("title", "Resistensassocierat mönster")
                pattern["level"] = "minor"
                pattern["title"] = "Resistensassocierat mönster i minoritetspopulation"
                pattern["interpretation"] = (
                    f"{original_title} påvisades, men den högsta rapporterade variantandelen "
                    f"är {maximum:.1f} %. Ingen dominerande resistensassocierad population "
                    "påvisades. Fyndet är en varningssignal och bör bedömas tillsammans med "
                    "behandlingshistorik och HBV-DNA-utveckling. " + pattern.get("interpretation", "")
                )
            else:
                pattern["level"] = "major"
                pattern["title"] = "Dominerande resistensassocierat mönster"
                pattern["interpretation"] = (
                    f"Ett etablerat resistensassocierat mönster påvisades i majoritetspopulation "
                    f"(högsta variantandel {maximum:.1f} %). Fyndet ger starkt genotypiskt stöd "
                    "för nedsatt läkemedelskänslighet. " + pattern.get("interpretation", "")
                )
        adjusted.append(pattern)
    return adjusted


def hbv_combination_interpretations(row):
    """Interpret established HBV RT mutation patterns without changing variant calls.

    Calls are not read-phased across distant codons. A detected combination therefore
    means that the component substitutions coexist in the sample, not necessarily on
    the same viral genome. Low-coverage samples are not assigned a definitive pattern.
    """
    if str(row.get("hbv_quality", "")) != "✓":
        return [{
            "category": "Teknisk begränsning",
            "level": "unknown",
            "title": "Kombinationsbedömning ej möjlig",
            "interpretation": (
                "En eller flera panelpositioner har otillräcklig täckning. "
                "Påvisade mutationer kan visas, men läkemedels- och kombinationsmönster "
                "ska inte bedömas som fullständiga."
            ),
            "mutations": "",
        }]

    idx = hbv_rt_call_index(row)
    present = set(idx)
    results = []

    m204 = [m for m in ("rtM204I", "rtM204V") if m in present]
    lam_background = bool(m204)
    l180 = "rtL180M" in present
    v173 = "rtV173L" in present
    a181 = [m for m in ("rtA181T", "rtA181V") if m in present]
    n236 = "rtN236T" in present
    etv_secondary = [m for m in present if re.match(r"^rt(?:T184|S202|M250)[A-Z*]$", m)]
    i169 = "rtI169T" in present

    # Lamivudine/telbivudine pathway.
    if lam_background and l180:
        muts = ["rtL180M"] + m204
        results.append({
            "category": "Lamivudin/telbivudin",
            "level": "high",
            "title": "Etablerat resistensmönster",
            "interpretation": (
                "Kombinationen rtL180M + rtM204I/V är ett klassiskt och väl etablerat "
                "resistensmönster mot lamivudin och telbivudin."
            ),
            "mutations": hbv_pattern_mutation_text(idx, muts),
        })
    elif lam_background:
        results.append({
            "category": "Lamivudin/telbivudin",
            "level": "high",
            "title": "Primär resistensmutation påvisad",
            "interpretation": (
                "rtM204I/V är en primär resistensmutation mot lamivudin/telbivudin. "
                "Avsaknad av rtL180M utesluter inte resistens."
            ),
            "mutations": hbv_pattern_mutation_text(idx, m204),
        })

    if v173 and l180 and "rtM204V" in present:
        muts = ["rtV173L", "rtL180M", "rtM204V"]
        results.append({
            "category": "Lamivudin/telbivudin",
            "level": "high",
            "title": "Kompenserat resistensmönster",
            "interpretation": (
                "rtV173L + rtL180M + rtM204V är ett etablerat lamivudinresistent mönster "
                "där rtV173L kan återställa replikationsförmåga."
            ),
            "mutations": hbv_pattern_mutation_text(idx, muts),
        })

    # Entecavir requires a lamivudine-resistant background plus secondary changes.
    if lam_background and etv_secondary:
        muts = (["rtL180M"] if l180 else []) + m204 + sorted(etv_secondary)
        results.append({
            "category": "Entekavir",
            "level": "high",
            "title": "Etablerat entekavirresistensmönster",
            "interpretation": (
                "Lamivudinresistent bakgrund med rtM204I/V och ytterligare substitution "
                "vid rtT184, rtS202 eller rtM250 är ett etablerat entekavirresistensmönster."
            ),
            "mutations": hbv_pattern_mutation_text(idx, muts),
        })
    elif etv_secondary:
        results.append({
            "category": "Entekavir",
            "level": "caution",
            "title": "Ofullständigt entekavirmönster",
            "interpretation": (
                "En entekavirassocierad sekundär mutation påvisades utan rtM204I/V. "
                "Fyndet ska inte ensamt tolkas som etablerad entekavirresistens."
            ),
            "mutations": hbv_pattern_mutation_text(idx, sorted(etv_secondary)),
        })

    if lam_background and i169:
        muts = (["rtL180M"] if l180 else []) + m204 + ["rtI169T"]
        results.append({
            "category": "Entekavir",
            "level": "caution",
            "title": "Entekavirassocierat kombinationsfynd",
            "interpretation": (
                "rtI169T tillsammans med lamivudinresistent bakgrund är associerat med "
                "minskad entekavirkänslighet, men bör bedömas tillsammans med övriga RT-fynd."
            ),
            "mutations": hbv_pattern_mutation_text(idx, muts),
        })
    elif i169:
        results.append({
            "category": "Entekavir",
            "level": "caution",
            "title": "Kontextberoende entekavirfynd",
            "interpretation": "rtI169T utan lamivudinresistent bakgrund är inte tillräckligt för etablerad entekavirresistens.",
            "mutations": hbv_pattern_mutation_text(idx, ["rtI169T"]),
        })

    # Adefovir pathway.
    adv_muts = a181 + (["rtN236T"] if n236 else [])
    if a181 and n236:
        results.append({
            "category": "Adefovir",
            "level": "high",
            "title": "Kombinerat adefovirresistensmönster",
            "interpretation": (
                "Substitution vid rtA181 tillsammans med rtN236T utgör ett starkt "
                "adefovirresistensmönster."
            ),
            "mutations": hbv_pattern_mutation_text(idx, adv_muts),
        })
    elif adv_muts:
        results.append({
            "category": "Adefovir",
            "level": "high",
            "title": "Etablerad adefovirresistensmutation",
            "interpretation": "rtA181T/V eller rtN236T är etablerat associerad med adefovirresistens.",
            "mutations": hbv_pattern_mutation_text(idx, adv_muts),
        })

    # Cross-/multidrug pattern.
    if lam_background and adv_muts:
        muts = (["rtL180M"] if l180 else []) + m204 + adv_muts
        results.append({
            "category": "Multiläkemedelsresistens",
            "level": "high",
            "title": "Mönster förenligt med resistens i flera läkemedelsklasser",
            "interpretation": (
                "Lamivudinresistent bakgrund tillsammans med rtA181T/V och/eller rtN236T "
                "är förenlig med kors- eller multiläkemedelsresistens."
            ),
            "mutations": hbv_pattern_mutation_text(idx, muts),
        })

    # Tenofovir remains cautionary.
    if adv_muts:
        results.append({
            "category": "Tenofovir",
            "level": "caution",
            "title": "Tenofovir – försiktighetsfynd",
            "interpretation": (
                "rtA181T/V och/eller rtN236T kan förekomma i behandlade populationer med "
                "minskad känslighet, men utgör inte ensamt bevis för klinisk tenofovirresistens."
            ),
            "mutations": hbv_pattern_mutation_text(idx, adv_muts),
        })

    if not results:
        results.append({
            "category": "Samlad RT-bedömning",
            "level": "none",
            "title": "Inget etablerat kombinationsmönster påvisat",
            "interpretation": (
                "Inget av de regelbaserade resistensmönstren påvisades över rapporteringströskeln. "
                "Enskilda associationer redovisas separat i mutationstabellen."
            ),
            "mutations": "",
        })
    return hbv_apply_population_context(results)


def hbv_combination_summary_text(row):
    patterns = hbv_combination_interpretations(row)
    important = [p for p in patterns if p.get("level") == "major"]
    selected = important or [p for p in patterns if p.get("level") == "minor"] or [p for p in patterns if p.get("level") == "caution"] or patterns
    return " | ".join(f"{p['category']}: {p['title']}" for p in selected)


def hbv_combination_rows_html(row):
    css = {"major": "pattern-high", "minor": "pattern-minor", "caution": "pattern-caution", "none": "pattern-none", "unknown": "pattern-unknown"}
    rows = []
    for pattern in hbv_combination_interpretations(row):
        rows.append(
            "<tr>"
            f"<td><strong>{html_escape(pattern['category'])}</strong></td>"
            f"<td><span class='pattern-status {css.get(pattern.get('level'), 'pattern-none')}'>{html_escape(pattern['title'])}</span></td>"
            f"<td>{html_escape(pattern.get('mutations') or '—')}</td>"
            f"<td>{html_escape(pattern['interpretation'])}</td>"
            "</tr>"
        )
    return "".join(rows)


def hbv_pedagogical_interpretation_html(row):
    """Render existing HBV rule results as concise Swedish teaching cards."""
    patterns = hbv_combination_interpretations(row)
    level_order = {"major": 0, "minor": 1, "caution": 2, "unknown": 3, "none": 4}
    patterns = sorted(patterns, key=lambda item: level_order.get(item.get("level"), 9))
    labels = {
        "major": ("Dominerande fynd", "interp-high"),
        "minor": ("Minoritetspopulation", "interp-minor"),
        "caution": ("Försiktighet", "interp-caution"),
        "unknown": ("Ej bedömbar", "interp-unknown"),
        "none": ("Inget mönster", "interp-none"),
    }
    cards = []
    for pattern in patterns:
        level = pattern.get("level", "none")
        badge, css_class = labels.get(level, labels["none"])
        mutation_text = pattern.get("mutations") or "Inga stödjande kombinationsmutationer"
        cards.append(
            f"<article class='interp-card {css_class}'>"
            f"<div class='interp-card-head'><div><div class='interp-kicker'>{html_escape(pattern['category'])}</div>"
            f"<h3>{html_escape(pattern['title'])}</h3></div>"
            f"<span class='interp-badge'>{html_escape(badge)}</span></div>"
            f"<p class='interp-meaning'><strong>Vad betyder fyndet?</strong> {html_escape(pattern['interpretation'])}</p>"
            f"<div class='interp-evidence'><strong>Stödjande fynd:</strong> {html_escape(mutation_text)}</div>"
            "</article>"
        )
    major = [p for p in patterns if p.get("level") == "major"]
    minor = [p for p in patterns if p.get("level") == "minor"]
    caution = [p for p in patterns if p.get("level") == "caution"]
    unknown = [p for p in patterns if p.get("level") == "unknown"]
    if unknown:
        lead = "Resultatet kan inte tolkas fullständigt eftersom en eller flera relevanta positioner har otillräcklig täckning."
        action = "Granska täckningen och överväg ny analys innan behandlingsrelaterad slutsats dras."
        lead_class = "lead-unknown"
    elif major:
        names = ", ".join(dict.fromkeys(p["category"] for p in major))
        lead = f"Ett eller flera etablerade resistensassocierade mönster påvisades i majoritetspopulation: {names}."
        action = "Fyndet ger starkt genotypiskt stöd för nedsatt känslighet och ska bedömas tillsammans med behandlingshistorik och HBV-DNA-utveckling."
        lead_class = "lead-high"
    elif minor:
        names = ", ".join(dict.fromkeys(p["category"] for p in minor))
        lead = f"Resistensassocierade mönster påvisades endast i minoritetspopulation: {names}. Ingen dominerande resistensassocierad population påvisades."
        action = "Gul varningssignal: dokumentera variantandelarna och bedöm fyndet tillsammans med behandling, följsamhet och HBV-DNA-utveckling."
        lead_class = "lead-minor"
    elif caution:
        names = ", ".join(dict.fromkeys(p["category"] for p in caution))
        lead = f"Kontextberoende eller ofullständiga resistensfynd påvisades: {names}."
        action = "Fynden ska inte ensamma användas för att fastställa klinisk resistens."
        lead_class = "lead-caution"
    else:
        lead = "Inget etablerat regelbaserat resistensmönster påvisades över rapporteringströskeln."
        action = "Enskilda mutationer och teknisk kvalitet ska ändå granskas i tabellerna nedan."
        lead_class = "lead-none"
    return (
        f"<section class='interpretation-overview'>"
        f"<div class='interpretation-lead {lead_class}'><strong>Samlad tolkning:</strong> {html_escape(lead)}"
        f"<div class='interpretation-action'><strong>Praktisk betydelse:</strong> {html_escape(action)}</div></div>"
        f"<div class='interp-grid'>{''.join(cards)}</div>"
        "</section>"
    )


def hbv_status_for_medication(row, medication):
    calls = []
    for item in hbv_detected_calls(row):
        evidence = item.get("medications", {}).get(medication)
        if evidence:
            calls.append((item, evidence))
    if not calls:
        return "Inga associerade mutationer påvisade", "", "none"
    max_frequency = max(float(item.get("frequency", 0) or 0) for item, _ in calls)
    level = "major" if max_frequency >= 50 else "minor"
    summary = "; ".join(
        f"{item['mutation']} ({item['frequency']} %, {hbv_swedish_evidence_label(evidence)})"
        for item, evidence in calls
    )
    if medication == "Tenofovir":
        status = "Försiktighet – litteraturassociation, ej ensam grund för resistensbedömning"
        level = "caution"
    elif level == "major":
        status = "Associerad mutation i majoritetspopulation"
    else:
        status = "Associerad mutation i minoritetspopulation"
    return status, summary, level


def safe_report_filename(sample):
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", str(sample or "prov")).strip("._")
    return safe or "prov"


def run_samples_folder(output_folder, run_tag):
    """Return the shared container for all per-sample/run subfolders."""
    return os.path.join(output_folder, f"Samples_{run_tag}")


def run_consensus_folder(output_folder, run_tag):
    """Return the run-level folder containing easy-to-find consensus copies."""
    return os.path.join(output_folder, f"Consensus_{run_tag}")


def is_valid_report_image(path):
    """Accept real PNG/JPEG files and reject macOS AppleDouble/resource-fork files."""
    if not path or not os.path.isfile(path):
        return False
    name = os.path.basename(path)
    if name.startswith("._") or name.startswith("."):
        return False
    try:
        with open(path, "rb") as handle:
            signature = handle.read(12)
        return (
            signature.startswith(b"\x89PNG\r\n\x1a\n")
            or signature.startswith(b"\xff\xd8\xff")
        )
    except Exception:
        return False


def export_run_consensus_files(result_rows, output_folder, run_tag):
    """Copy each sample's final consensus files into one run-level folder."""
    consensus_dir = run_consensus_folder(output_folder, run_tag)
    ensure_dir(consensus_dir, must_be_writable=True)
    copied = 0
    for row in result_rows:
        sample = safe_report_filename(row.get("sample", "sample"))
        sources = [
            (row.get("consensus_file"), f"{sample}.consensus.fasta"),
            (row.get("iub_consensus_file"), f"{sample}.consensus.iub.fasta"),
        ]
        for source, filename in sources:
            if not source or not os.path.isfile(source) or os.path.getsize(source) <= 0:
                continue
            destination = os.path.join(consensus_dir, filename)
            try:
                shutil.copy2(source, destination)
                copied += 1
            except Exception as exc:
                print(f"Could not copy consensus file {source} to {destination}: {exc}")
    if copied == 0:
        try:
            os.rmdir(consensus_dir)
        except OSError:
            pass
        print("No per-sample consensus files were available for run-level export.")
        return ""
    print(f"Per-sample consensus files saved to {consensus_dir}")
    return consensus_dir


REPORT_LOGO_FILENAMES = (
    "report_logo.png", "aurora_logo.png", "logo.png", "lab_logo.png", "cvr_logo.png"
)


def find_hbv_report_logo(row, report_path):
    """Find a PNG/JPEG logo near the reference, script, or report output."""
    candidates = []
    explicit = str(row.get("report_logo", "") or "").strip()
    if explicit:
        candidates.append(explicit)
    directories = [
        os.path.dirname(str(row.get("reference_file", "") or "")),
        app_base_dir(),
        os.path.dirname(report_path),
        os.path.dirname(os.path.dirname(report_path)),
    ]
    for directory in directories:
        if not directory:
            continue
        for filename in REPORT_LOGO_FILENAMES:
            candidates.append(os.path.join(directory, filename))
    for path in candidates:
        if is_valid_report_image(path):
            return path
    # Fallback: accept a clearly named logo PNG, or the only PNG in the reference folder.
    for directory in directories:
        if not directory or not os.path.isdir(directory):
            continue
        pngs = sorted(
            os.path.join(directory, name)
            for name in os.listdir(directory)
            if name.lower().endswith((".png", ".jpg", ".jpeg"))
            and not name.startswith("._")
            and not name.startswith(".")
            and is_valid_report_image(os.path.join(directory, name))
        )
        named_logos = [path for path in pngs if "logo" in os.path.basename(path).lower()]
        if named_logos:
            return named_logos[0]
        if len(pngs) == 1:
            return pngs[0]
    return ""


def image_file_to_data_uri(path):
    if not path:
        return ""
    try:
        mime = mimetypes.guess_type(path)[0] or "image/png"
        with open(path, "rb") as handle:
            encoded = base64.b64encode(handle.read()).decode("ascii")
        return f"data:{mime};base64,{encoded}"
    except Exception as exc:
        print(f"Kunde inte läsa rapportlogotyp {path}: {exc}")
        return ""


def hbv_overall_summary(row, calls):
    if row.get("hbv_quality") != "✓":
        return "Resultatet är inte fullständigt bedömbart på grund av otillräcklig täckning eller analysfel."
    if not calls:
        return "Inga resistens- eller escape-associerade panelmutationer påvisades över rapporteringströskeln."
    mutations = ", ".join(f"{item['mutation']} ({item['frequency']} %)" for item in calls)
    return f"Påvisade HBV-panelmutationer: {mutations}."


def hbv_panel_coverage_rows(row):
    rows = []
    for detail in row.get("hbv_mutation_details", []) or []:
        depth = int(detail.get("depth", 0) or 0)
        calls = detail.get("calls", []) or []
        nonsyn_calls = []
        for call in calls:
            alt = str(call.get("aa", "") or "")
            if alt and alt != detail.get("reference_aa"):
                nonsyn_calls.append(f"{alt} ({float(call.get('frequency', 0) or 0):.1f} %)")
        result = ", ".join(nonsyn_calls) or "Inget panelrelevant fynd"
        qc = "Godkänd" if depth >= HBV_MIN_COVERAGE else "Otillräcklig"
        rows.append(
            "<tr>"
            f"<td>{html_escape(str(detail.get('position','')))}</td>"
            f"<td>{html_escape(str(detail.get('reference_aa','')))}</td>"
            f"<td>{depth}</td>"
            f"<td>{html_escape(result)}</td>"
            f"<td>{qc}</td>"
            "</tr>"
        )
    return "".join(rows) or "<tr><td colspan='5' class='empty'>Paneldata saknas.</td></tr>"


def write_hbv_html_report(row, report_path, run_tag):
    calls = hbv_detected_calls(row)
    rt_calls = [x for x in calls if x["gene"] == "RT"]
    s_calls = [x for x in calls if x["gene"] == "S"]
    logo_uri = image_file_to_data_uri(find_hbv_report_logo(row, report_path))
    logo_html = f"<img class='report-logo' src='{logo_uri}' alt='Logotyp'>" if logo_uri else ""
    summary_text = hbv_overall_summary(row, calls)

    def status_class():
        if row.get("hbv_quality") != "✓":
            return "status-unknown"
        if any(float(x.get("frequency", 0) or 0) >= 50 for x in calls):
            return "status-high"
        if calls:
            return "status-medium"
        return "status-good"

    medication_rows = []
    swedish_drugs = {
        "Lamivudine": "Lamivudin",
        "Adefovir": "Adefovir",
        "Entecavir": "Entekavir",
        "Tenofovir": "Tenofovir",
    }
    for drug in HBV_MEDICATIONS:
        status, mutations, level = hbv_status_for_medication(row, drug)
        medication_rows.append(
            f"<tr><td><strong>{html_escape(swedish_drugs.get(drug, drug))}</strong></td>"
            f"<td><span class='drug-status drug-{level}'>{html_escape(status)}</span></td>"
            f"<td>{html_escape(mutations or '—')}</td></tr>"
        )

    def mutation_rows(items, empty_text):
        if not items:
            return f"<tr><td colspan='6' class='empty'>{html_escape(empty_text)}</td></tr>"
        rows_html = []
        for item in items:
            meds = ", ".join(
                f"{drug}: {hbv_swedish_evidence_label(evidence)}"
                for drug, evidence in item.get("medications", {}).items()
            ) or "—"
            flag = item.get("flag", "")
            flag_sv = {
                "PRIORITY": "Prioriterat fynd",
                "ESTABLISHED": "Etablerat fynd",
                "CAUTION": "Kontextberoende/försiktighet",
                "ESCAPE": "Escape-associerat fynd",
            }.get(flag, flag or "—")
            rows_html.append(
                "<tr>"
                f"<td><strong>{html_escape(str(item['mutation']))}</strong></td>"
                f"<td>{item['frequency']} %</td>"
                f"<td>{html_escape(item['population'])}</td>"
                f"<td>{item['count']} / {item['depth']}</td>"
                f"<td>{html_escape(flag_sv)}</td>"
                f"<td>{html_escape(meds)}</td>"
                "</tr>"
            )
        return "".join(rows_html)

    generated_date = datetime.now().strftime("%Y-%m-%d %H:%M")
    html = f"""<!DOCTYPE html>
<html lang='sv'>
<head>
<meta charset='UTF-8'>
<meta name='viewport' content='width=device-width, initial-scale=1'>
<title>HBV-resistensrapport – {html_escape(str(row.get('sample', '')))}</title>
<style>
:root {{ --blue:#dbeeff; --blue-dark:#185b8d; --line:#c9d3dc; --text:#1e2a32; --muted:#5d6b75; }}
* {{ box-sizing:border-box; }}
body {{ font-family:Arial,Helvetica,sans-serif; color:var(--text); margin:0; background:#eef2f5; }}
.report {{ max-width:1050px; margin:24px auto; background:#fff; padding:34px 42px; box-shadow:0 3px 18px rgba(0,0,0,.12); }}
h1 {{ margin:0 0 5px; font-size:28px; color:#123e5f; }}
.header {{ display:flex; justify-content:space-between; align-items:flex-start; gap:24px; margin-bottom:20px; }}
.title-block {{ flex:1; }}
.report-logo {{ max-width:220px; max-height:105px; object-fit:contain; }}
.subtitle {{ color:var(--muted); margin-bottom:0; }}
.summary {{ background:#f4f8fb; border-left:6px solid var(--blue-dark); padding:12px 15px; margin:14px 0 18px; }}
h2 {{ font-size:18px; margin:28px 0 10px; border-bottom:2px solid var(--blue-dark); padding-bottom:6px; }}
table {{ width:100%; border-collapse:collapse; font-size:14px; }}
th {{ background:var(--blue); text-align:left; }}
th,td {{ border:1px solid var(--line); padding:9px 10px; vertical-align:top; }}
.meta td:first-child {{ width:28%; background:#f4f8fb; font-weight:bold; }}
.status {{ padding:14px 16px; border-radius:6px; font-weight:bold; margin:16px 0; }}
.status-good {{ background:#dff3e3; border-left:7px solid #3e9d54; }}
.status-medium {{ background:#fff4c7; border-left:7px solid #d0a600; }}
.status-high {{ background:#ffe1d6; border-left:7px solid #d86032; }}
.status-unknown {{ background:#e3e6e8; border-left:7px solid #6f777d; }}
.drug-status {{ font-weight:bold; }}
.drug-none {{ color:#2c7b3e; }} .drug-minor {{ color:#a17900; }} .drug-major {{ color:#b23b20; }} .drug-caution {{ color:#8a5a00; }}
.pattern-status {{ font-weight:bold; }} .pattern-high {{ color:#b23b20; }} .pattern-minor {{ color:#8a6a00; }} .pattern-caution {{ color:#8a5a00; }} .pattern-none {{ color:#2c7b3e; }} .pattern-unknown {{ color:#6f777d; }}
.interpretation-overview {{ margin:12px 0 20px; }}
.interpretation-lead {{ padding:15px 17px; border-radius:7px; border-left:7px solid #6f777d; background:#f4f6f7; line-height:1.45; }}
.interpretation-action {{ margin-top:6px; font-size:13px; font-weight:normal; }}
.lead-high {{ background:#ffe8df; border-left-color:#c94d28; }}
.lead-minor {{ background:#fff5cf; border-left-color:#d2a600; }}
.lead-caution {{ background:#fff5cf; border-left-color:#c99b00; }}
.lead-none {{ background:#e5f4e8; border-left-color:#438f55; }}
.lead-unknown {{ background:#e8ebed; border-left-color:#6f777d; }}
.interp-grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(310px,1fr)); gap:12px; margin-top:12px; }}
.interp-card {{ border:1px solid var(--line); border-left:7px solid #6f777d; border-radius:7px; padding:14px 15px; background:#fff; break-inside:avoid; }}
.interp-card-head {{ display:flex; justify-content:space-between; gap:12px; align-items:flex-start; }}
.interp-card h3 {{ margin:3px 0 8px; font-size:16px; color:#183f5c; }}
.interp-kicker {{ font-size:12px; font-weight:bold; text-transform:uppercase; letter-spacing:.04em; color:var(--muted); }}
.interp-badge {{ white-space:nowrap; border-radius:999px; padding:4px 9px; font-size:11px; font-weight:bold; background:#eef1f3; }}
.interp-meaning {{ margin:7px 0 10px; line-height:1.42; }}
.interp-evidence {{ border-top:1px solid #e1e6ea; padding-top:8px; font-size:12px; color:#46545e; }}
.interp-high {{ border-left-color:#c94d28; background:#fff9f7; }} .interp-high .interp-badge {{ background:#f8d7cc; color:#8d2e17; }}
.interp-minor {{ border-left-color:#d2a600; background:#fffdf2; }} .interp-minor .interp-badge {{ background:#ffefad; color:#705700; }}
.interp-caution {{ border-left-color:#c99b00; background:#fffdf5; }} .interp-caution .interp-badge {{ background:#fff0b7; color:#795d00; }}
.interp-none {{ border-left-color:#438f55; background:#fbfefb; }} .interp-none .interp-badge {{ background:#dcefe0; color:#28683a; }}
.interp-unknown {{ border-left-color:#6f777d; background:#fafbfb; }} .interp-unknown .interp-badge {{ background:#e2e5e7; color:#50585d; }}
.rule-details {{ margin-top:12px; border:1px solid var(--line); border-radius:6px; padding:10px 12px; background:#fafcfd; }}
.rule-details summary {{ cursor:pointer; color:#185b8d; }}
.empty {{ color:var(--muted); font-style:italic; }}
.notice {{ background:#f5f5f5; border-left:5px solid #777; padding:12px 14px; margin-top:26px; font-size:12px; }}
.small {{ font-size:12px; color:var(--muted); }}
.report-footer {{
    margin-top: 24px;
    padding-top: 10px;
    border-top: 1px solid var(--line);
    color: var(--muted);
    font-size: 10px;
    line-height: 1.35;
    text-align: center;
}}
@media print {{ body {{ background:#fff; }} .report {{ box-shadow:none; margin:0; max-width:none; }} .report-logo {{ max-height:85px; }} }}
</style>
</head>
<body><main class='report'>
<div class='header'><div class='title-block'>
<h1>Hepatit B-virus – resistens och immunologisk escape</h1>
<div class='subtitle'>Läsnivåbaserad mutationsanalys</div>
</div>{logo_html}</div>
<table class='meta'>
<tr><td>Prov-ID</td><td>{html_escape(str(row.get('sample','')))}</td></tr>
<tr><td>Bästa referens</td><td>{html_escape(str(row.get('best_reference','')))}</td></tr>
<tr><td>Rapportdatum</td><td>{generated_date}</td></tr>
<tr><td>Körning</td><td>{html_escape(str(run_tag))}</td></tr>
<tr><td>Rapporteringströskel</td><td>≥ {HBV_MIN_VARIANT_FREQUENCY * 100:.1f} %</td></tr>
<tr><td>Minsta kodontäckning</td><td>{HBV_MIN_COVERAGE} läsningar</td></tr>
<tr><td>Genomtäckning</td><td>{row.get('target_or_genome_coverage_percent', 0)} %</td></tr>
<tr><td>Mediandjup</td><td>{row.get('median_depth', 0)}×</td></tr>
<tr><td>Antal mappade läsningar</td><td>{row.get('mapped_reads', 0)}</td></tr>
</table>
<div class='summary'><strong>Sammanfattning:</strong> {html_escape(summary_text)}</div>
<div class='status {status_class()}'>{html_escape(hbv_swedish_interpretation(row.get('hbv_interpretation','')))}</div>

<h2>Pedagogisk resistens- och kombinationsbedömning</h2>
{hbv_pedagogical_interpretation_html(row)}
<details class='rule-details'>
<summary><strong>Visa fullständig regelbaserad bedömning</strong></summary>
<table style='margin-top:10px'><thead><tr><th>Läkemedel/mönster</th><th>Bedömning</th><th>Mutationer i provet</th><th>Tolkning</th></tr></thead>
<tbody>{hbv_combination_rows_html(row)}</tbody></table>
</details>
<div class='small' style='margin-top:8px'><strong>Viktigt:</strong> Mutationerna analyseras på läsnivå per kodon men är inte fasade över hela RT-regionen. Ett kombinationsfynd visar att komponenterna finns i samma prov, men inte säkert på samma virusgenom.</div>

<h2>Läkemedelsöversikt</h2>
<table><thead><tr><th>Läkemedel</th><th>Bedömning</th><th>Påvisade mutationer</th></tr></thead>
<tbody>{''.join(medication_rows)}</tbody></table>

<h2>RT-resistensassocierade mutationer</h2>
<table><thead><tr><th>Mutation</th><th>Andel</th><th>Population</th><th>Mutantläsningar / djup</th><th>Typ</th><th>Läkemedelsassociation</th></tr></thead>
<tbody>{mutation_rows(rt_calls, 'Inga RT-panelmutationer påvisades.')}</tbody></table>

<h2>Ytantigen- och escape-associerade mutationer</h2>
<table><thead><tr><th>Mutation</th><th>Andel</th><th>Population</th><th>Mutantläsningar / djup</th><th>Typ</th><th>Läkemedelsassociation</th></tr></thead>
<tbody>{mutation_rows(s_calls, 'Inga escape-associerade panelmutationer påvisades.')}</tbody></table>

<h2>Teknisk kvalitet</h2>
<table class='meta'>
<tr><td>Kvalitetsstatus</td><td>{html_escape(str(row.get('hbv_quality','')))}</td></tr>
<tr><td>Övergripande QC</td><td>{html_escape(str(row.get('qc_verdict','')))}</td></tr>
<tr><td>QC-kommentar</td><td>{html_escape(str(row.get('qc_reason','') or '—'))}</td></tr>
<tr><td>Positioner med otillräcklig täckning</td><td>{html_escape(str(row.get('hbv_low_coverage_positions',' ') or '—'))}</td></tr>
<tr><td>Medeldjup</td><td>{row.get('mean_depth', 0)}×</td></tr>
<tr><td>Minimidjup</td><td>{row.get('min_depth', 0)}×</td></tr>
<tr><td>Positioner ≥10×</td><td>{row.get('positions_ge_10x_percent', 0)} %</td></tr>
<tr><td>Positioner ≥30×</td><td>{row.get('positions_ge_30x_percent', 0)} %</td></tr>
<tr><td>Konsensuslängd</td><td>{row.get('consensus_length', 0)} nt</td></tr>
<tr><td>N-baser</td><td>{row.get('n_count', 0)} ({row.get('n_percent', 0)} %)</td></tr>
<tr><td>Konsensusåtervinning</td><td>{row.get('consensus_recovery_percent', 0)} %</td></tr>
<tr><td>RT-koordinat</td><td>start {html_escape(str(row.get('hbv_rt_codon1_start','')))}, metod {html_escape(str(row.get('hbv_coordinate_method','metadata')))}</td></tr>
<tr><td>S-koordinat</td><td>start {html_escape(str(row.get('hbv_s_codon1_start','')))}, metod {html_escape(str(row.get('hbv_coordinate_method','metadata')))}</td></tr>
<tr><td>Medaka-modell</td><td>{html_escape(str(row.get('medaka_model','')))}</td></tr>
<tr><td>FASTQ-fil</td><td class='small'>{html_escape(str(row.get('input_fastq','')))}</td></tr>
<tr><td>Konsensusfil</td><td class='small'>{html_escape(str(row.get('consensus_file','')))}</td></tr>
<tr><td>BAM-fil</td><td class='small'>{html_escape(str(row.get('bam_file','')))}</td></tr>
<tr><td>Referensfil</td><td class='small'>{html_escape(str(row.get('reference_file','')))}</td></tr>
<tr><td>Scriptversion</td><td>{html_escape(str(row.get('hbv_mutation_script_version','')))}</td></tr>
</table>

<h2>Paneltäckning och positionskontroll</h2>
<table><thead><tr><th>Position</th><th>Referens-AA</th><th>Kodontäckning</th><th>Panelrelevant resultat</th><th>QC</th></tr></thead>
<tbody>{hbv_panel_coverage_rows(row)}</tbody></table>
<div class='notice'><strong>Viktig begränsning:</strong> Resultatet är en litteraturbaserad screening och är inte en fristående klinisk resistensbedömning. Regelbaserade kombinationsfynd ska bedömas tillsammans med variantandelar, behandlingshistorik, HBV-DNA-utveckling, genotyp och kliniska data. Tenofovirassociationer är försiktighetsfynd och inte ensamma bevis för klinisk resistens.</div>

<footer class='report-footer'>
Vid frågor, tekniska problem eller oväntade analysresultat, kontakta skriptets författare Joakim Bedner Stenbäck (joakim.stenback@vgregion.se).
</footer>

</main></body></html>"""
    with open(report_path, "w", encoding="utf-8") as handle:
        handle.write(html)


def log_hbv_mutation_results(rows, output_folder, run_tag):
    """Write a streamlined Swedish Excel workbook and one Swedish HTML report per sample."""
    ensure_dir(output_folder, must_be_writable=True)
    excel_path = os.path.join(output_folder, f"hbv_resistensresultat_{run_tag}.xlsx")
    html_dir = os.path.join(run_samples_folder(output_folder, run_tag), "HBV_HTML_reports")
    ensure_dir(html_dir, must_be_writable=True)

    wb = Workbook()
    overview = wb.active
    overview.title = "Översikt"
    medications = wb.create_sheet("Läkemedel")
    mutations = wb.create_sheet("Mutationer")
    combinations = wb.create_sheet("Kombinationsbedömning")
    technical = wb.create_sheet("Teknisk QC")
    panel = wb.create_sheet("Paneldefinition")

    dark_blue = "1F4E78"
    light_blue = "D9EAF7"
    pale_blue = "EAF3F8"
    green = "C6EFCE"
    yellow = "FFEB9C"
    orange = "FCE4D6"
    red = "FFC7CE"
    grey = "E7E6E6"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    overview.merge_cells("A1:K1")
    overview["A1"] = "Hepatit B-virus – resistens och immunologisk escape"
    overview["A1"].font = Font(size=18, bold=True, color=white)
    overview["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    overview["A1"].alignment = Alignment(horizontal="left", vertical="center")
    overview.row_dimensions[1].height = 30
    overview.merge_cells("A2:K2")
    overview["A2"] = (
        f"Läsnivåtröskel ≥{HBV_MIN_VARIANT_FREQUENCY * 100:.1f}% | "
        f"Minsta kodontäckning {HBV_MIN_COVERAGE} | Körning {run_tag}"
    )
    overview["A2"].font = Font(italic=True, color="4F5B62")

    overview_headers = [
        "Prov-ID", "Kvalitet", "Samlad bedömning", "Bästa referens",
        "Prioriterade RT-fynd", "Övriga RT-fynd", "Escape-fynd",
        "Kombinationsbedömning", "Läkemedelsöversikt", "Otillräcklig täckning", "HTML-rapport",
    ]
    overview.append([])
    overview.append(overview_headers)

    medications.append(["Prov-ID", "Läkemedel", "Bedömning", "Påvisade mutationer", "Högsta andel (%)"])
    combinations.append(["Prov-ID", "Läkemedel/mönster", "Bedömning", "Mutationer i provet", "Tolkning", "Evidensnivå"])
    mutations.append([
        "Prov-ID", "Gen", "Mutation", "Andel (%)", "Population", "Mutantläsningar",
        "Kodondjup", "Fyndtyp", "Kategori", "Läkemedelsassociationer", "Bästa referens",
    ])
    technical.append([
        "Prov-ID", "Kvalitet", "Samlad bedömning", "Bästa referens", "Otillräcklig täckning",
        "BAM-fil", "Minsta kodontäckning", "Rapporteringströskel", "Scriptversion", "Körning",
    ])
    panel.append(["Gen", "Position", "Rapporterbara substitutioner", "Fyndtyp", "Kategori/läkemedel"])

    html_links = []
    swedish_drugs = {"Lamivudine": "Lamivudin", "Adefovir": "Adefovir", "Entecavir": "Entekavir", "Tenofovir": "Tenofovir"}

    for row in rows:
        sample = row.get("sample", "")
        calls = hbv_detected_calls(row)
        rt_calls = [x for x in calls if x["gene"] == "RT"]
        escape_calls = [x for x in calls if x["gene"] == "S"]
        priority = [x for x in rt_calls if x.get("flag") == "PRIORITY"]
        other_rt = [x for x in rt_calls if x.get("flag") != "PRIORITY"]

        def compact(items):
            return "; ".join(f"{x['mutation']} ({x['frequency']} %)" for x in items) or "—"

        medication_parts = []
        for drug in HBV_MEDICATIONS:
            _, mutation_text, _ = hbv_status_for_medication(row, drug)
            if mutation_text:
                medication_parts.append(f"{swedish_drugs.get(drug, drug)}: {mutation_text}")
        medication_overview = " | ".join(medication_parts) or "Inga läkemedelsassocierade mutationer"

        filename = f"{safe_report_filename(sample)}_{run_tag}_HBV.html"
        html_path = os.path.join(html_dir, filename)
        write_hbv_html_report(row, html_path, run_tag)
        html_links.append((sample, filename))

        overview.append([
            sample,
            row.get("hbv_quality", ""),
            hbv_swedish_interpretation(row.get("hbv_interpretation", "")),
            row.get("best_reference", ""),
            compact(priority), compact(other_rt), compact(escape_calls), hbv_combination_summary_text(row), medication_overview,
            row.get("hbv_low_coverage_positions", " ") or "—",
            filename,
        ])
        link_cell = overview.cell(overview.max_row, 11)
        link_cell.hyperlink = os.path.relpath(html_path, output_folder)
        link_cell.style = "Hyperlink"

        for pattern in hbv_combination_interpretations(row):
            combinations.append([
                sample, pattern.get("category", ""), pattern.get("title", ""),
                pattern.get("mutations", "") or "—", pattern.get("interpretation", ""),
                {"major": "Dominerande etablerat mönster", "minor": "Etablerat mönster i minoritetspopulation", "caution": "Försiktighet/kontextberoende", "none": "Inget mönster", "unknown": "Ej bedömbar"}.get(pattern.get("level"), pattern.get("level", "")),
            ])

        for drug in HBV_MEDICATIONS:
            status, mutation_text, _ = hbv_status_for_medication(row, drug)
            frequencies = [x["frequency"] for x in calls if drug in x.get("medications", {})]
            medications.append([
                sample, swedish_drugs.get(drug, drug), status, mutation_text or "—",
                max(frequencies) if frequencies else "",
            ])

        if calls:
            for call in calls:
                meds = "; ".join(
                    f"{swedish_drugs.get(drug, drug)}: {hbv_swedish_evidence_label(evidence)}"
                    for drug, evidence in call.get("medications", {}).items()
                ) or "—"
                flag_sv = {
                    "PRIORITY": "Prioriterat",
                    "ESTABLISHED": "Etablerat",
                    "CAUTION": "Kontextberoende/försiktighet",
                    "ESCAPE": "Escape-associerat",
                }.get(call.get("flag"), call.get("flag") or "—")
                mutations.append([
                    sample, call["gene"], call["mutation"], call["frequency"], call["population"],
                    call["count"], call["depth"], flag_sv, hbv_translate_free_text(call.get("category", "")), meds,
                    row.get("best_reference", ""),
                ])
        else:
            mutations.append([sample, "—", "Inga panelmutationer påvisade", "", "", "", "", "", "", "", row.get("best_reference", "")])

        technical.append([
            sample, row.get("hbv_quality", ""), hbv_swedish_interpretation(row.get("hbv_interpretation", "")),
            row.get("best_reference", ""), row.get("hbv_low_coverage_positions", " ") or "—",
            row.get("bam_file", ""), HBV_MIN_COVERAGE, f">={HBV_MIN_VARIANT_FREQUENCY * 100:.1f}%",
            row.get("hbv_mutation_script_version", ""), run_tag,
        ])

    for aa_pos, rule in HBV_RT_PANEL.items():
        panel.append(["RT", rule.get("label", f"rt{aa_pos}"), ", ".join(sorted(rule.get("aas", set()))), rule.get("flag", ""), hbv_translate_free_text(rule.get("drugs", ""))])
    for aa_pos, rule in HBV_S_ESCAPE_PANEL.items():
        panel.append(["S", rule.get("label", f"s{aa_pos}"), ", ".join(sorted(rule.get("aas", set()))), "Escape-associerat", hbv_translate_free_text(rule.get("category", ""))])

    for ws in [overview, medications, combinations, mutations, technical, panel]:
        header_row = 4 if ws is overview else 1
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=dark_blue)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row_cells in ws.iter_rows(min_row=header_row + 1):
            for cell in row_cells:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.sheet_view.showGridLines = False

    for row_idx in range(5, overview.max_row + 1):
        quality = overview.cell(row_idx, 2)
        assessment = overview.cell(row_idx, 3)
        quality.fill = PatternFill("solid", fgColor=green if quality.value == "✓" else red)
        text_value = str(assessment.value or "").lower()
        if "inga panelmutationer" in text_value:
            assessment.fill = PatternFill("solid", fgColor=green)
        elif "ej bedömbar" in text_value:
            assessment.fill = PatternFill("solid", fgColor=grey)
        elif "minoritet" in text_value:
            assessment.fill = PatternFill("solid", fgColor=yellow)
        else:
            assessment.fill = PatternFill("solid", fgColor=orange)

    overview.column_dimensions["A"].width = 36
    overview.column_dimensions["B"].width = 11
    overview.column_dimensions["C"].width = 35
    overview.column_dimensions["D"].width = 16
    for col in ["E", "F", "G"]:
        overview.column_dimensions[col].width = 30
    overview.column_dimensions["H"].width = 48
    overview.column_dimensions["I"].width = 58
    overview.column_dimensions["J"].width = 35
    overview.column_dimensions["K"].width = 34

    for ws in [medications, combinations, mutations, technical, panel]:
        autosize_worksheet(ws)
    medications.column_dimensions["D"].width = 65
    combinations.column_dimensions["A"].width = 28
    combinations.column_dimensions["B"].width = 26
    combinations.column_dimensions["C"].width = 38
    combinations.column_dimensions["D"].width = 55
    combinations.column_dimensions["E"].width = 85
    combinations.column_dimensions["F"].width = 28
    mutations.column_dimensions["I"].width = 38
    mutations.column_dimensions["J"].width = 52
    technical.column_dimensions["F"].width = 75
    panel.column_dimensions["E"].width = 60

    wb.save(excel_path)

    index_path = os.path.join(html_dir, f"index_{run_tag}.html")
    links = "".join(
        f"<li><a href='{html_escape(filename)}'>{html_escape(str(sample))}</a></li>"
        for sample, filename in html_links
    )
    with open(index_path, "w", encoding="utf-8") as handle:
        handle.write(
            "<!DOCTYPE html><html lang='sv'><head><meta charset='UTF-8'>"
            "<title>HBV-rapporter</title><style>body{font-family:Arial;margin:40px;max-width:900px}"
            "li{margin:8px 0}a{color:#185b8d}</style></head><body>"
            f"<h1>HBV-rapporter – {html_escape(run_tag)}</h1><ul>{links}</ul></body></html>"
        )

    print(f"Svensk HBV-resistensrapport sparad: {excel_path}")
    print(f"HTML-rapporter sparade i: {html_dir}")

#Logging

def autosize_worksheet(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 60)

def log_reference_usage(result_rows, output_folder, virus_type, run_tag):
    try:
        ensure_dir(output_folder, must_be_writable=True)
        excel_path = os.path.join(output_folder, f"reference_usage_{run_tag}.xlsx")

        wb = Workbook()

        ws = wb.active
        ws.title = "Summary"

        details_ws = wb.create_sheet("Details")

        summary_headers = [
            "Sample",
            "Virus Type",
            "Best Reference (SN)",
            "Mapped Reads",
            "Target/Genome Coverage (%)",
            "Mean Depth",
            "Median Depth",
            "Minimum Depth",
            "% Positions >=10x",
            "% Positions >=30x",
            "Consensus Length",
            "N Count",
            "N %",
            "Consensus Recovery vs Expected (%)",
            "QC Verdict",
            "QC Reason",
            "BAM File"
        ]

        if virus_type == "Hepatitis B Virus":
            summary_headers.extend([
                "HBV System 1 Full Median Depth",
                "HBV System 1 Full % >=10x",
                "HBV System 1 Full % >=30x",
                "HBV System 2 Full Median Depth",
                "HBV System 2 Full % >=10x",
                "HBV System 2 Full % >=30x",
                "HBV System 1 Unique Median Depth",
                "HBV System 1 Unique % >=10x",
                "HBV System 1 Unique % >=30x",
                "HBV System 2 Unique Median Depth",
                "HBV System 2 Unique % >=10x",
                "HBV System 2 Unique % >=30x",
                "HBV Overlap Median Depth",
                "HBV Overlap % >=10x",
                "HBV Overlap % >=30x",
                "HBV System Balance Ratio",
                "HBV System QC Warning"
            ])

        detail_headers = [
            "Sample",
            "Reference Length",
            "Detected Target Start",
            "Detected Target End",
            "Detected Target Length",
            "Expected Target/Genome Length",
            "Coverage Type"
        ]

        ws.append(summary_headers)
        details_ws.append(detail_headers)

        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

        for row in result_rows:
            summary_row = [
                row.get("sample", ""),
                row.get("virus_type", ""),
                row.get("best_reference", ""),
                row.get("mapped_reads", 0),
                row.get("target_or_genome_coverage_percent", 0),
                row.get("mean_depth", 0),
                row.get("median_depth", 0),
                row.get("min_depth", 0),
                row.get("positions_ge_10x_percent", 0),
                row.get("positions_ge_30x_percent", 0),
                row.get("consensus_length", 0),
                row.get("n_count", 0),
                row.get("n_percent", 0),
                row.get("consensus_recovery_percent", 0),
                row.get("qc_verdict", ""),
                row.get("qc_reason", ""),
                row.get("bam_file", ""),
            ]

            if virus_type == "Hepatitis B Virus":
                summary_row.extend([
                    row.get("hbv_system1_full_median_depth", 0),
                    row.get("hbv_system1_full_positions_ge_10x_percent", 0),
                    row.get("hbv_system1_full_positions_ge_30x_percent", 0),
                    row.get("hbv_system2_full_median_depth", 0),
                    row.get("hbv_system2_full_positions_ge_10x_percent", 0),
                    row.get("hbv_system2_full_positions_ge_30x_percent", 0),
                    row.get("hbv_system1_unique_median_depth", 0),
                    row.get("hbv_system1_unique_positions_ge_10x_percent", 0),
                    row.get("hbv_system1_unique_positions_ge_30x_percent", 0),
                    row.get("hbv_system2_unique_median_depth", 0),
                    row.get("hbv_system2_unique_positions_ge_10x_percent", 0),
                    row.get("hbv_system2_unique_positions_ge_30x_percent", 0),
                    row.get("hbv_overlap_median_depth", 0),
                    row.get("hbv_overlap_positions_ge_10x_percent", 0),
                    row.get("hbv_overlap_positions_ge_30x_percent", 0),
                    row.get("hbv_system_balance_ratio", 0),
                    row.get("hbv_system_qc_warning", ""),
                ])

            ws.append(summary_row)

            details_ws.append([
                row.get("sample", ""),
                row.get("reference_length", 0),
                row.get("detected_target_start", ""),
                row.get("detected_target_end", ""),
                row.get("detected_target_length", 0),
                row.get("expected_target_length", 0),
                row.get("coverage_type", ""),
            ])

            mapped_cell = ws.cell(row=ws.max_row, column=4)
            mapped_reads = row.get("mapped_reads", 0)
            mapped_cell.fill = red_fill if mapped_reads < 100 else green_fill

            coverage_cell = ws.cell(row=ws.max_row, column=5)
            coverage = row.get("target_or_genome_coverage_percent", 0)
            coverage_fill(coverage_cell, coverage, red_fill, yellow_fill, green_fill)

            qc_cell = ws.cell(row=ws.max_row, column=15)
            qc_fill(qc_cell, row.get("qc_verdict", ""), red_fill, yellow_fill, green_fill)

            if virus_type == "Hepatitis B Virus":
                # Highlight the unique-region system metrics because those drive QC.
                for col in [24, 25, 26, 27, 28, 29]:
                    cell = ws.cell(row=ws.max_row, column=col)
                    try:
                        value = float(cell.value)
                    except Exception:
                        value = 0.0

                    if "Median Depth" in str(ws.cell(row=1, column=col).value):
                        if value >= 30:
                            cell.fill = green_fill
                        elif value >= 10:
                            cell.fill = yellow_fill
                        else:
                            cell.fill = red_fill
                    else:
                        coverage_fill(cell, value, red_fill, yellow_fill, green_fill)

        autosize_worksheet(ws)
        autosize_worksheet(details_ws)

        wb.save(excel_path)
        print(f"Reference usage Excel saved to {excel_path}")

    except Exception as e:
        print(f"Error logging reference usage: {e}")

# Per-sample processing (loop)

def add_hbv_system_stats_to_row(row, hbv_system_stats, hbv_flags):
    s1_full = hbv_system_stats.get("system1_full", empty_depth_stats())
    s2_full = hbv_system_stats.get("system2_full", empty_depth_stats())
    s1_unique = hbv_system_stats.get("system1_unique", empty_depth_stats())
    s2_unique = hbv_system_stats.get("system2_unique", empty_depth_stats())
    overlap = hbv_system_stats.get("overlap_only", empty_depth_stats())

    row["hbv_system1_full_median_depth"] = s1_full.get("median_depth", 0)
    row["hbv_system1_full_positions_ge_10x_percent"] = s1_full.get("positions_ge_10x_percent", 0)
    row["hbv_system1_full_positions_ge_30x_percent"] = s1_full.get("positions_ge_30x_percent", 0)

    row["hbv_system2_full_median_depth"] = s2_full.get("median_depth", 0)
    row["hbv_system2_full_positions_ge_10x_percent"] = s2_full.get("positions_ge_10x_percent", 0)
    row["hbv_system2_full_positions_ge_30x_percent"] = s2_full.get("positions_ge_30x_percent", 0)

    row["hbv_system1_unique_median_depth"] = s1_unique.get("median_depth", 0)
    row["hbv_system1_unique_positions_ge_10x_percent"] = s1_unique.get("positions_ge_10x_percent", 0)
    row["hbv_system1_unique_positions_ge_30x_percent"] = s1_unique.get("positions_ge_30x_percent", 0)

    row["hbv_system2_unique_median_depth"] = s2_unique.get("median_depth", 0)
    row["hbv_system2_unique_positions_ge_10x_percent"] = s2_unique.get("positions_ge_10x_percent", 0)
    row["hbv_system2_unique_positions_ge_30x_percent"] = s2_unique.get("positions_ge_30x_percent", 0)

    row["hbv_overlap_median_depth"] = overlap.get("median_depth", 0)
    row["hbv_overlap_positions_ge_10x_percent"] = overlap.get("positions_ge_10x_percent", 0)
    row["hbv_overlap_positions_ge_30x_percent"] = overlap.get("positions_ge_30x_percent", 0)

    row["hbv_system_balance_ratio"] = hbv_flags.get("balance_ratio", 0)
    row["hbv_system_qc_warning"] = hbv_flags.get("warning", "")

    return row

def process_sample(
    input_file,
    reference_files,
    output_folder,
    medaka_model,
    virus_type,
    final_consensus_list,
    result_rows,
    auto_detect_models=False
):
    try:
        base_name = os.path.splitext(os.path.basename(input_file))[0].replace(".fastq", "").replace(".gz", "")
        sample_dir = os.path.join(output_folder, base_name)
        ensure_dir(sample_dir, must_be_writable=True)
        selected_medaka_model = choose_medaka_model(input_file, medaka_model, auto_detect_models)

        best = {
            "ref_path": None,
            "ref_sn": None,
            "approx_coverage": -1.0,
            "sam_path": None,
            "mapped": 0,
            "aligned_query_bases": 0,
            "ref_length": 0
        }

        hcv_best_mapped_by_subtype = {"1A": 0, "1B": 0}

        for reference_file in reference_files:
            ref_name = os.path.splitext(os.path.basename(reference_file))[0]
            sam_tmp = os.path.join(sample_dir, f"{ref_name}.sam")

            cmd = ["minimap2", "-ax", "map-ont", reference_file, input_file]
            with open(sam_tmp, "w", encoding="utf-8") as sam_out:
                subprocess.run(cmd, check=True, stdout=sam_out)

            approx_cov, mapped, sn, aligned_query_bases, ref_length = evaluate_alignment_quality(sam_tmp)

            if virus_type == "Hepatitis C Virus":
                subtype = hcv_subtype_from_name(sn or ref_name)
                if subtype in hcv_best_mapped_by_subtype:
                    hcv_best_mapped_by_subtype[subtype] = max(hcv_best_mapped_by_subtype[subtype], mapped)

            if approx_cov > best["approx_coverage"]:
                safe_remove(best["sam_path"])
                best.update({
                    "ref_path": reference_file,
                    "ref_sn": sn,
                    "approx_coverage": approx_cov,
                    "sam_path": sam_tmp,
                    "mapped": mapped,
                    "aligned_query_bases": aligned_query_bases,
                    "ref_length": ref_length
                })
            else:
                safe_remove(sam_tmp)

        if not (best["ref_path"] and best["sam_path"]):
            print(f"No suitable reference found for {input_file}. Skipping.")
            return

        detected_target_start, detected_target_end, detected_target_length = detect_aligned_target_region_from_sam(
            best["sam_path"],
            best["ref_sn"]
        )

        tmp_ref = os.path.join(sample_dir, f"{base_name}_best_reference.fasta")
        found = False

        with open(best["ref_path"], "r", encoding="utf-8", errors="replace") as ref_in, \
             open(tmp_ref, "w", encoding="utf-8") as ref_out:
            write = False
            for line in ref_in:
                if line.startswith(">"):
                    header = line[1:].strip()
                    name_only = header.split()[0]
                    write = (name_only == best["ref_sn"])
                if write:
                    ref_out.write(line)
                    found = True

        if not found:
            print(f"Could not find sequence '{best['ref_sn']}' in {best['ref_path']}. Using full reference instead.")
            tmp_ref = best["ref_path"]

        bam_sorted = os.path.join(sample_dir, f"{base_name}.sorted.bam")
        subprocess.run(["samtools", "sort", "-O", "bam", "-o", bam_sorted, best["sam_path"]], check=True)
        subprocess.run(["samtools", "index", bam_sorted], check=True)

        reference_coverage_percent = calculate_reference_breadth_coverage_percent(
            bam_file=bam_sorted,
            ref_name=best["ref_sn"],
            ref_length=best["ref_length"]
        )

        hbv_system_stats = empty_hbv_system_stats()
        hbv_flags = {
            "high_ok": True,
            "medium_ok": True,
            "warning": "",
            "balance_ratio": 0,
        }

        if virus_type == "Adenovirus" and detected_target_start and detected_target_end:
            depth_region = (detected_target_start, detected_target_end)
            coverage_type = "Detected Adenovirus Target/Amplicon Coverage"
            expected_target_length = expected_target_length_for_qc(virus_type, detected_target_length)

            depth_stats = calculate_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                region=depth_region
            )

            target_or_genome_coverage_percent = 100.0 if detected_target_length > 0 else 0.0

        elif virus_type == "Hepatitis B Virus":
            depth_region = None
            coverage_type = "Approximate HBV Whole-Genome Coverage"
            expected_target_length = HBV_EXPECTED_GENOME_LENGTH

            depth_stats = calculate_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                region=None
            )

            target_or_genome_coverage_percent = reference_coverage_percent

            hbv_system_stats = calculate_hbv_system_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"]
            )
            hbv_flags = hbv_system_qc_flags(hbv_system_stats)

        elif virus_type == "Hepatitis C Virus":
            depth_region = (
                (detected_target_start, detected_target_end)
                if detected_target_start and detected_target_end
                else None
            )
            coverage_type = "Detected HCV Target/Amplicon Coverage"
            expected_target_length = detected_target_length or 0

            depth_stats = calculate_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                region=depth_region
            )

            target_or_genome_coverage_percent = calculate_hcv_amplicon_breadth_coverage_percent(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                start_1based=detected_target_start,
                end_1based=detected_target_end,
            )

        else:
            depth_region = None
            coverage_type = "Reference Coverage"
            expected_target_length = expected_target_length_for_qc(virus_type, detected_target_length)

            depth_stats = calculate_depth_stats(
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                region=None
            )

            target_or_genome_coverage_percent = reference_coverage_percent

        print(
            f"Best reference for {input_file}: {best['ref_sn']} "
            f"with {best['mapped']} mapped reads."
        )

        if virus_type == "Adenovirus":
            print(
                f"Detected adenovirus target region: "
                f"{detected_target_start}-{detected_target_end} "
                f"({detected_target_length} bp)."
            )
        elif virus_type == "Hepatitis C Virus":
            print(
                f"Detected HCV target amplicon: "
                f"{detected_target_start}-{detected_target_end} "
                f"({detected_target_length} bp)."
            )

        print(
            f"{coverage_type}: {target_or_genome_coverage_percent}% | "
            f"Median depth: {depth_stats.get('median_depth', 0)}x"
        )

        if virus_type == "Hepatitis B Virus":
            print(
                "HBV unique-region QC: "
                f"System 1 median {hbv_system_stats['system1_unique'].get('median_depth', 0)}x, "
                f"System 2 median {hbv_system_stats['system2_unique'].get('median_depth', 0)}x | "
                f"{hbv_flags.get('warning', '')}"
            )

        draft_consensus = os.path.join(sample_dir, f"{base_name}.draft.fasta")
        subprocess.run(["samtools", "consensus", "-f", "fasta", "-o", draft_consensus, bam_sorted], check=True)

        if not (os.path.exists(draft_consensus) and os.path.getsize(draft_consensus) > 0):
            print(f"Draft consensus missing/empty for {base_name}; skipping Medaka.")

            fasta_stats = {
                "consensus_length": 0,
                "n_count": 0,
                "n_percent": 100.0,
                "non_n_length": 0,
            }

            qc_verdict, qc_reason, consensus_recovery_percent = determine_qc_verdict(
                virus_type=virus_type,
                mapped_reads=best["mapped"],
                coverage_percent=target_or_genome_coverage_percent,
                median_depth=depth_stats.get("median_depth", 0),
                n_percent=fasta_stats["n_percent"],
                consensus_length=fasta_stats["consensus_length"],
                expected_length=expected_target_length,
                hbv_system_high_ok=hbv_flags.get("high_ok", True),
                hbv_system_medium_ok=hbv_flags.get("medium_ok", True),
                hbv_system_warning=hbv_flags.get("warning", "")
            )

            row = {
                "sample": base_name,
                "virus_type": virus_type,
                "best_reference": best["ref_sn"] or "NA",
                "mapped_reads": best["mapped"],
                "reference_length": best["ref_length"],
                "detected_target_start": detected_target_start or "",
                "detected_target_end": detected_target_end or "",
                "detected_target_length": detected_target_length,
                "expected_target_length": expected_target_length,
                "coverage_type": coverage_type,
                "target_or_genome_coverage_percent": target_or_genome_coverage_percent,
                "mean_depth": depth_stats.get("mean_depth", 0),
                "median_depth": depth_stats.get("median_depth", 0),
                "min_depth": depth_stats.get("min_depth", 0),
                "positions_ge_10x_percent": depth_stats.get("positions_ge_10x_percent", 0),
                "positions_ge_30x_percent": depth_stats.get("positions_ge_30x_percent", 0),
                "consensus_length": fasta_stats["consensus_length"],
                "n_count": fasta_stats["n_count"],
                "n_percent": fasta_stats["n_percent"],
                "consensus_recovery_percent": consensus_recovery_percent,
                "qc_verdict": qc_verdict,
                "qc_reason": qc_reason,
                "bam_file": bam_sorted,
                "input_fastq": input_file,
                "consensus_file": draft_consensus,
                "reference_file": best["ref_path"],
                "medaka_model": "Ej körd",
            }

            if virus_type == "Hepatitis B Virus":
                row = add_hbv_system_stats_to_row(row, hbv_system_stats, hbv_flags)
                row.update(analyse_hbv_mutations(base_name, bam_sorted, best["ref_sn"], best["ref_path"]))

            if virus_type == "Hepatitis C Virus":
                row.update(analyse_hcv_resistance(
                    sample=base_name,
                    input_fastq=input_file,
                    bam_file=bam_sorted,
                    ref_name=best["ref_sn"],
                    ref_path=best["ref_path"],
                    mapped_reads_1a=hcv_best_mapped_by_subtype.get("1A", 0),
                    mapped_reads_1b=hcv_best_mapped_by_subtype.get("1B", 0),
                    median_nonzero_coverage=calculate_median_nonzero_depth(
                        bam_sorted, best["ref_sn"], depth_region
                    ),
                ))

            result_rows.append(row)

            safe_remove(best["sam_path"])
            if tmp_ref != best["ref_path"]:
                safe_remove(tmp_ref)
            return

        iub_consensus_path = generate_iub_consensus_from_bam(
            bam_file=bam_sorted,
            ref_name=best["ref_sn"],
            ref_path=best["ref_path"],
            output_folder=sample_dir,
            sample_name=base_name,
            variant_frequency=HBV_MIN_VARIANT_FREQUENCY if virus_type == "Hepatitis B Virus" else 0.15,
        )

        medaka_result = generate_consensus(input_file, draft_consensus, selected_medaka_model, sample_dir, base_name)
        final_consensus_path_for_analysis = medaka_result if medaka_result else draft_consensus

        if medaka_result:
            final_consensus_list.append(medaka_result)

        fasta_stats = fasta_sequence_stats(final_consensus_path_for_analysis)
        if virus_type == "Hepatitis C Virus":
            fasta_stats = hcv_amplicon_consensus_stats(
                final_consensus_path_for_analysis,
                detected_target_start,
                detected_target_end,
            )

        qc_verdict, qc_reason, consensus_recovery_percent = determine_qc_verdict(
            virus_type=virus_type,
            mapped_reads=best["mapped"],
            coverage_percent=target_or_genome_coverage_percent,
            median_depth=depth_stats.get("median_depth", 0),
            n_percent=fasta_stats["n_percent"],
            consensus_length=fasta_stats["consensus_length"],
            expected_length=expected_target_length,
            hbv_system_high_ok=hbv_flags.get("high_ok", True),
            hbv_system_medium_ok=hbv_flags.get("medium_ok", True),
            hbv_system_warning=hbv_flags.get("warning", "")
        )

        row = {
            "sample": base_name,
            "virus_type": virus_type,
            "best_reference": best["ref_sn"] or "NA",
            "mapped_reads": best["mapped"],
            "reference_length": best["ref_length"],
            "detected_target_start": detected_target_start or "",
            "detected_target_end": detected_target_end or "",
            "detected_target_length": detected_target_length,
            "expected_target_length": expected_target_length,
            "coverage_type": coverage_type,
            "target_or_genome_coverage_percent": target_or_genome_coverage_percent,
            "mean_depth": depth_stats.get("mean_depth", 0),
            "median_depth": depth_stats.get("median_depth", 0),
            "min_depth": depth_stats.get("min_depth", 0),
            "positions_ge_10x_percent": depth_stats.get("positions_ge_10x_percent", 0),
            "positions_ge_30x_percent": depth_stats.get("positions_ge_30x_percent", 0),
            "consensus_length": fasta_stats["consensus_length"],
            "n_count": fasta_stats["n_count"],
            "n_percent": fasta_stats["n_percent"],
            "consensus_recovery_percent": consensus_recovery_percent,
            "qc_verdict": qc_verdict,
            "qc_reason": qc_reason,
            "bam_file": bam_sorted,
            "input_fastq": input_file,
            "consensus_file": final_consensus_path_for_analysis,
            "iub_consensus_file": iub_consensus_path or "",
            "reference_file": best["ref_path"],
            "medaka_model": selected_medaka_model,
        }

        if virus_type == "Hepatitis B Virus":
            row = add_hbv_system_stats_to_row(row, hbv_system_stats, hbv_flags)
            row.update(analyse_hbv_mutations(base_name, bam_sorted, best["ref_sn"], best["ref_path"]))

        if virus_type == "Hepatitis C Virus":
            row.update(analyse_hcv_resistance(
                sample=base_name,
                input_fastq=input_file,
                bam_file=bam_sorted,
                ref_name=best["ref_sn"],
                ref_path=best["ref_path"],
                mapped_reads_1a=hcv_best_mapped_by_subtype.get("1A", 0),
                mapped_reads_1b=hcv_best_mapped_by_subtype.get("1B", 0),
                median_nonzero_coverage=calculate_median_nonzero_depth(
                        bam_sorted, best["ref_sn"], depth_region
                    ),
            ))

        result_rows.append(row)

        print(f"QC Verdict for {base_name}: {qc_verdict} ({qc_reason})")

        safe_remove(best["sam_path"])
        safe_remove(draft_consensus)

        if tmp_ref != best["ref_path"]:
            safe_remove(tmp_ref)

    except subprocess.CalledProcessError as e:
        print(f"Subprocess error while processing {input_file}: {e}")
    except Exception as e:
        print(f"Unexpected error in processing sample {input_file}: {e}")

def concatenate_fasta_files(fasta_files, output_file):
    if not fasta_files:
        print("No consensus FASTA files to concatenate.")
        return False

    try:
        ensure_dir(os.path.dirname(output_file), must_be_writable=True)
        with open(output_file, "w", encoding="utf-8") as outfile:
            for fasta in fasta_files:
                if os.path.isfile(fasta) and os.path.getsize(fasta) > 0:
                    with open(fasta, "r", encoding="utf-8", errors="replace") as infile:
                        outfile.write(infile.read())
                else:
                    print(f"Skipping empty or missing FASTA file: {fasta}")

        print(f"Concatenated FASTA file saved to {output_file}")
        return True

    except Exception as e:
        print(f"Error concatenating FASTA files: {e}")
        return False

def split_fasta(output_fasta, split_folder):
    ensure_dir(split_folder, must_be_writable=True)

    with open(output_fasta, "r", encoding="utf-8", errors="replace") as f:
        current = None
        for line in f:
            if line.startswith(">"):
                if current:
                    current.close()
                header = line.strip()[1:].replace(" ", "_")
                current = open(os.path.join(split_folder, f"{header}.fasta"), "w", encoding="utf-8")
            if current:
                current.write(line)

        if current:
            current.close()

    print(f"Split FASTA sequences saved to {split_folder}")

# NCBI pipeline (Not active, work in progress,decapritate)

def run_ncbi_database_script(input_folder):
    try:
        base = app_base_dir()

        ncbi_script_path = path_exists_any(
            os.path.join(base, "AdenoPipeline", "BLASTscript.py"),
            "/Volumes/Newt-Vault/AdenoPipeline/BLASTscript.py",
            "/Volumes/NewtVault/AdenoPipeline/BLASTscript.py"
        )

        if not ncbi_script_path:
            print("NCBI script not found. Put BLASTscript.py in ./AdenoPipeline/ next to this script or use the mounted volume path.")
            return

        output_excel = os.path.join(input_folder, "BLAST_results.xlsx")
        invalid_files_csv = os.path.join(input_folder, "invalid_files.csv")

        blastplus_path = path_exists_any(
            os.path.join(base, "blastplus"),
            "/Users/xsjoak/Documents/Entero/blastplus"
        )

        if not blastplus_path:
            print("BLAST+ folder not found. Put it in ./blastplus next to the script or adjust the path.")
            return

        subprocess.run([
            sys.executable,
            ncbi_script_path,
            input_folder,
            blastplus_path,
            output_excel,
            invalid_files_csv
        ], check=True)

        print(f"NCBI database script executed successfully. Results saved to {output_excel}")

    except subprocess.CalledProcessError as e:
        print(f"Error running the NCBI database script: {e}")
    except Exception as e:
        print(f"Unexpected error while running the NCBI script: {e}")

#Pipeline runners

def run_full_pipeline(input_folder, output_folder, reference_folder, medaka_model, virus_type, run_ncbi=False, auto_detect_models=False):
    print("\nRunning full pipeline...\n")
    print(f"Aurora script file: {os.path.abspath(__file__)}")
    print(f"HCV script version: {HCV_SCRIPT_VERSION if virus_type == 'Hepatitis C Virus' else 'not HCV run'}")
    print(f"Virus type: {virus_type}")
    print(f"Input: {input_folder}")
    print(f"Output: {output_folder}")
    print(f"Reference folder: {reference_folder}")
    print(f"Medaka fallback model: {medaka_model}")
    print(f"Auto-detect Medaka models: {auto_detect_models}")
    print("Keep BAM files: True\n")

    try:
        check_required_tools()
    except Exception as e:
        print(f"[ERROR] {e}")
        return

    ensure_dir(output_folder, must_be_writable=True)
    try:
        run_date, run_number, run_tag, run_marker_path = allocate_run_identity(output_folder)
    except FileExistsError:
        # Extremely unlikely collision (for example two launches at the same time).
        run_date, run_number, run_tag, run_marker_path = allocate_run_identity(output_folder)
    print(f"Run identifier: {run_tag}")
    print(f"Run marker: {run_marker_path}\n")

    input_files = find_files_recursively(input_folder, is_real_fastq_gz)
    reference_files = [
        os.path.join(reference_folder, f)
        for f in os.listdir(reference_folder)
        if is_real_fasta(os.path.join(reference_folder, f))
    ]

    if not input_files:
        print("No .fastq.gz input files found.")
        return

    if not reference_files:
        print("No reference .fa/.fasta files found.")
        return

    result_rows = []
    final_consensus_list = []
    samples_output_folder = run_samples_folder(output_folder, run_tag)
    ensure_dir(samples_output_folder, must_be_writable=True)

    for input_file in input_files:
        print(f"\n--- Processing sample: {input_file} ---")
        process_sample(
            input_file=input_file,
            reference_files=reference_files,
            output_folder=samples_output_folder,
            medaka_model=medaka_model,
            virus_type=virus_type,
            final_consensus_list=final_consensus_list,
            result_rows=result_rows,
            auto_detect_models=auto_detect_models
        )

    if result_rows:
        if virus_type == "Hepatitis C Virus":
            log_hcv_resistance_results(result_rows, output_folder, run_tag)
            log_reference_usage(result_rows, output_folder, virus_type, run_tag)
        elif virus_type == "Hepatitis B Virus":
            log_hbv_mutation_results(result_rows, output_folder, run_tag)
            log_reference_usage(result_rows, output_folder, virus_type, run_tag)
        else:
            log_reference_usage(result_rows, output_folder, virus_type, run_tag)
    else:
        print("No reference usage data recorded.")

    if result_rows:
        export_run_consensus_files(result_rows, output_folder, run_tag)

    safe_virus_name = virus_type.replace(" ", "_")
    final_fasta_path = os.path.join(
        output_folder,
        f"{safe_virus_name}_{run_tag}_{len(final_consensus_list)}_consensus.fasta"
    )

    if final_consensus_list:
        success = concatenate_fasta_files(final_consensus_list, final_fasta_path)
        if success and run_ncbi:
            split_folder = os.path.join(samples_output_folder, "split_fasta")
            split_fasta(final_fasta_path, split_folder)
            run_ncbi_database_script(split_folder)
    else:
        print("No consensus files produced; skipping concatenation and optional NCBI step.")

# GUI PySide6

MEDAKA_MODEL_CHOICES = [
    "Auto-detect from FASTQ header",
    "r1041_e82_400bps_sup_v5.2.0",
    "r1041_e82_400bps_hac_v5.0.0",
    "r1041_e82_400bps_sup_v4.3.0",
    "r1041_e82_400bps_sup_v5.0.0",
    "r1041_e82_400bps_sup_v4.2.0",
    "r1041_e82_400bps_sup_v3.5.1",
    "r1041_e82_400bps_sup_variant_v4.2.0",
    "r1041_e82_400bps_hac_v4.2.0",
    "r1041_e82_400bps_hac_variant_v4.2.0",
    "r1041_e82_400bps_fast_g632",
    "r1041_e82_400bps_fast_variant_g632",
    "r1041_e82_260bps_sup_v4.1.0",
    "r1041_e82_260bps_sup_variant_v4.1.0",
    "r1041_e82_260bps_hac_v4.1.0",
    "r1041_e82_260bps_hac_variant_v4.1.0",
    "r1041_e82_260bps_fast_g632",
    "r1041_e82_260bps_fast_variant_g632",
    "r941_min_fast_g507",
    "r941_min_fast_variant_g507",
    "r941_min_hac_g507",
    "r941_min_hac_variant_g507",
    "r941_min_sup_g507",
    "r941_min_sup_variant_g507",
    "r941_prom_fast_g507",
    "r941_prom_fast_variant_g507",
    "r941_prom_hac_g507",
    "r941_prom_hac_variant_g507",
    "r941_prom_sup_g507",
    "r941_prom_sup_variant_g507",
]


class ConsoleStream(QObject):
    """Small file-like object that forwards print output to the GUI."""

    text_written = Signal(str)

    def write(self, text):
        if text:
            self.text_written.emit(str(text))
        return len(text or "")

    def flush(self):
        pass

    def isatty(self):
        return False


class NativeOutputCapture:
    """Capture output written directly to file descriptors 1 and 2.

    Redirecting only ``sys.stdout`` does not catch output from programs such as
    minimap2, samtools, and Medaka. On POSIX systems this class temporarily
    redirects the process stdout/stderr file descriptors into a pipe drained by
    a small reader thread, then forwards that text to the Qt log widget.
    """

    def __init__(self, console_stream):
        self.console_stream = console_stream
        self._saved_stdout_fd = None
        self._saved_stderr_fd = None
        self._read_fd = None
        self._reader_thread = None
        self._active = False

    def start(self):
        if os.name != "posix" or self._active:
            return

        read_fd, write_fd = os.pipe()
        self._saved_stdout_fd = os.dup(1)
        self._saved_stderr_fd = os.dup(2)
        self._read_fd = read_fd

        os.dup2(write_fd, 1)
        os.dup2(write_fd, 2)
        os.close(write_fd)

        self._active = True
        self._reader_thread = threading.Thread(
            target=self._read_loop,
            name="AuroraNativeOutputReader",
            daemon=True,
        )
        self._reader_thread.start()

    def _read_loop(self):
        try:
            while True:
                chunk = os.read(self._read_fd, 4096)
                if not chunk:
                    break
                self.console_stream.write(chunk.decode("utf-8", errors="replace"))
        except OSError:
            pass
        finally:
            if self._read_fd is not None:
                try:
                    os.close(self._read_fd)
                except OSError:
                    pass
                self._read_fd = None

    def stop(self):
        if not self._active:
            return

        try:
            os.dup2(self._saved_stdout_fd, 1)
            os.dup2(self._saved_stderr_fd, 2)
        finally:
            for saved_fd in (self._saved_stdout_fd, self._saved_stderr_fd):
                if saved_fd is not None:
                    try:
                        os.close(saved_fd)
                    except OSError:
                        pass
            self._saved_stdout_fd = None
            self._saved_stderr_fd = None
            self._active = False

        if self._reader_thread is not None:
            self._reader_thread.join(timeout=2.0)
            self._reader_thread = None


class PipelineWorker(QObject):
    finished = Signal()
    failed = Signal(str)

    def __init__(self, parameters, console_stream):
        super().__init__()
        self.parameters = parameters
        self.console_stream = console_stream

    @Slot()
    def run(self):
        old_stdout = sys.stdout
        old_stderr = sys.stderr
        native_capture = NativeOutputCapture(self.console_stream)

        # Python print() output goes directly to the GUI. Native/subprocess
        # output is captured at the file-descriptor level below.
        sys.stdout = self.console_stream
        sys.stderr = self.console_stream
        native_capture.start()
        try:
            # Fail visibly in the GUI instead of letting run_full_pipeline()
            # print an error and then look like a successful completion.
            check_required_tools()
            run_full_pipeline(**self.parameters)
        except Exception:
            import traceback
            traceback.print_exc()
            self.failed.emit("The pipeline stopped because of an unexpected error.")
        else:
            self.finished.emit()
        finally:
            native_capture.stop()
            sys.stdout = old_stdout
            sys.stderr = old_stderr


class AuroraWindow(QMainWindow):
    """Controller for the window designed in Qt Widgets Designer."""

    def __init__(self):
        super().__init__()
        self.worker_thread = None
        self.worker = None

        self.ui = Ui_MainWindow()
        self.ui.setupUi(self)

        # Respect the compact 573 x 795 geometry from Qt Designer instead of
        # overriding it with a wide desktop layout.
        self.resize(590, 820)
        self.setMinimumSize(540, 700)
        self.setMaximumWidth(680)

        # Short aliases keep the pipeline/controller code readable while using
        # the exact object names already present in the Designer file.
        self.virus_type_combo = self.ui.virus_combo
        self.input_folder_edit = self.ui.input_edit
        self.output_folder_edit = self.ui.output_edit
        self.reference_folder_edit = self.ui.reference_edit
        self.medaka_model_combo = self.ui.medaka_combo
        self.run_ncbi_checkbox = self.ui.ncbi_checkbox
        self.progress_bar = self.ui.progress
        self.console = self.ui.log
        self.run_button = self.ui.run_button
        self.clear_button = self.ui.clear_button
        self.check_button = self.ui.check_button

        # Let the form fields use the available compact width without imposing
        # the large minimum widths that caused the stretched window.
        self.ui.settingsForm.setFieldGrowthPolicy(
            QFormLayout.FieldGrowthPolicy.AllNonFixedFieldsGrow
        )
        self.ui.settingsBox.setMinimumWidth(0)
        self.ui.settingsBox.setMaximumWidth(16777215)

        for field_widget in (
            self.ui.inputFieldWidget,
            self.ui.outputFieldWidget,
            self.ui.referenceFieldWidget,
        ):
            field_widget.setMinimumWidth(0)
            field_widget.setSizePolicy(
                QSizePolicy.Policy.Expanding,
                QSizePolicy.Policy.Preferred,
            )

        for line_edit in (
            self.input_folder_edit,
            self.output_folder_edit,
            self.reference_folder_edit,
        ):
            line_edit.setMinimumWidth(150)
            line_edit.setSizePolicy(
                QSizePolicy.Policy.Expanding,
                QSizePolicy.Policy.Fixed,
            )

        for browse_button in (
            self.ui.input_browse_button,
            self.ui.output_browse_button,
            self.ui.reference_browse_button,
        ):
            browse_button.setFixedWidth(86)

        for button in (self.check_button, self.clear_button, self.run_button):
            button.setMinimumWidth(0)
            button.setMaximumWidth(180)

        # Keep the lower part compact while still leaving enough room for logs.
        self.console.setMinimumHeight(105)
        self.console.setMaximumHeight(150)
        self.progress_bar.setMaximumWidth(16777215)
        self.ui.buttonSpacer.changeSize(
            0,
            0,
            QSizePolicy.Policy.Minimum,
            QSizePolicy.Policy.Fixed,
        )
        self.ui.outerLayout.invalidate()

        # Keep the original full-width banner size and add only a soft glow.
        # The artwork remains controlled by the Designer label rather than being
        # redrawn into a smaller rounded card.
        self.ui.bannerLabel.setScaledContents(True)
        self.ui.bannerLabel.setMinimumHeight(150)
        self.ui.bannerLabel.setMaximumHeight(170)
        self.ui.bannerLabel.setSizePolicy(
            QSizePolicy.Policy.Expanding,
            QSizePolicy.Policy.Fixed,
        )
        self._create_banner_effect()

        self.status_label = QLabel("Ready", self)
        self.ui.statusBar.addWidget(self.status_label, 1)

        self.console_stream = ConsoleStream(self)
        self.console_stream.text_written.connect(self.append_console_text)

        # Start with an explicit prompt instead of silently selecting a virus.
        virus_prompt = "Choose Virus..."
        virus_prompt_index = self.virus_type_combo.findText(virus_prompt)
        if virus_prompt_index < 0:
            self.virus_type_combo.insertItem(0, virus_prompt)
            virus_prompt_index = 0
        self.virus_type_combo.setCurrentIndex(virus_prompt_index)

        self.ui.input_browse_button.clicked.connect(
            lambda: self.choose_directory(self.input_folder_edit)
        )
        self.ui.output_browse_button.clicked.connect(
            lambda: self.choose_directory(self.output_folder_edit)
        )
        self.ui.reference_browse_button.clicked.connect(
            lambda: self.choose_directory(self.reference_folder_edit)
        )
        self.run_button.clicked.connect(self.start_pipeline)
        self.clear_button.clicked.connect(self.console.clear)
        self.check_button.clicked.connect(self.check_dependencies)

        # Keep the optional automatic model detection available even if the
        # current Designer combobox only lists concrete Medaka models.
        auto_model_text = "Auto-detect from FASTQ header"
        auto_model_index = self.medaka_model_combo.findText(auto_model_text)
        if auto_model_index < 0:
            self.medaka_model_combo.insertItem(0, auto_model_text)
            auto_model_index = 0
        self.medaka_model_combo.setCurrentIndex(auto_model_index)

        # The progress duration is unknown, so it is switched to busy mode
        # only while the pipeline is running.
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setTextVisible(False)
        self.progress_bar.setVisible(False)

        self._create_progress_glow()

        banner_pixmap = self.ui.bannerLabel.pixmap()
        if banner_pixmap is None or banner_pixmap.isNull():
            self.append_console_text(
                "[GUI WARNING] The banner resource did not load. "
                "Rebuild resources_rc.py and verify the :/ resource path.\n"
            )

    def _create_banner_effect(self):
        """Apply a full-size, softly pulsing aurora glow to the banner."""
        # Keep a crisp teal edge around the original Designer-sized image.
        self.ui.bannerLabel.setStyleSheet(
            "QLabel#bannerLabel {"
            " background: #08162c;"
            " border: 1px solid #2dd4bf;"
            " padding: 0;"
            "}"
        )

        self.banner_glow = QGraphicsDropShadowEffect(self.ui.bannerLabel)
        self.banner_glow.setBlurRadius(24)
        self.banner_glow.setOffset(0, 0)
        self.banner_glow.setColor(QColor(34, 211, 238, 150))
        self.ui.bannerLabel.setGraphicsEffect(self.banner_glow)

        self.banner_glow_animation = QPropertyAnimation(
            self.banner_glow,
            b"blurRadius",
            self,
        )
        self.banner_glow_animation.setDuration(2800)
        self.banner_glow_animation.setKeyValueAt(0.0, 18.0)
        self.banner_glow_animation.setKeyValueAt(0.5, 32.0)
        self.banner_glow_animation.setKeyValueAt(1.0, 18.0)
        self.banner_glow_animation.setEasingCurve(QEasingCurve.Type.InOutSine)
        self.banner_glow_animation.setLoopCount(-1)
        self.banner_glow_animation.start()

    @Slot()
    def check_dependencies(self):
        missing = [
            tool
            for tool in ("minimap2", "samtools", "medaka_consensus")
            if which_or_none(tool) is None
        ]
        if missing:
            message = "Missing tools on PATH: " + ", ".join(missing)
            self.status_label.setText("Dependencies missing")
            self.append_console_text(f"[DEPENDENCY ERROR] {message}\n")
            QMessageBox.warning(self, "Missing dependencies", message)
            return

        message = "minimap2, samtools, and medaka_consensus were found on PATH."
        if pysam is None:
            message += "\n\npysam is not installed; HCV minor-variant and IUB features may be unavailable."
        self.status_label.setText("Dependencies checked")
        self.append_console_text(f"[DEPENDENCIES] {message}\n")
        QMessageBox.information(self, "Dependencies", message)

    @Slot()
    def choose_directory(self, target_edit):
        start = target_edit.text().strip() or app_base_dir()
        selected = QFileDialog.getExistingDirectory(self, "Select folder", start)
        if selected:
            target_edit.setText(selected)

    def _create_progress_glow(self):
        self.progress_glow = QGraphicsDropShadowEffect(self.progress_bar)
        self.progress_glow.setBlurRadius(20)
        self.progress_glow.setOffset(0, 0)
        self.progress_glow.setColor(QColor(34, 211, 238, 190))
        self.progress_bar.setGraphicsEffect(self.progress_glow)

        self.glow_animation = QPropertyAnimation(
            self.progress_glow,
            b"blurRadius",
            self,
        )
        self.glow_animation.setDuration(1500)
        self.glow_animation.setKeyValueAt(0.0, 14.0)
        self.glow_animation.setKeyValueAt(0.5, 32.0)
        self.glow_animation.setKeyValueAt(1.0, 14.0)
        self.glow_animation.setEasingCurve(QEasingCurve.Type.InOutSine)
        self.glow_animation.setLoopCount(-1)

    @Slot(str)
    def append_console_text(self, text):
        self.console.moveCursor(QTextCursor.MoveOperation.End)
        self.console.insertPlainText(text)
        self.console.ensureCursorVisible()

    def _validated_parameters(self):
        input_folder = self.input_folder_edit.text().strip()
        output_folder = self.output_folder_edit.text().strip()
        reference_folder = self.reference_folder_edit.text().strip()

        selected_virus = self.virus_type_combo.currentText().strip()

        missing = []
        if not selected_virus or selected_virus == "Choose Virus...":
            missing.append("Virus type")
        if not input_folder:
            missing.append("Input folder")
        if not output_folder:
            missing.append("Output folder")
        if not reference_folder:
            missing.append("Reference folder")
        if missing:
            raise ValueError("Please select: " + ", ".join(missing))

        if not os.path.isdir(input_folder):
            raise ValueError("The input folder does not exist.")
        if not os.path.isdir(reference_folder):
            raise ValueError("The reference folder does not exist.")

        selected_model = self.medaka_model_combo.currentText()
        auto_detect_models = selected_model == "Auto-detect from FASTQ header"
        fallback_model = (
            "r1041_e82_400bps_sup_v4.2.0"
            if auto_detect_models
            else selected_model
        )

        return {
            "input_folder": input_folder,
            "output_folder": output_folder,
            "reference_folder": reference_folder,
            "medaka_model": fallback_model,
            "virus_type": selected_virus,
            "run_ncbi": self.run_ncbi_checkbox.isChecked(),
            "auto_detect_models": auto_detect_models,
        }

    def _set_input_controls_enabled(self, enabled):
        for widget in (
            self.virus_type_combo,
            self.ui.mode_combo,
            self.input_folder_edit,
            self.output_folder_edit,
            self.reference_folder_edit,
            self.medaka_model_combo,
            self.run_ncbi_checkbox,
            self.ui.input_browse_button,
            self.ui.output_browse_button,
            self.ui.reference_browse_button,
            self.check_button,
        ):
            widget.setEnabled(enabled)

    @Slot()
    def start_pipeline(self):
        if self.worker_thread is not None:
            return

        try:
            parameters = self._validated_parameters()
        except ValueError as exc:
            QMessageBox.warning(self, "Missing or invalid settings", str(exc))
            return

        self.console.clear()
        self._set_input_controls_enabled(False)
        self.run_button.setEnabled(False)
        self.clear_button.setEnabled(False)
        self.status_label.setText("Running pipeline...")
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setTextVisible(False)
        self.progress_bar.setVisible(True)
        self.glow_animation.start()

        self.worker_thread = QThread(self)
        self.worker = PipelineWorker(parameters, self.console_stream)
        self.worker.moveToThread(self.worker_thread)

        self.worker_thread.started.connect(self.worker.run)
        self.worker.finished.connect(self.pipeline_finished)
        self.worker.failed.connect(self.pipeline_failed)
        self.worker.finished.connect(self.worker_thread.quit)
        self.worker.failed.connect(self.worker_thread.quit)
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.failed.connect(self.worker.deleteLater)
        self.worker_thread.finished.connect(self._thread_finished)
        self.worker_thread.finished.connect(self.worker_thread.deleteLater)
        self.worker_thread.start()

    @Slot()
    def pipeline_finished(self):
        self.status_label.setText("Pipeline finished")
        self.append_console_text("\n[GUI] Pipeline finished.\n")
        self._stop_busy_indicator()

    @Slot(str)
    def pipeline_failed(self, message):
        self.status_label.setText("Pipeline failed")
        self.append_console_text(f"\n[GUI ERROR] {message}\n")
        self._stop_busy_indicator()
        QMessageBox.critical(self, "Pipeline error", message)

    def _stop_busy_indicator(self):
        self.glow_animation.stop()
        self.progress_bar.setVisible(False)
        self.run_button.setEnabled(True)
        self.clear_button.setEnabled(True)
        self._set_input_controls_enabled(True)
        # Analysis mode is intentionally fixed in the Designer file.
        self.ui.mode_combo.setEnabled(False)

    @Slot()
    def _thread_finished(self):
        self.worker = None
        self.worker_thread = None

    def closeEvent(self, event):
        if self.worker_thread is not None and self.worker_thread.isRunning():
            QMessageBox.information(
                self,
                "Pipeline still running",
                "The pipeline is still running. Close the window after it finishes.",
            )
            event.ignore()
            return
        event.accept()

def main():
    app = QApplication(sys.argv)
    app.setApplicationName("Aurora")
    window = AuroraWindow()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
